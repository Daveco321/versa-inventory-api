"""
inventory_aging.py: the Inventory Aging tool's data (Oct 5 2026).

Pure functions only: no Flask, no S3, no network. app.py feeds it the hourly
Inventory_ATS.xlsx bytes it already downloads, and the local trend backfill
imports this same file, so the live page and the history use one code path.

The workbook carries lot detail next to the ATS sheet. A raw sheet holds the
receive date and a *_INVENTORY sheet holds the platform SKU, container, lot
text and units, row for row:

    JTW  + JTW_INVENTORY   date 'Inventory Date' (since Mar 23 2026) or
                           'Date Received' (older files), units 'Quantity'
    TR   + TR_INVENTORY    date 'Receipt Date', SKU 'SKU', units 'QTY'
    DCW  + DCW_INVENTORY   date 'Receipt Date', SKU 'Style', units 'QTY'

Every column is looked up by header name (first occurrence: the *_INVENTORY
sheets repeat names in a pivot block on the right). Row i of a raw sheet pairs
with row i of its *_INVENTORY sheet (trailing rows with no style or SKU are not
lots and are dropped first). When the counts differ the lot detail of
that warehouse is unusable and reported, never guessed. One exception, verified
row by row: both JTW formats list every receipt including shipped-out lines,
and JTW_INVENTORY keeps only the lines with a positive 'Ballance' (old format)
or a positive 'OnHand' / 'Received' (new format; 'OnHand' when the listing has
that column). Those lines pair one to one (style, container and cartons must
all agree, or JTW is unusable for that file).

Public API (the trend backfill depends on these signatures):
    parse_workbook(source, sku_rows=None, is_sized=None, tick=None) -> parsed
    restore_map_from_jtw_sheet(source, source_label='', built_at=None) -> restore map
    apply_corrections(lots, restore_map, today, jtw_format='new') -> (lots, restore_stats)
    analyze(parsed, restore_map, today) -> analysis
    trend_record(analysis, src='live') -> record | None
    build_payload(analysis, is_admin, wh_ats=None, rev=None, as_of=None, built_at=None) -> dict
    validate_seed(body) -> (kind, value, error)
    merge_trend(doc, series, dropped, mode, updated_at) / upsert_trend(doc, record, updated_at)
"""

import io
import math
import re
from datetime import date, datetime, timedelta

# ── constants (keys and labels are part of the API contract) ────────────────
BUCKETS = (('lt1m', 'Under 1 month', 0), ('m1_3', '1 to 3 months', 30), ('m3_6', '3 to 6 months', 90),
           ('m6_12', '6 to 12 months', 180), ('y1_2', '1 to 2 years', 365), ('y2_3', '2 to 3 years', 730),
           ('y3p', '3 years or more', 1095))
# Nuri's report: 360-day years, "oldest lot" per SKU.
NURI_BUCKETS = (('lt1m', '< 1 Month', 0), ('m1', '> 1 Month', 30), ('m3', '> 3 Months', 90),
                ('m6', '> 6 Months', 180), ('y1', '> 1 Year', 360), ('y2', '> 2 Years', 720),
                ('y3', '> 3 Years', 1080))
WAREHOUSES = ('JTW', 'TR', 'DCW')
LOT_COLS = ['sku', 'wh', 'units', 'date', 'container', 'lot', 'flags', 'origDate']
SKU_COLS = ['sku', 'base', 'brand', 'prefix', 'jtw', 'tr', 'dcw', 'qa', 'nj', 'abfi', 'committed', 'allocated',
            'receiveDate']
BASE_COLS = ['base', 'brand', 'whAts']
TREND_KEYS = ('d', 'units', 'wavg', 'brand', 'wh', 'prefix', 'nuri', 'restored', 'countDate', 'src')

MIN_VALID_DATE = date(2010, 1, 1)
COUNT_DATES = (date(2026, 3, 18), date(2026, 3, 20))     # JTW's 3/18/2026 physical count re-stamp
TREND_TOP_PREFIXES = 15
SEED_MAX_BYTES = 5 * 1024 * 1024
TREND_MAX_RECORDS = 5000

CONTAINER_RE = re.compile(r'^[A-Z]{4}\d{7}$')
RESET_RE = re.compile(r'TRANSFER|TRUCK|CARSON|INVENTORY ?20\d\d|\bINV ?\d\d|CYCLE COUNT|\bRMA\b|RESTACK|AISLE|INVADJ')
_PREFIX_RE = re.compile(r'^[A-Z]{4}[A-Z0-9]{2}\d{3}')
_ISO_RE = re.compile(r'^\d{4}-\d{2}-\d{2}$')
_EMPTY_TEXT = ('', 'nan', 'none', 'n/a', 'nat', 'null')

# Raw/inventory sheet layout per warehouse. Header names: alternatives in order.
_SHEETS = {
    'JTW': {'raw': 'JTW', 'inv': 'JTW_INVENTORY', 'sku': ('Style',), 'qty': ('Quantity',),
            'container': ('Container #', 'Container'), 'lot': ()},
    'TR': {'raw': 'TR', 'inv': 'TR_INVENTORY', 'sku': ('SKU',), 'qty': ('QTY',),
           'container': ('Container',), 'lot': ('Lot Number',)},
    'DCW': {'raw': 'DCW', 'inv': 'DCW_INVENTORY', 'sku': ('Style',), 'qty': ('QTY',),
            'container': ('Container',), 'lot': ('Lot Number',)},
}


# ── small helpers ────────────────────────────────────────────────────────────
def clean_key(value):
    """Container/style key: every whitespace (incl. non-breaking) removed, upper case."""
    return re.sub(r'[\s\xa0]+', '', str(value if value is not None else '')).upper()


def is_container(value):
    """A real ocean container number after clean_key (never NA, N/A, NONE, DELIVERY...)."""
    return bool(CONTAINER_RE.match(clean_key(value)))


def _num(value):
    if value is None or isinstance(value, bool):
        return 0.0
    if isinstance(value, (int, float)):
        return float(value) if math.isfinite(value) else 0.0
    try:
        f = float(str(value).replace(',', '').strip())
    except (TypeError, ValueError):
        return 0.0
    return f if math.isfinite(f) else 0.0


def _int(value):
    return int(_num(value))


def _text(value):
    """Cell text, '' for empty and the usual 'nan' / 'N/A' fillers."""
    if value is None:
        return ''
    s = str(value).strip()
    return '' if s.lower() in _EMPTY_TEXT else s


_DATE_PATTERNS = (
    (re.compile(r'^(\d{4})-(\d{1,2})-(\d{1,2})(?:[ T]\d{1,2}:\d{2}(?::\d{2}(?:\.\d+)?)?)?$'), 'ymd'),
    (re.compile(r'^(\d{1,2})-(\d{1,2})-(\d{4})$'), 'mdy'),
    (re.compile(r'^(\d{1,2})/(\d{1,2})/(\d{4})(?: \d{1,2}:\d{2}(?::\d{2})?(?: ?[AP]M)?)?$', re.I), 'mdy'),
    (re.compile(r'^(\d{1,2})/(\d{1,2})/(\d{2})$'), 'mdyy'),
)
# An Excel serial day number left in a date column (the old JTW listing held a few, as text,
# in early Dec 2025). Only serials from 2010-01-01 (40179) up to 60000 (2064) are read as
# dates; any other bare number is not a receive date.
EXCEL_SERIAL_MIN, EXCEL_SERIAL_MAX = 40179, 60000
_EXCEL_EPOCH = date(1899, 12, 30)
_SERIAL_RE = re.compile(r'^\d{5}(?:\.\d+)?$')


def excel_serial_date(n):
    """date for an Excel serial day number in [EXCEL_SERIAL_MIN, EXCEL_SERIAL_MAX), else None."""
    try:
        f = float(n)
    except (TypeError, ValueError):
        return None
    if not math.isfinite(f) or not (EXCEL_SERIAL_MIN <= f < EXCEL_SERIAL_MAX):
        return None
    return _EXCEL_EPOCH + timedelta(days=int(f))


def parse_date(value):
    """date | None from a datetime/date cell or text in any of these forms:
    YYYY-MM-DD HH:MM:SS, YYYY-MM-DD, MM-DD-YYYY, M/D/YYYY, M/D/YY; or an Excel serial day
    number (a number or digits-only text) between 2010 and 2064."""
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)):
        return excel_serial_date(value)
    s = str(value).strip()
    if s.lower() in _EMPTY_TEXT:
        return None
    if _SERIAL_RE.match(s):
        return excel_serial_date(s)
    for rx, kind in _DATE_PATTERNS:
        m = rx.match(s)
        if not m:
            continue
        a, b, c = (int(x) for x in m.groups()[:3])
        if kind == 'ymd':
            y, mo, d = a, b, c
        elif kind == 'mdy':
            mo, d, y = a, b, c
        else:                # two-digit year, same pivot as strptime %y
            mo, d, y = a, b, (2000 + c if c < 69 else 1900 + c)
        try:
            return date(y, mo, d)
        except ValueError:
            return None
    return None


def _iso(d):
    return d.isoformat() if d else None


def _mdy(d):
    return d.strftime('%m-%d-%Y') if d else None


def receive_text(value):
    """The ATS 'Receive Date' as the MM-DD-YYYY text the oldest-lot method reads (None when
    blank). A datetime cell or an Excel serial number becomes MM-DD-YYYY; other text is kept."""
    if isinstance(value, bool):
        return None
    if isinstance(value, (datetime, date, int, float)):
        return _mdy(parse_date(value))
    s = _text(value)
    if s and _SERIAL_RE.match(s):
        return _mdy(parse_date(s)) or s
    return s or None


# ── style keys (mirror app.py; keep in sync) ────────────────────────────────
_SIZED_ROW_BASE_RE = re.compile(r'^[A-Z0-9]{2}[A-Z]{4}[A-Z0-9]{3}[A-Z]{2,4}$')
_SIZED_ROW_ALPHA = {'XS', 'S', 'M', 'L', 'XL', 'XXL', 'XXXL', '2XL', '3XL', '4XL', '5XL',
                    'ST', 'MT', 'LT', 'XLT', 'XXLT', '2XLT', '3XLT'}
_SIZED_ROW_TOKEN_RES = [re.compile(p) for p in (
    r'^\d{2}(\.\d)?$', r'^\d{4}\.\d{4,6}$', r'^\d{2}(\.\d)?\d{2}/\d{2}$',
    r'^\d{2}/\d{2}$', r'^\d{2}X\d{2}$', r'^\d{2}W(\d{2}L)?$')]


def is_sized_sku(sku):
    """Copy of app._is_sized_sku (a test asserts the two agree). app.py passes its own."""
    s = str(sku or '').upper().strip()
    di = s.find('-')
    if di < 0:
        return False
    if not _SIZED_ROW_BASE_RE.match(s[:di]):
        return False
    suf = s[di + 1:]
    if suf.endswith('-FBA'):
        suf = suf[:-4]
    if suf.endswith('-V'):
        suf = suf[:-2]
    if suf.startswith('V-'):
        suf = suf[2:]
    if not suf or suf in ('V', 'FBA'):
        return False

    def tok(t):
        return t in _SIZED_ROW_ALPHA or any(r.match(t) for r in _SIZED_ROW_TOKEN_RES)
    if tok(suf):
        return True
    parts = [p for p in suf.split('-') if p]
    return bool(parts) and all(tok(p) for p in parts)


def base_style(sku, is_sized=None):
    """Strip a trailing -V; then a by-size row rolls up to the part before the first '-'."""
    s = str(sku or '').strip()
    if s.upper().endswith('-V'):
        s = s[:-2]
    if (is_sized or is_sized_sku)(s):
        return s.split('-')[0]
    return s


def customer_prefix(sku):
    s = str(sku or '').strip().upper()
    return s[:2] if _PREFIX_RE.match(s) else 'Other'


def bucket_index(days, buckets=BUCKETS):
    """Index of the bucket a lot of this age (days) falls in; ages below 0 count as 0."""
    i = 0
    for j, b in enumerate(buckets):
        if days >= b[2]:
            i = j
    return i


def bucket_key(days, buckets=BUCKETS):
    return buckets[bucket_index(days, buckets)][0]


def bucket_defs(buckets=BUCKETS):
    return [{'key': k, 'label': lab, 'min': m} for k, lab, m in buckets]


# ── workbook access ──────────────────────────────────────────────────────────
def load_workbook(source):
    """openpyxl workbook (read only, values) from bytes, a path or an open workbook."""
    if hasattr(source, 'sheetnames'):
        return source
    import openpyxl
    if isinstance(source, (bytes, bytearray, memoryview)):
        return openpyxl.load_workbook(io.BytesIO(bytes(source)), read_only=True, data_only=True)
    return openpyxl.load_workbook(source, read_only=True, data_only=True)


def read_sheet(wb, name, tick=None):
    """(header index {name: first column}, data rows) or (None, None) when the sheet is missing.
    Trailing all-empty rows are dropped. tick(), when given, is called every 500 rows (the
    server passes a cooperative yield so a gevent worker keeps answering during a parse)."""
    if name not in wb.sheetnames:
        return None, None
    rows = []
    for row in wb[name].iter_rows(values_only=True):
        rows.append(row)
        if tick is not None and len(rows) % 500 == 0:
            tick()
    if not rows:
        return {}, []
    ix = {}
    for i, h in enumerate(rows[0]):
        if h is None:
            continue
        k = str(h).strip()
        if k and k not in ix:
            ix[k] = i
    data = rows[1:]
    while data and all(v is None or (isinstance(v, str) and not v.strip()) for v in data[-1]):
        data.pop()
    return ix, data


def _col(ix, names):
    for n in names:
        if n in ix:
            return ix[n]
    return None


def _cell(row, i):
    if i is None or i >= len(row):
        return None
    return row[i]


# ── ATS sheet ────────────────────────────────────────────────────────────────
_ATS_COLS = {
    'sku': ('SKU',), 'brand': ('Brand',), 'jtw': ('JTW',), 'tr': ('TR',), 'dcw': ('DCW',),
    'qa': ('QA', 'Q/A', 'Quality'), 'nj': ('NJ',), 'abfi': ('ABFI',),
    'committed': ('Committed', 'Committed Inventory'), 'allocated': ('Allocated',),
    'ats': ('Total ATS', 'Total_ATS', 'TotalATS'), 'receive': ('Receive Date', 'ReceiveDate'),
}


def _new_sku_row(sku, brand, is_sized):
    return {'sku': sku, 'base': base_style(sku, is_sized), 'brand': brand or 'OTHER',
            'prefix': customer_prefix(sku), 'jtw': 0, 'tr': 0, 'dcw': 0, 'qa': 0, 'nj': 0, 'abfi': 0,
            'committed': 0, 'allocated': 0, 'ats': 0, 'receiveDate': None}


def _fold_sku_row(rows, order, sku, brand, vals, receive, is_sized):
    """One row per SKU: stock sums; committed/allocated keep the largest magnitude (a duplicate
    row of the SAME sku repeats them, so they are never summed); the oldest receive date wins."""
    r = rows.get(sku)
    if r is None:
        r = rows[sku] = _new_sku_row(sku, brand, is_sized)
        order.append(sku)
    for k in ('jtw', 'tr', 'dcw', 'qa', 'nj', 'abfi', 'ats'):
        r[k] += vals.get(k, 0)
    for k in ('committed', 'allocated'):
        v = vals.get(k, 0)
        if abs(v) > abs(r[k]):
            r[k] = v
    if receive:
        old = parse_date(r['receiveDate'])
        new = parse_date(receive)
        if r['receiveDate'] is None or (new and (old is None or new < old)):
            r['receiveDate'] = receive


def _restricted(sku, v):
    """NJ/ABFI the way app.parse_inventory_excel keeps them: a negative on-hand never counts and
    a '-FBA' earmarked row is never shippable restricted stock."""
    if v < 0:
        return 0
    if v and str(sku or '').strip().upper().split(' ')[0].endswith('-FBA'):
        return 0
    return v


def parse_ats_sheet(wb, is_sized=None, tick=None):
    """(sku rows {sku: row}, has Receive Date column, data row count, error | None) from the
    first sheet ('ATS'), normalised like app.parse_inventory_excel (brand upper case, NT is
    NAUTICA, rows without SKU or brand skipped, NJ/ABFI negatives and FBA rows zeroed)."""
    name = 'ATS' if 'ATS' in wb.sheetnames else wb.sheetnames[0]
    ix, rows = read_sheet(wb, name, tick)
    if not ix or _col(ix, _ATS_COLS['sku']) is None:
        return {}, False, 0, 'the ATS sheet has no SKU column'
    c = {k: _col(ix, v) for k, v in _ATS_COLS.items()}
    out, order = {}, []
    for r in rows:
        sku = str(_cell(r, c['sku']) or '').strip()
        brand = str(_cell(r, c['brand']) or '').strip().upper()
        if brand == 'NT':
            brand = 'NAUTICA'
        if not sku or sku == 'N/A' or not brand:
            continue
        vals = {k: _int(_cell(r, c[k])) for k in ('jtw', 'tr', 'dcw', 'qa', 'nj', 'abfi', 'committed',
                                                     'allocated', 'ats')}
        vals['nj'] = _restricted(sku, vals['nj'])
        vals['abfi'] = _restricted(sku, vals['abfi'])
        receive = receive_text(_cell(r, c['receive']))
        _fold_sku_row(out, order, sku, brand, vals, receive, is_sized)
    return {k: out[k] for k in order}, c['receive'] is not None, len(rows), None


def sku_rows_from_items(items, is_sized=None):
    """Sku rows from the platform's parsed inventory items (app.py _inventory['items_raw']), so
    the ATS sheet is never parsed twice on the server."""
    out, order = {}, []
    for it in items or []:
        if not isinstance(it, dict):
            continue
        sku = str(it.get('sku') or '').strip()
        if not sku:
            continue
        vals = {k: _int(it.get(k)) for k in ('jtw', 'tr', 'dcw', 'qa', 'nj', 'abfi', 'committed', 'allocated')}
        vals['ats'] = _int(it.get('total_ats'))
        receive = receive_text(it.get('receive_date'))
        _fold_sku_row(out, order, sku, str(it.get('brand') or '').strip().upper(), vals, receive, is_sized)
    return {k: out[k] for k in order}


# ── lot sheets ───────────────────────────────────────────────────────────────
def _drop_trailing_keyless(rows, col):
    """rows without the trailing rows that carry no style/SKU in column col. Such rows are not
    lots: the raw JTW listing often ends in rows holding only 'Available Inv' 0 and SourceFile
    (Oct 2 to 4 2026, mid Jul 2026), which would otherwise break the row-for-row pairing.
    Only the tail is trimmed; a style-less row in the middle still makes the counts differ."""
    if col is None:
        return rows
    n = len(rows)
    while n and not clean_key(_cell(rows[n - 1], col)):
        n -= 1
    return rows[:n]


def _new_jtw_on_hand(rx, raw, ix, inv):
    """The NEW JTW listing (since 3/23/2026) keeps fully shipped lines, like the old one did;
    JTW_INVENTORY holds, in order, the lines with a positive 'OnHand' (listings that carry that
    column, 3/31 to 4/6/2026) or else 'Received', with QTY = that column. Returns (the raw lines
    to pair, column used, None) when a candidate verifies on EVERY line (style, container,
    cartons = QTY), else (None, None, reason). Never guessed."""
    rstyle, rcont = rx.get('Style'), rx.get('Container #')
    istyle, icont, icart = ix.get('Style'), ix.get('Container #'), ix.get('QTY')
    if None in (rstyle, rcont, istyle, icont, icart):
        return None, None, 'the JTW sheets have no Style, Container # or QTY column'
    errors = []
    for qname in ('OnHand', 'Received'):
        qc = rx.get(qname)
        if qc is None:
            continue
        keep = [r for r in raw if _num(_cell(r, qc)) > 0]
        if len(keep) != len(inv):
            errors.append(f'{len(keep):,} JTW lines with {qname} > 0 but {len(inv):,} JTW_INVENTORY rows')
            continue
        bad = None
        for k, (a, b) in enumerate(zip(keep, inv)):
            if (clean_key(_cell(a, rstyle)) != clean_key(_cell(b, istyle))
                    or clean_key(_text(_cell(a, rcont))) != clean_key(_text(_cell(b, icont)))
                    or _num(_cell(a, qc)) != _num(_cell(b, icart))):
                bad = f'JTW_INVENTORY row {k + 2} does not match the JTW line with {qname} > 0'
                break
        if bad:
            errors.append(bad)
            continue
        return keep, qname, None
    return None, None, '; '.join(errors) or 'the JTW sheet has no OnHand or Received column'


def _parse_wh_lots(wb, wh, tick=None):
    """(lots, error | None, format) for one warehouse. Lots are dicts with the raw date."""
    spec = _SHEETS[wh]
    rx, raw = read_sheet(wb, spec['raw'], tick)
    ix, inv = read_sheet(wb, spec['inv'], tick)
    if rx is None or ix is None:
        missing = spec['raw'] if rx is None else spec['inv']
        return [], f'the {missing} sheet is missing', None
    fmt = None
    if wh == 'JTW':
        if 'Inventory Date' in rx:
            fmt, dcol = 'new', rx['Inventory Date']
        elif 'Date Received' in rx:
            fmt, dcol = 'old', rx['Date Received']
        else:
            return [], 'the JTW sheet has no Inventory Date or Date Received column', None
    else:
        dcol = rx.get('Receipt Date')
        if dcol is None:
            return [], f'the {spec["raw"]} sheet has no Receipt Date column', None
    skucol = _col(ix, spec['sku'])
    qcol = _col(ix, spec['qty'])
    if skucol is None or qcol is None:
        return [], (f'the {spec["inv"]} sheet has no {spec["sku"][0]} or {spec["qty"][0]} column'), fmt
    ccol = _col(ix, spec['container'])
    lcol = _col(ix, spec['lot'])
    raw_style = rx.get('Style') if wh == 'JTW' else None
    if wh == 'JTW' and raw_style is None:
        return [], 'the JTW sheet has no Style column', fmt
    raw = _drop_trailing_keyless(raw, raw_style if wh == 'JTW' else rx.get('SKU'))
    inv = _drop_trailing_keyless(inv, skucol)
    if wh == 'JTW' and fmt == 'old' and len(raw) != len(inv):
        # The old JTW listing kept shipped-out receipts; JTW_INVENTORY holds the lines with a
        # positive balance, in order. Verified below line by line, never assumed.
        bcol = rx.get('Ballance', rx.get('Balance'))
        if bcol is not None:
            raw = [r for r in raw if _num(_cell(r, bcol)) > 0]
            if len(raw) == len(inv):
                rcont = rx.get('Container')
                icart = ix.get('QTY')
                for k, (a, b) in enumerate(zip(raw, inv)):
                    if (clean_key(_cell(a, raw_style)) != clean_key(_cell(b, skucol))
                            or (rcont is not None and ccol is not None
                                and clean_key(_cell(a, rcont)) != clean_key(_cell(b, ccol)))
                            or (icart is not None and _num(_cell(a, bcol)) != _num(_cell(b, icart)))):
                        return [], (f'JTW row {k + 2} of JTW_INVENTORY does not match the JTW listing '
                                    'line with a balance, so the two sheets cannot be paired'), fmt
    if wh == 'JTW' and fmt == 'new' and len(raw) != len(inv):
        # Same pattern in the new layout (fully shipped lines stay listed). When no candidate
        # verifies, the counts still differ and the warehouse is reported below, never guessed.
        kept = _new_jtw_on_hand(rx, raw, ix, inv)[0]
        if kept is not None:
            raw = kept
    if len(raw) != len(inv):
        return [], (f'{spec["raw"]} has {len(raw):,} rows but {spec["inv"]} has {len(inv):,}, '
                    'so the receive dates cannot be paired with the lots'), fmt
    lots = []
    for k, (a, b) in enumerate(zip(raw, inv)):
        sku = str(_cell(b, skucol) or '').strip()
        units = _int(_cell(b, qcol))
        if units <= 0 or not sku:
            continue
        rstyle = clean_key(_cell(a, raw_style)) if raw_style is not None else None
        if wh == 'JTW' and rstyle != clean_key(sku):
            return [], (f'JTW_INVENTORY row {k + 2} has style {sku} but the JTW row next to it has '
                        f'{str(_cell(a, raw_style) or "").strip() or "no style"}, so the two sheets '
                        'cannot be paired'), fmt
        rawd = _cell(a, dcol)
        lots.append({'sku': sku, 'wh': wh, 'units': units, 'date': parse_date(rawd),
                     'rawDate': None if rawd is None else str(rawd).strip(),
                     'container': str(_cell(b, ccol) or '').strip() if ccol is not None else '',
                     'lot': _text(_cell(b, lcol)) if lcol is not None else '',
                     'rawStyle': rstyle, 'flags': [], 'origDate': None})
    return lots, None, fmt


def parse_workbook(source, sku_rows=None, is_sized=None, tick=None):
    """Parse the lot sheets (and the ATS sheet unless sku_rows is given).

    source:   workbook bytes, a path, or an openpyxl workbook.
    sku_rows: {sku: row} from sku_rows_from_items (server); None parses the ATS sheet.
    is_sized: the by-size test for base styles (default: this module's copy of app._is_sized_sku).
    tick:     optional callable, called every 500 rows read (cooperative yield).
    Returns {
      'skus': {sku: row}, 'atsRows': n, 'atsHasReceiveDate': bool, 'atsError': str | None,
      'lots': [lot], 'whErrors': {wh: str}, 'lotsOk': bool, 'lotsError': str | None,
      'jtwFormat': 'new' | 'old' | None, 'lotTotals': {wh: units}, 'whTotals': {wh: units},
      'notInAts': units, 'notInAtsSkus': [sku]}
    Lots whose SKU is not an ATS SKU are dropped and counted in notInAts. Lots of a warehouse
    whose sheets do not pair are left out and named in whErrors."""
    wb = load_workbook(source)
    try:
        if sku_rows is None:
            skus, has_rd, n_rows, ats_err = parse_ats_sheet(wb, is_sized, tick)
        else:
            skus = sku_rows
            has_rd = any(r.get('receiveDate') for r in skus.values())
            n_rows, ats_err = len(skus), None
        lots, wh_err, fmt = [], {}, None
        for wh in WAREHOUSES:
            got, err, f = _parse_wh_lots(wb, wh, tick)
            if wh == 'JTW':
                fmt = f
            if err:
                wh_err[wh] = err
                continue
            lots.extend(got)
    finally:
        if not hasattr(source, 'sheetnames'):
            try:
                wb.close()
            except Exception:
                pass
    upper = {}
    for k in skus:
        upper.setdefault(k.upper(), k)
    kept, not_in, not_in_skus = [], 0, set()
    for lot in lots:
        sku = lot['sku'] if lot['sku'] in skus else upper.get(lot['sku'].upper())
        if sku is None:
            not_in += lot['units']
            not_in_skus.add(lot['sku'])
            continue
        lot['sku'] = sku
        kept.append(lot)
    lot_totals = {wh: 0 for wh in WAREHOUSES}
    for lot in kept:
        lot_totals[lot['wh']] += lot['units']
    wh_totals = {wh: sum(max(0, r[wh.lower()]) for r in skus.values()) for wh in WAREHOUSES}
    lots_error = None
    if wh_err:
        lots_error = '. '.join(f'{wh}: {msg}' for wh, msg in wh_err.items()) + '.'
    elif ats_err:
        lots_error = ats_err
    return {'skus': skus, 'atsRows': n_rows, 'atsHasReceiveDate': bool(has_rd), 'atsError': ats_err,
            'lots': kept, 'whErrors': wh_err, 'lotsOk': not wh_err and not ats_err, 'lotsError': lots_error,
            'jtwFormat': fmt, 'lotTotals': lot_totals, 'whTotals': wh_totals,
            'notInAts': not_in, 'notInAtsSkus': sorted(not_in_skus)}


# ── restore map (JTW listing before the 3/18/2026 count) ────────────────────
def restore_map_from_jtw_sheet(source, source_label='', built_at=None):
    """{'pairs': {'CONTAINER|STYLE': iso}, 'containers': {'CONTAINER': iso}, 'source', 'builtAt'}:
    the oldest 'Date Received' per (container, style) and per container from an OLD-format
    'JTW' sheet ('Date Received', 'Container', 'Style'). source: bytes, path or workbook.
    Only real container numbers and valid dates (2010 or later) are kept."""
    wb = load_workbook(source)
    try:
        ix, rows = read_sheet(wb, 'JTW')
    finally:
        if not hasattr(source, 'sheetnames'):
            try:
                wb.close()
            except Exception:
                pass
    if ix is None:
        raise ValueError('the workbook has no JTW sheet')
    for col in ('Date Received', 'Container', 'Style'):
        if col not in ix:
            raise ValueError(f'the JTW sheet has no {col} column (not the old JTW format)')
    pairs, conts = {}, {}
    for r in rows:
        d = parse_date(_cell(r, ix['Date Received']))
        c = clean_key(_cell(r, ix['Container']))
        st = clean_key(_cell(r, ix['Style']))
        if not d or d < MIN_VALID_DATE or not CONTAINER_RE.match(c):
            continue
        if st:
            k = c + '|' + st
            if k not in pairs or d < pairs[k]:
                pairs[k] = d
        if c not in conts or d < conts[c]:
            conts[c] = d
    return {'pairs': {k: v.isoformat() for k, v in sorted(pairs.items())},
            'containers': {k: v.isoformat() for k, v in sorted(conts.items())},
            'source': str(source_label or ''), 'builtAt': built_at or datetime.utcnow().strftime('%Y-%m-%dT%H:%M:%SZ')}


# ── corrections ──────────────────────────────────────────────────────────────
def _valid(d, today):
    return d is not None and MIN_VALID_DATE <= d <= today + timedelta(days=1)


def copy_lots(lots):
    return [dict(lot, flags=list(lot.get('flags') or [])) for lot in lots]


def apply_corrections(lots, restore_map, today, jtw_format='new'):
    """Corrected copies of the lots, in this order:
      1. bad dates (before 2010, more than a day ahead, or unreadable): the oldest valid date of
         another lot with the same SKU and container ('bad_date_fixed'), else undated ('bad_date');
      2. JTW lots on the 3/18 or 3/20/2026 count date (new JTW format only): the restore map by
         (container, style), then by container ('restored' when older), else 'count_date';
      3. 'reset' when the container or lot text says transfer, truck, recount, RMA...
    Returns (lots, {'available', 'source', 'matchedUnits', 'unmatchedUnits'})."""
    out = copy_lots(lots)
    # 1. bad dates
    good = {}
    for lot in out:
        if _valid(lot['date'], today):
            k = (lot['sku'], clean_key(lot['container']))
            if k not in good or lot['date'] < good[k]:
                good[k] = lot['date']
    for lot in out:
        if _valid(lot['date'], today):
            continue
        alt = good.get((lot['sku'], clean_key(lot['container'])))
        lot['origDate'] = lot['date']
        lot['date'] = alt
        lot['flags'].append('bad_date_fixed' if alt else 'bad_date')
    # 2. JTW count re-date
    rm = restore_map if isinstance(restore_map, dict) else None
    pairs = (rm or {}).get('pairs') or {}
    conts = (rm or {}).get('containers') or {}
    matched = unmatched = 0
    if jtw_format != 'old':
        for lot in out:
            if lot['wh'] != 'JTW' or lot['date'] not in COUNT_DATES:
                continue
            rd = None
            c = clean_key(lot['container'])
            if rm and CONTAINER_RE.match(c):
                st = lot.get('rawStyle') or clean_key(lot['sku'])
                rd = parse_date(pairs.get(c + '|' + st)) or parse_date(conts.get(c))
            if rd and rd < lot['date'] and _valid(rd, today):
                if lot['origDate'] is None:
                    lot['origDate'] = lot['date']
                lot['date'] = rd
                lot['flags'].append('restored')
                matched += lot['units']
            else:
                lot['flags'].append('count_date')
                unmatched += lot['units']
    # 3. transfer / recount dates
    for lot in out:
        if RESET_RE.search((str(lot['container']) + ' ' + str(lot['lot'])).upper()):
            lot['flags'].append('reset')
    return out, {'available': rm is not None, 'source': (rm or {}).get('source') if rm else None,
                 'matchedUnits': matched, 'unmatchedUnits': unmatched}


def lot_mismatches(lots, skus, warehouses=WAREHOUSES):
    """[{'sku','wh','lots','ats'}] where a SKU's lot units in a warehouse are not its ATS column."""
    got = {}
    for lot in lots:
        k = (lot['sku'], lot['wh'])
        got[k] = got.get(k, 0) + lot['units']
    out = []
    for sku, r in skus.items():
        for wh in warehouses:
            ats = r[wh.lower()]
            have = got.get((sku, wh), 0)
            if have != ats:
                out.append({'sku': sku, 'wh': wh, 'lots': have, 'ats': ats})
    return out


def analyze(parsed, restore_map, today):
    """parsed + corrections for one day: {'parsed', 'lots', 'restore', 'today', 'mismatches'}."""
    lots, restore = apply_corrections(parsed['lots'], restore_map, today, parsed.get('jtwFormat') or 'new')
    ok_wh = tuple(wh for wh in WAREHOUSES if wh not in (parsed.get('whErrors') or {}))
    return {'parsed': parsed, 'lots': lots, 'restore': restore, 'today': today,
            'mismatches': lot_mismatches(lots, parsed['skus'], ok_wh) if not parsed.get('atsError') else []}


# ── aggregates ───────────────────────────────────────────────────────────────
def _age(d, today):
    return max(0, (today - d).days)


def lot_aggregate(lots, skus, today):
    """By-lot units after corrections, dated lots only: totals, by brand / warehouse / prefix,
    the unit-weighted average age, and undated lot units."""
    tot = [0] * len(BUCKETS)
    by_brand, by_wh, by_prefix = {}, {wh: [0] * len(BUCKETS) for wh in WAREHOUSES}, {}
    wsum = usum = undated = 0
    for lot in lots:
        if lot['date'] is None:
            undated += lot['units']
            continue
        days = _age(lot['date'], today)
        i = bucket_index(days)
        u = lot['units']
        tot[i] += u
        wsum += days * u
        usum += u
        row = skus.get(lot['sku']) or {}
        by_brand.setdefault(row.get('brand') or 'OTHER', [0] * len(BUCKETS))[i] += u
        by_wh.setdefault(lot['wh'], [0] * len(BUCKETS))[i] += u
        by_prefix.setdefault(row.get('prefix') or customer_prefix(lot['sku']), [0] * len(BUCKETS))[i] += u
    return {'units': tot, 'brand': by_brand, 'wh': by_wh, 'prefix': by_prefix, 'datedUnits': usum,
            'wavg': (wsum / usum) if usum else None, 'undatedLotUnits': undated}


def nuri_buckets(skus, today, has_receive_date=True):
    """Nuri's oldest-lot method: per SKU, the ATS Receive Date and JTW+TR+DCW units, 360-day
    years. SKUs without a Receive Date are left out. None when the file has no Receive Date."""
    if not has_receive_date:
        return None
    out = [0] * len(NURI_BUCKETS)
    seen = False
    for r in skus.values():
        d = parse_date(r.get('receiveDate'))
        if d is None:
            continue
        seen = True
        out[bucket_index(_age(d, today), NURI_BUCKETS)] += max(0, r['jtw']) + max(0, r['tr']) + max(0, r['dcw'])
    return out if seen else None


def _top_prefixes(by_prefix, top=TREND_TOP_PREFIXES):
    ranked = sorted((p for p in by_prefix if p != 'Other'), key=lambda p: (-sum(by_prefix[p]), p))
    keep = ranked[:top]
    out = {p: list(by_prefix[p]) for p in keep}
    other = [0] * len(BUCKETS)
    for p, v in by_prefix.items():
        if p not in out:
            other = [a + b for a, b in zip(other, v)]
    if sum(other):
        out['Other'] = other
    return out


def trend_record(analysis, src='live'):
    """One trend record for analysis['today'] (staff safe: dated JTW/TR/DCW lots only, no NJ or
    ABFI). None when the lot detail is not usable."""
    parsed = analysis['parsed']
    if not parsed.get('lotsOk'):
        return None
    today = analysis['today']
    agg = lot_aggregate(analysis['lots'], parsed['skus'], today)
    return {'d': today.isoformat(), 'units': agg['units'],
            'wavg': round(agg['wavg'], 1) if agg['wavg'] is not None else None,
            'brand': {k: v for k, v in sorted(agg['brand'].items())},
            'wh': {wh: agg['wh'].get(wh, [0] * len(BUCKETS)) for wh in WAREHOUSES},
            'prefix': _top_prefixes(agg['prefix']),
            'nuri': nuri_buckets(parsed['skus'], today, parsed.get('atsHasReceiveDate', True)),
            'restored': sum(l['units'] for l in analysis['lots'] if 'restored' in l['flags']),
            'countDate': sum(l['units'] for l in analysis['lots'] if 'count_date' in l['flags']),
            'src': src}


# ── payload ──────────────────────────────────────────────────────────────────
def _fmt(n):
    return f'{int(n):,}'


def _check(code, severity, title, detail, units, skus, **extra):
    c = {'code': code, 'severity': severity, 'title': title, 'detail': detail, 'units': int(units),
         'skus': sorted(set(skus))}
    c.update(extra)
    return c


def build_checks(analysis, is_admin):
    parsed, lots, restore = analysis['parsed'], analysis['lots'], analysis['restore']
    skus = parsed['skus']
    checks = []
    if not parsed.get('lotsOk'):
        bad = sorted(parsed.get('whErrors') or {})
        what = ', '.join(bad) if bad else 'the ATS sheet'
        detail = f'The lot sheets could not be read for {what}.'
        if parsed.get('lotsError'):
            detail += ' ' + parsed['lotsError']
        detail += " Only the oldest lot method (Nuri's report) is complete."
        checks.append(_check('lots_unavailable', 'high', 'Lot detail unavailable', detail, 0, [],
                             warehouses=bad))
    sel = [l for l in lots if 'bad_date' in l['flags'] or 'bad_date_fixed' in l['flags']]
    if sel:
        fixed = sum(l['units'] for l in sel if 'bad_date_fixed' in l['flags'])
        checks.append(_check('bad_date', 'high', 'Impossible receive dates',
                             'Lots dated before 2010, more than a day in the future, or with no readable date. '
                             'Fixed from another lot of the same style in the same container when possible '
                             f'({_fmt(fixed)} units fixed), otherwise left undated.',
                             sum(l['units'] for l in sel), [l['sku'] for l in sel]))
    sel = [l for l in lots if 'count_date' in l['flags']]
    if sel:
        detail = ('These JTW lots could not be matched to the JTW listing from before the 3/18/2026 physical '
                  'count, so the date shown is the count date and the stock is probably older.'
                  if restore.get('available') else
                  'The JTW listing from before the 3/18/2026 physical count is not loaded, so every JTW lot '
                  'dated by the count keeps the count date. The stock is probably older.')
        checks.append(_check('count_date', 'medium', 'JTW stock still dated by the 3/18 count', detail,
                             sum(l['units'] for l in sel), [l['sku'] for l in sel]))
    sel = [l for l in lots if 'restored' in l['flags']]
    if sel:
        checks.append(_check('restored', 'info', 'JTW dates restored',
                             'Dates taken from the JTW listing before the 3/18/2026 physical count.',
                             sum(l['units'] for l in sel), [l['sku'] for l in sel]))
    sel = [l for l in lots if 'reset' in l['flags']]
    if sel:
        by_wh = {wh: sum(l['units'] for l in sel if l['wh'] == wh) for wh in WAREHOUSES}
        checks.append(_check('reset', 'info', 'Dated by a transfer, recount or return',
                             'The date is when the stock was moved or recounted, not when it first arrived, so '
                             'the stock is probably older. ' +
                             ', '.join(f'{wh} {_fmt(by_wh[wh])}' for wh in WAREHOUSES if by_wh[wh]) + ' units.',
                             sum(by_wh.values()), [l['sku'] for l in sel], byWh=by_wh))
    parts = {'qa': sum(max(0, r['qa']) for r in skus.values())}
    und_skus = [s for s, r in skus.items() if r['qa'] > 0]
    if is_admin:
        parts['nj'] = sum(max(0, r['nj']) for r in skus.values())
        parts['abfi'] = sum(max(0, r['abfi']) for r in skus.values())
        und_skus += [s for s, r in skus.items() if r['nj'] > 0 or r['abfi'] > 0]
    lot_und = [l for l in lots if l['date'] is None]
    parts['lots'] = sum(l['units'] for l in lot_und)
    und_skus += [l['sku'] for l in lot_und]
    if sum(parts.values()):
        names = 'QA, NJ and ABFI stock carries' if is_admin else 'QA stock carries'
        bits = [f'QA {_fmt(parts["qa"])}']
        if is_admin:
            bits += [f'NJ {_fmt(parts["nj"])}', f'ABFI {_fmt(parts["abfi"])}']
        if parts['lots']:
            bits.append(f'lots with no usable date {_fmt(parts["lots"])}')
        checks.append(_check('undated', 'info', 'Stock with no receive date',
                             f'{names} no receive date, so it is not in the age buckets. ' + ', '.join(bits) + ' units.',
                             sum(parts.values()), und_skus, parts=parts))
    mm = analysis.get('mismatches') or []
    if mm:
        off = sum(abs(m['lots'] - m['ats']) for m in mm)
        checks.append(_check('lot_mismatch', 'high', 'Lot detail does not match the ATS sheet',
                             f'For {_fmt(len(mm))} style and warehouse pairs the lots do not add up to the ATS sheet '
                             f'column for that warehouse ({_fmt(off)} units apart in total). Ages for those styles '
                             'follow the lots.', off, [m['sku'] for m in mm], pairs=mm[:500]))
    if parsed.get('notInAts'):
        checks.append(_check('not_in_ats', 'info', 'Lots for items that are not on the ATS sheet',
                             'These lot lines name an item the ATS sheet does not list, so they are left out.',
                             parsed['notInAts'], parsed.get('notInAtsSkus') or []))
    return checks


def build_payload(analysis, is_admin, wh_ats=None, rev=None, as_of=None, built_at=None):
    """The GET /api/inventory-aging body. wh_ats: {base: units available now} (None: unknown).
    A non-admin payload carries no NJ or ABFI number anywhere."""
    parsed, lots, restore, today = analysis['parsed'], analysis['lots'], analysis['restore'], analysis['today']
    skus = parsed['skus']
    sku_rows, bases = [], {}
    for s, r in skus.items():
        sku_rows.append([s, r['base'], r['brand'], r['prefix'], r['jtw'], r['tr'], r['dcw'], r['qa'],
                         r['nj'] if is_admin else 0, r['abfi'] if is_admin else 0,
                         r['committed'], r['allocated'], r['receiveDate']])
        if r['base'] not in bases:
            bases[r['base']] = r['brand']
    wa = None
    if isinstance(wh_ats, dict):
        wa = {str(k).upper(): v for k, v in wh_ats.items()}
    base_rows = [[b, br, (int(wa.get(b.upper(), 0)) if wa is not None else None)] for b, br in bases.items()]
    excluded = {'qa': sum(max(0, r['qa']) for r in skus.values())}
    if is_admin:
        excluded['nj'] = sum(max(0, r['nj']) for r in skus.values())
        excluded['abfi'] = sum(max(0, r['abfi']) for r in skus.values())
    excluded['notInAts'] = parsed.get('notInAts') or 0
    return {
        'ready': True, 'rev': rev, 'asOf': as_of,
        'builtAt': built_at or datetime.utcnow().strftime('%Y-%m-%dT%H:%M:%SZ'),
        'today': today.isoformat(), 'isAdmin': bool(is_admin),
        'lotsOk': bool(parsed.get('lotsOk')), 'lotsError': parsed.get('lotsError'),
        'buckets': bucket_defs(BUCKETS), 'nuriBuckets': bucket_defs(NURI_BUCKETS),
        'lotCols': list(LOT_COLS),
        'lots': [[l['sku'], l['wh'], l['units'], _iso(l['date']), l['container'], l['lot'], list(l['flags']),
                  _iso(l['origDate'])] for l in lots],
        'skuCols': list(SKU_COLS), 'skus': sku_rows,
        'baseCols': list(BASE_COLS), 'bases': base_rows,
        'checks': build_checks(analysis, is_admin),
        'restore': {'available': bool(restore.get('available')), 'source': restore.get('source'),
                    'matchedUnits': restore.get('matchedUnits', 0), 'unmatchedUnits': restore.get('unmatchedUnits', 0)},
        'excluded': excluded,
    }


# ── trend store documents ────────────────────────────────────────────────────
def _is_int(v):
    return isinstance(v, int) and not isinstance(v, bool)


def _is_num(v):
    return (isinstance(v, (int, float)) and not isinstance(v, bool) and math.isfinite(v))


def _iso_ok(v):
    if not isinstance(v, str) or not _ISO_RE.match(v):
        return False
    try:
        date.fromisoformat(v)
        return True
    except ValueError:
        return False


def _vec_ok(v):
    return isinstance(v, list) and len(v) == len(BUCKETS) and all(_is_int(x) and x >= 0 for x in v)


def _map_ok(m, max_keys=500):
    return (isinstance(m, dict) and len(m) <= max_keys
            and all(isinstance(k, str) and 0 < len(k) <= 40 and _vec_ok(v) for k, v in m.items()))


def validate_trend_record(rec):
    """None when rec is a well-formed trend record, else a short reason."""
    if not isinstance(rec, dict):
        return 'record is not an object'
    if set(rec) != set(TREND_KEYS):
        return 'record keys must be exactly ' + ', '.join(TREND_KEYS)
    if not _iso_ok(rec['d']):
        return 'd must be a YYYY-MM-DD date'
    if not _vec_ok(rec['units']):
        return f'units must be {len(BUCKETS)} non-negative integers'
    if rec['wavg'] is not None and not (_is_num(rec['wavg']) and rec['wavg'] >= 0):
        return 'wavg must be a non-negative number or null'
    for k in ('brand', 'wh', 'prefix'):
        if not _map_ok(rec[k]):
            return f'{k} must map names to {len(BUCKETS)} non-negative integers'
    if rec['nuri'] is not None and not _vec_ok(rec['nuri']):
        return f'nuri must be null or {len(NURI_BUCKETS)} non-negative integers'
    for k in ('restored', 'countDate'):
        if not (_is_int(rec[k]) and rec[k] >= 0):
            return f'{k} must be a non-negative integer'
    if rec['src'] not in ('backfill', 'live'):
        return 'src must be backfill or live'
    return None


def _validate_dropped(rows):
    if not isinstance(rows, list) or len(rows) > TREND_MAX_RECORDS:
        return 'dropped must be a list'
    for r in rows:
        if not isinstance(r, dict) or set(r) != {'d', 'reason'}:
            return 'each dropped entry must be {"d", "reason"}'
        if not _iso_ok(r['d']) or not isinstance(r['reason'], str) or len(r['reason']) > 500:
            return 'dropped entries need a YYYY-MM-DD d and a reason of at most 500 characters'
    return None


def _validate_restore(m):
    if not isinstance(m, dict):
        return 'restore must be an object'
    extra = set(m) - {'pairs', 'containers', 'source', 'builtAt'}
    if extra or not isinstance(m.get('pairs'), dict) or not isinstance(m.get('containers'), dict):
        return 'restore must be {"pairs", "containers", "source", "builtAt"}'
    for k in ('source', 'builtAt'):
        if k in m and m[k] is not None and (not isinstance(m[k], str) or len(m[k]) > 300):
            return f'restore.{k} must be a string of at most 300 characters'
    if len(m['pairs']) > 200000 or len(m['containers']) > 50000:
        return 'restore map is too large'
    for k, v in m['pairs'].items():
        c, sep, st = str(k).partition('|')
        if not sep or not CONTAINER_RE.match(c) or not st or len(st) > 80 or re.search(r'\s', st) or st != st.upper():
            return 'restore.pairs keys must be CONTAINER|STYLE (cleaned, upper case)'
        if not _iso_ok(v):
            return 'restore dates must be YYYY-MM-DD'
    for k, v in m['containers'].items():
        if not isinstance(k, str) or not CONTAINER_RE.match(k):
            return 'restore.containers keys must be container numbers like ABCU1234567'
        if not _iso_ok(v):
            return 'restore dates must be YYYY-MM-DD'
    return None


def validate_seed(body):
    """POST /admin/aging/seed body -> (kind, value, error).
    {"restore": {...}}                                        -> ('restore', map, None)
    {"trend": {"series": [...], "dropped": [...]}, "mode": "replace" | "merge"}
                                                              -> ('trend', {'series','dropped','mode'}, None)
    Anything else -> (None, None, reason)."""
    if not isinstance(body, dict):
        return None, None, 'body must be a JSON object'
    if 'restore' in body:
        if set(body) != {'restore'}:
            return None, None, 'a restore seed carries only "restore"'
        err = _validate_restore(body['restore'])
        if err:
            return None, None, err
        r = body['restore']
        return 'restore', {'pairs': dict(r['pairs']), 'containers': dict(r['containers']),
                           'source': r.get('source') or '', 'builtAt': r.get('builtAt') or ''}, None
    if 'trend' in body:
        if not set(body) <= {'trend', 'mode'}:
            return None, None, 'a trend seed carries only "trend" and "mode"'
        mode = body.get('mode', 'replace')
        if mode not in ('replace', 'merge'):
            return None, None, 'mode must be replace or merge'
        t = body['trend']
        if not isinstance(t, dict) or not set(t) <= {'series', 'dropped', 'updatedAt'} or 'series' not in t:
            return None, None, 'trend must be {"series": [...], "dropped": [...]}'
        series = t['series']
        if not isinstance(series, list) or len(series) > TREND_MAX_RECORDS:
            return None, None, f'series must be a list of at most {TREND_MAX_RECORDS} records'
        seen = set()
        for i, rec in enumerate(series):
            err = validate_trend_record(rec)
            if err:
                return None, None, f'series[{i}]: {err}'
            if rec['d'] in seen:
                return None, None, f'series[{i}]: duplicate date {rec["d"]}'
            seen.add(rec['d'])
        dropped = t.get('dropped') or []
        err = _validate_dropped(dropped)
        if err:
            return None, None, err
        return 'trend', {'series': series, 'dropped': dropped, 'mode': mode}, None
    return None, None, 'body must carry "restore" or "trend"'


def empty_trend():
    return {'series': [], 'dropped': [], 'updatedAt': None}


def _norm_doc(doc):
    if not isinstance(doc, dict):
        return empty_trend()
    return {'series': [r for r in (doc.get('series') or []) if isinstance(r, dict) and r.get('d')],
            'dropped': [r for r in (doc.get('dropped') or []) if isinstance(r, dict) and r.get('d')],
            'updatedAt': doc.get('updatedAt')}


def merge_trend(doc, series, dropped, mode, updated_at):
    """replace: the stored document becomes exactly series + dropped. merge: incoming records
    replace the stored ones of the same date, the rest stay. A date with a record is never also
    listed as dropped."""
    if mode == 'replace':
        by_d = {r['d']: r for r in series}
        drop = {r['d']: r for r in dropped}
    else:
        cur = _norm_doc(doc)
        by_d = {r['d']: r for r in cur['series']}
        by_d.update({r['d']: r for r in series})
        drop = {r['d']: r for r in cur['dropped']}
        drop.update({r['d']: r for r in dropped})
    return {'series': [by_d[d] for d in sorted(by_d)],
            'dropped': [drop[d] for d in sorted(drop) if d not in by_d],
            'updatedAt': updated_at}


def upsert_trend(doc, record, updated_at):
    """(new document, changed) with record stored for its date. Never touches another date's
    record (backfill or live); a dropped entry for the same date is removed."""
    cur = _norm_doc(doc)
    by_d = {r['d']: r for r in cur['series']}
    if by_d.get(record['d']) == record:
        return cur, False
    by_d[record['d']] = record
    return {'series': [by_d[d] for d in sorted(by_d)],
            'dropped': [r for r in cur['dropped'] if r['d'] != record['d']],
            'updatedAt': updated_at}, True
