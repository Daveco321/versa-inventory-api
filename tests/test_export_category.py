"""Category column (David, Oct 28 2026): every line-sheet export says what each style is
(Dress Shirt, Dress Pants, Shorts, Sportswear, Blazer, Vest, Accessories), right after
Brand, on admin and customer exports alike: single and multi-brand exports, the TJX
colour-key layout, the AI / mailbox breakdown sheets, caller-defined column sets (ship
plans, Customer Selling, the weekly APO report) and the Overseas Summary.
Synthetic rows only, offline. Run from tests/: python -B -m unittest test_export_category
"""
import io
import unittest

import openpyxl

from _support import load_app
from test_tjx_workbook import Base, row


class Labels(unittest.TestCase):

    @classmethod
    def setUpClass(cls):
        cls.app = load_app()

    def lab(self, sku, brand='', fabrication=''):
        return self.app._export_category(sku, brand, fabrication)

    def test_dress_shirts(self):
        for sku, brand in (('ROEBFT007SLS', 'EB'),          # long sleeve
                           ('AMDKPK001SSS', 'DKNY'),        # short sleeve
                           ('HANASU201BBB', 'NAUTICA'),     # Big & Tall
                           ('BUCHOX001TBS', 'CHAPS'),       # short sleeve Big & Tall
                           ('AMNASU820SLP-V', 'NAUTICA'),   # partial-carton variant
                           ('AMNASU100SLP-L-FBA', 'NAUTICA'),   # by-size row
                           ('KHGBWS002SLS 17', 'BEENE')):   # size after a space
            self.assertEqual(self.lab(sku, brand), 'Dress Shirt', sku)

    def test_pants_and_shorts(self):
        for sku in ('AMUSDPP01SRS', 'BIGBRPP01SHS', 'HPNATBP01SRP', 'BUEBRPP01SRS', 'WMUSAWP06SRS',
                    'ZZUSDPP01TRS',       # Big & Tall pants fit
                    'AMUSDPP01SRS-32X30'):
            self.assertEqual(self.lab(sku), 'Dress Pants', sku)
        self.assertEqual(self.lab('ZZUSDPP01BAS'), 'Shorts')   # a Bermuda fit in a dress fabric

    def test_young_men_bottoms_are_sportswear(self):
        # BA/BC/BR/BH are the Young Men bottoms fabrics: Von Dutch pinstripe pants and Bermudas
        for sku in ('TMVDBAP07SCP', 'BUVDBA010SCP', 'ROVDBA008SCP', 'BUVDBAP14THS',
                    'BUVDBAP17BHS', 'BUVDBAP18BKS', 'ZZGBBCP01SRS'):
            self.assertEqual(self.lab(sku), 'Sportswear', sku)

    def test_von_dutch_is_sportswear(self):
        # graphic wovens in S-XL, Big & Tall included; the brand key decides, so a
        # DV typo in the style number still reads Von Dutch
        for sku in ('TTVDTS075PSP', 'TTVDAW003PSD', 'CIVDTS128SSS', 'ROVDTS120WBK',
                    'NOVDTS056PSP M', 'PMDVBA008SCP'):
            self.assertEqual(self.lab(sku, 'VD'), 'Sportswear', sku)
        self.assertEqual(self.lab('TTVDTS075PSP'), 'Sportswear')   # brand read from the style number
        self.assertEqual(self.lab('SSVDNP001C', 'VD'), 'Accessories')   # jewelry still wins

    def test_sportswear(self):
        for sku in ('BRGBPO001SSM', 'BRGBPL002SLM', 'BRGBPW027SSZ', 'ROGBRB001SLM',   # polos, rugby
                    'AMVDSL007SLC', 'AMVDCO027SLS', 'BUVDSH042TBE'):                  # shackets, overshirts
            self.assertEqual(self.lab(sku), 'Sportswear', sku)

    def test_blazers_and_vests(self):
        self.assertEqual(self.lab('BUGBRPB05BRV'), 'Blazer')
        self.assertEqual(self.lab('HPNAKSB01DBV'), 'Blazer')
        self.assertEqual(self.lab('BUGBRPV01SLS'), 'Vest')

    def test_accessories(self):
        self.assertEqual(self.lab('CTH01', 'CHAPS'), 'Accessories')        # Chaps tie & hanky
        self.assertEqual(self.lab('CTH08-V', 'CHAPS'), 'Accessories')
        self.assertEqual(self.lab('JCTSHPP031', 'SHAQ'), 'Accessories')    # Shaq tie
        for sku in ('SSVDNP001C', 'BVDNP007R', 'BVDR002', 'SSVDR001-SZ10', 'STVDR003'):
            self.assertEqual(self.lab(sku, 'VD'), 'Accessories', sku)       # Von Dutch jewelry
        self.assertEqual(self.lab('ZZVDXX001', 'VD', 'Jewelry'), 'Accessories')   # override fabrication
        self.assertEqual(self.lab('CTH01', '', '100% Polyester - Tie & Hanky'), 'Accessories')

    def test_shaq_shirts_with_a_t_customer_code_are_not_ties(self):
        # the T of the JT customer code is not the tie T (JCTSH...)
        for sku in ('JTSHST049BTP', 'JTSHST002BBS', 'JTSHST005TTP'):
            self.assertEqual(self.lab(sku, 'SHAQ'), 'Dress Shirt', sku)

    def test_old_keys_are_dress_shirts(self):
        for sku, brand in (('WACHA18', 'CHAPS'), ('1661337', 'NAUTICA'), ('NAU-008-BLACK XL', 'NAUTICA'),
                           ('BEN-22-116', 'BEN'), ('T-WHT-SL', 'TAYION'), ('CU-NAUSP272', 'NAUTICA'),
                           ('BJNASSP27-4', 'NAUTICA')):     # "SS" here is not a shacket fabric
            self.assertEqual(self.lab(sku, brand), 'Dress Shirt', sku)

    def test_blank_sku(self):
        self.assertEqual(self.lab(''), '')
        self.assertEqual(self.lab(None), '')


class Headers(unittest.TestCase):

    @classmethod
    def setUpClass(cls):
        cls.app = load_app()

    def headers(self, **kw):
        import xlsxwriter
        wb = xlsxwriter.Workbook(io.BytesIO(), {'in_memory': True})
        ws = wb.add_worksheet('T')
        _, h = self.app._setup_worksheet(wb, ws, has_color=True, **kw)
        wb.close()
        return h

    def assertAfterBrand(self, h):
        self.assertIn('Category', h)
        self.assertEqual(h.index('Category'), h.index('Brand') + 1, h)
        self.assertEqual(h.count('Category'), 1)

    def test_every_layout(self):
        for kw in ({'view_mode': 'all'}, {'view_mode': 'ats'}, {'view_mode': 'incoming'},
                   {'view_mode': 'incoming', 'flow_mode': True},
                   {'view_mode': 'all', 'incoming_only': True},
                   {'view_mode': 'all', 'is_order': True},
                   {'view_mode': 'all', 'catalog_mode': True},
                   {'view_mode': 'ats', 'catalog_mode': True},
                   {'view_mode': 'incoming', 'catalog_mode': True, 'flow_mode': True},
                   {'view_mode': 'ats', 'catalog_mode': True, 'wh_breakdown': True},
                   {'view_mode': 'ats', 'catalog_mode': True, 'tjx_layout': True},
                   {'view_mode': 'incoming', 'catalog_mode': True, 'tjx_layout': True, 'flow_mode': True}):
            with self.subTest(**kw):
                self.assertAfterBrand(self.headers(**kw))

    def test_mailbox_breakdown_order(self):
        self.assertEqual(self.headers(view_mode='ats', catalog_mode=True, wh_breakdown=True),
                         ['IMAGE', 'SKU', 'Brand', 'Category', 'Color', 'Fit', 'Fabrication',
                          'JTW ATS', 'TR ATS', 'DCW ATS', 'QA ATS', 'Warehouse ATS'])

    def test_caller_defined_sets(self):
        self.assertAfterBrand(self.headers(headers_override=self.app.SHIP_PLAN_HEADERS))
        self.assertAfterBrand(self.headers(headers_override=['IMAGE', 'SKU', 'Brand', 'Color', 'Trajectory']))
        self.assertNotIn('Category', self.app.SHIP_PLAN_HEADERS)          # the caller's list is not touched
        self.assertNotIn('Category', self.headers(headers_override=['IMAGE', 'SKU', 'Units']))   # no Brand, no Category
        h = self.headers(headers_override=['IMAGE', 'SKU', 'Category', 'Brand'])   # a caller's own Category stays
        self.assertEqual(h, ['IMAGE', 'SKU', 'Category', 'Brand'])


ITEMS = [row('ROEBFT007SLS', 'EB', 'Eddie Bauer', 1044),
         row('BUGBRPB05BRV', 'BEENE', 'Geoffrey Beene', 900),
         row('AMUSDPP01SRS', 'USPA', 'U.S. Polo Assn.', 500),
         row('BRGBPO001SSM', 'BEENE', 'Geoffrey Beene', 300),
         row('CTH01', 'CHAPS', 'Chaps', 60)]
WANT = ['Dress Shirt', 'Blazer', 'Dress Pants', 'Sportswear', 'Accessories']


def col_of(ws, name, header_row=1):
    for c in range(1, ws.max_column + 1):
        if ws.cell(header_row, c).value == name:
            return c
    raise AssertionError(f'{name} not in the header row')


class Workbooks(Base):

    def values(self, ws, name, first, n, header_row=1):
        c = col_of(ws, name, header_row)
        return [ws.cell(r, c).value for r in range(first, first + n)]

    def test_single_brand_admin_and_customer(self):
        for catalog_mode in (False, True):
            with self.subTest(catalog_mode=catalog_mode):
                data = self.app.build_brand_excel('Mixed', [dict(r) for r in ITEMS], 'http://photos.test/',
                                                  view_mode='ats', catalog_mode=catalog_mode)
                ws = openpyxl.load_workbook(io.BytesIO(data))['Mixed']
                self.assertEqual(col_of(ws, 'Category'), col_of(ws, 'Brand') + 1)
                self.assertEqual(self.values(ws, 'Category', 2, 5), WANT)
                self.assertEqual(self.values(ws, 'SKU', 2, 5), [r['sku'] for r in ITEMS])

    def test_multi_brand_tabs(self):
        tabs = [{'brand_name': 'A', 'items': [dict(r) for r in ITEMS[:2]]},
                {'brand_name': 'B', 'items': [dict(r) for r in ITEMS[2:]], 'view_mode': 'incoming'}]
        wb = openpyxl.load_workbook(io.BytesIO(self.app.build_multi_brand_excel(
            tabs, 'http://photos.test/', catalog_mode=True, view_mode='ats')))
        self.assertEqual(self.values(wb['A'], 'Category', 2, 2), WANT[:2])
        self.assertEqual(self.values(wb['B'], 'Category', 2, 3), WANT[2:])

    def test_tjx_colour_key_still_points_at_its_columns(self):
        tab = [{'brand_name': 'Warehouse', 'tab_name': 'Warehouse', 'view_mode': 'ats', 'keep_order': True,
                'items': [dict(r) for r in ITEMS]}]
        wb = openpyxl.load_workbook(io.BytesIO(self.app.build_multi_brand_excel(
            tab, 'http://photos.test/', catalog_mode=True, view_mode='ats', tjx_layout=True, color_key=True)))
        ws = wb['Warehouse']
        head = self.app.TJX_KEY_HEADER_ROW + 1
        self.assertEqual(col_of(ws, 'Category', head), col_of(ws, 'Brand', head) + 1)
        fam = openpyxl.utils.get_column_letter(col_of(ws, 'Color Family', head))
        ats = openpyxl.utils.get_column_letter(col_of(ws, 'Total ATS', head))
        formulas = ' '.join(str(ws.cell(r, c).value) for r in range(3, 9) for c in range(2, 5))
        self.assertIn(f'${fam}$', formulas)
        self.assertIn(f'${ats}$', formulas)
        cats = [v for v in (ws.cell(r, col_of(ws, 'Category', head)).value for r in range(head + 1, ws.max_row + 1))
                if v in WANT]
        self.assertEqual(sorted(cats), sorted(WANT))

    def test_ship_plan_and_selling_sheets(self):
        tabs = [{'name': 'Plan', 'items': [dict(r, units_ship=10) for r in ITEMS]},
                {'name': 'Best Sellers', 'items': [dict(r) for r in ITEMS[:1]],
                 'headers': ['IMAGE', 'SKU', 'Brand', 'Color', 'Trajectory']}]
        wb = openpyxl.load_workbook(io.BytesIO(self.app.build_ship_plan_excel(tabs, 'http://photos.test/')))
        self.assertEqual(self.values(wb['Plan'], 'Category', 2, 5), WANT)
        self.assertEqual([c.value for c in wb['Best Sellers'][1]], ['IMAGE', 'SKU', 'Brand', 'Category', 'Color', 'Trajectory'])
        self.assertEqual(wb['Best Sellers'].cell(2, 4).value, 'Dress Shirt')

    def test_overseas_summary(self):
        items = [dict(r, production='TF26001', po='PO1', ex_factory='Oct 1, 2026', arrival='Nov 25, 2026',
                      produced=1200, deducted=-200, flow_ats=1000) for r in ITEMS]
        ws = openpyxl.load_workbook(io.BytesIO(self.app.build_overseas_summary_excel(
            'OS', items, 'http://photos.test/')))['Overseas Summary']
        head = [c.value for c in ws[1]]
        self.assertEqual(head[:7], ['IMAGE', 'SKU', 'Brand', 'Category', 'Color', 'Fit', 'Fabrication'])
        data_rows = [r for r in ws.iter_rows(min_row=2, values_only=True) if r[1]]
        self.assertEqual([r[3] for r in data_rows], WANT)
        self.assertEqual([r[1] for r in data_rows], [r['sku'] for r in ITEMS])
        self.assertEqual([r[head.index('Flow ATS')] for r in data_rows], [1000] * 5)
        self.assertEqual([r[head.index('Production')] for r in data_rows], ['TF26001'] * 5)


if __name__ == '__main__':
    unittest.main()
