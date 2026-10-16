"""Overseas line sheet tabs print OVERSEAS availability only (Oct 16 2026).

The Overseas ATS column on a stock:'overseas' tab printed the style's TOTAL
ATS, warehouse stock included: a style with 6,732 units in the warehouse and
468 incoming read 7,200 overseas. The number must come from the same smart
routing split the overseas photo deck uses (_pres_stock_cards ->
_pres_route_sku), so a sheet and a deck can never disagree.

Everything here is synthetic and offline: made-up style numbers (serial 9xx,
fabric ZZ), a made-up feed, ledger and order book, S3 replaced by a recorder.
"""
import io
import os
import sys
import unittest
from datetime import datetime

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from _support import load_app  # noqa: E402

NOW = datetime(2026, 10, 16, 10, 0, 0)

S_SPLIT = 'TJNAZZ931SLS'      # warehouse 6,732 + incoming 468, nothing committed
S_ROUTED = 'TJNAZZ932SLS'     # one order only the warehouse can make, one only production can
S_WHCOVER = 'TJNAZZ933SLS'    # one dated order the warehouse covers
S_AS_OS = 'TJDKZZ934SLS'      # deduction assigned to overseas
S_AS_WH = 'TJDKZZ935SLS'      # deduction assigned to warehouse
S_NJLOT = 'TJDKZZ936SLS'      # one lot lands NJ (admin only), one lands TR
S_NONE = 'TJNAZZ937SLS'       # every incoming unit is spoken for
S_WHONLY = 'TJNAZZ938SLS'     # warehouse only, part of it NJ
S_VAR = 'TJDKZZ939SLS'        # a base and its -V variant both incoming


def _item(sku, brand, jtw=0, nj=0, inc=0, committed=0):
    return {'sku': sku, 'brand_abbr': brand, 'brand': brand, 'jtw': jtw, 'tr': 0, 'dcw': 0, 'qa': 0,
            'nj': nj, 'abfi': 0, 'incoming': inc, 'committed': committed, 'allocated': 0,
            'total_ats': jtw + nj + inc + committed}


FEED = [
    _item(S_SPLIT, 'NAUTICA', jtw=6732, inc=468),
    _item(S_ROUTED, 'NAUTICA', jtw=1000, inc=2000, committed=-2100),
    _item(S_WHCOVER, 'NAUTICA', jtw=1000, inc=2000, committed=-600),
    _item(S_AS_OS, 'DKNY', jtw=1000, inc=2000, committed=-600),
    _item(S_AS_WH, 'DKNY', jtw=1000, inc=2000, committed=-600),
    _item(S_NJLOT, 'DKNY', jtw=100, inc=3000),
    _item(S_NONE, 'NAUTICA', jtw=5000, inc=500, committed=-500),
    _item(S_WHONLY, 'NAUTICA', jtw=300, nj=50),
    _item(S_VAR, 'DKNY', inc=1000),
    _item(S_VAR + '-V', 'DKNY', inc=400),
]


def _lot(style, units, arrival, wh='TR', prod='ZZ26900'):
    return {'production': prod, 'poName': 'P', 'style': style, 'units': units, 'etd': '',
            'arrival': arrival, 'port_dated': True, 'fob_flag': False, 'warehouse': wh, 'shipmentNo': ''}


LEDGER = [
    _lot(S_SPLIT, 468, '2026-12-15'),
    _lot(S_ROUTED, 2000, '2026-12-15'),
    _lot(S_WHCOVER, 2000, '2026-12-15'),
    _lot(S_AS_OS, 2000, '2026-12-15'),
    _lot(S_AS_WH, 2000, '2026-12-15'),
    _lot(S_NJLOT, 1000, '2026-12-01', wh='NJ', prod='ZZ26901'),
    _lot(S_NJLOT, 2000, '2026-12-15', prod='ZZ26902'),
    _lot(S_NONE, 500, '2026-12-15'),
    _lot(S_VAR, 1000, '2026-12-15'),
    _lot(S_VAR + '-V', 400, '2026-12-20'),
]


def _order(style, qty, start, cancel, po):
    return {'orderNo': po, 'ctrlNo': po, 'customer': 'ZCUST', 'customerFull': 'Z Test Stores',
            'style': style, 'openQty': qty, 'pickQty': 0, 'isPipeline': False,
            'startDate': start + 'T00:00:00', 'cancelDate': cancel + 'T00:00:00'}


ORDERS = [
    # Ships before the December lot lands: only the warehouse can make it.
    _order(S_ROUTED, 600, '2026-10-20', '2026-10-30', '9100001'),
    # Ships after the lot lands: smart routing charges it to production.
    _order(S_ROUTED, 1500, '2027-01-10', '2027-01-20', '9100002'),
    _order(S_WHCOVER, 600, '2026-10-20', '2026-10-30', '9100003'),
    _order(S_NONE, 500, '2027-01-10', '2027-01-20', '9100004'),
]

ASSIGN = {S_AS_OS: 'overseas', S_AS_WH: 'warehouse'}


class _Base(unittest.TestCase):
    # A subclass adds its own synthetic styles here, so the pinned numbers of the
    # other classes never move.
    EXTRA_FEED = []
    EXTRA_LEDGER = []
    EXTRA_ORDERS = []

    @classmethod
    def setUpClass(cls):
        cls.app = load_app()

    def setUp(self):
        a = self.app
        names = ('build_multi_brand_excel', 'get_s3', '_production_data', '_fresh_prepack_defaults',
                 '_ai_agent_enrich', '_pres_now_et', '_fetch_all_open_orders', 'load_manual_allocations_from_s3',
                 'load_deduction_assignments_from_s3', 'load_suppression_overrides_from_s3',
                 'load_allocation_from_s3', '_manual_allocations', '_deduction_assignments',
                 '_suppression_overrides', '_apo_data', 'download_images_for_items', '_pres_stock_cards')
        self._saved = {k: getattr(a, k) for k in names}
        self._inv_items = a._inventory.get('items')
        a._inventory['items'] = [dict(i) for i in FEED + self.EXTRA_FEED]
        a._production_data = [dict(p) for p in LEDGER + self.EXTRA_LEDGER]
        a._fresh_prepack_defaults = lambda: None
        a._ai_agent_enrich = lambda r: ('White Solid', 'Slim Fit', 'Poly')
        a._pres_now_et = lambda: NOW
        a._fetch_all_open_orders = lambda: ([dict(o) for o in ORDERS + self.EXTRA_ORDERS], True)
        for k in ('load_manual_allocations_from_s3', 'load_deduction_assignments_from_s3',
                  'load_suppression_overrides_from_s3'):
            setattr(a, k, lambda: None)
        a.load_allocation_from_s3 = lambda: []
        a._manual_allocations = []
        a._deduction_assignments = dict(ASSIGN)
        a._suppression_overrides = []
        a._apo_data = []
        a.download_images_for_items = lambda items, url, use_cache=True: {}
        self.built = {}
        self.real_excel = self._saved['build_multi_brand_excel']

        def fake_excel(tabs, *args, **kw):
            self.built['tabs'] = tabs
            self.built['kw'] = kw
            return b'xlsx'
        a.build_multi_brand_excel = fake_excel

        class S3:
            def put_object(_s, **kw):
                self.built['body'] = kw['Body']
        a.get_s3 = lambda: S3()

    def tearDown(self):
        for k, v in self._saved.items():
            setattr(self.app, k, v)
        self.app._inventory['items'] = self._inv_items

    # helpers
    def sheet(self, tabs, customer_view=False, breakdown=False):
        return self.app._ai_tool_build_line_sheet({'tabs': tabs, 'customer_view': customer_view,
                                                   'warehouse_breakdown': breakdown})

    def items(self, tab_index=0):
        return {it['sku']: it for it in self.built['tabs'][tab_index]['items']}

    def printed(self, tab_index=0):
        return {s: it['total_ats'] for s, it in self.items(tab_index).items()}

    def deck(self, customer_view, params=None):
        """What the overseas photo deck prints per base style for the same inputs."""
        got = self.app._pres_stock_cards('overseas', customer_view, self.app._pres_filter_spec({}), params or {})
        out = {}
        for c in got['cards']:
            out[c['base']] = out.get(c['base'], 0) + c['number']
        return out

    @staticmethod
    def entry(out, title):
        return next(e for e in out['tabs'] if e.get('tab') == title)


class OverseasTabs(_Base):

    # a. warehouse stock never reaches the Overseas ATS column
    def test_overseas_tab_prints_only_the_overseas_share(self):
        out = self.sheet([{'title': 'OS', 'skus': [S_SPLIT], 'stock': 'overseas'}])
        self.assertNotIn('error', out)
        self.assertEqual(self.built['tabs'][0]['view_mode'], 'incoming')
        self.assertEqual(self.printed(), {S_SPLIT: 468})     # was 7,200 (6,732 + 468)

    def test_customer_view_overseas_tab_prints_only_the_overseas_share(self):
        self.sheet([{'title': 'OS', 'skus': [S_SPLIT], 'stock': 'overseas'}], customer_view=True)
        self.assertEqual(self.printed(), {S_SPLIT: 468})

    # b. smart routing decides which side pays for committed demand
    def test_order_only_the_warehouse_can_make_leaves_incoming_whole(self):
        self.sheet([{'title': 'OS', 'skus': [S_WHCOVER], 'stock': 'overseas'}])
        self.assertEqual(self.printed(), {S_WHCOVER: 2000})   # was 2,400 total ATS

    def test_order_only_production_can_make_is_charged_to_incoming(self):
        # 600 ships in October (warehouse), 1,500 in January (December lot):
        # overseas = 2,000 - 1,500. Warehouse-first FIFO would say 900.
        self.sheet([{'title': 'OS', 'skus': [S_ROUTED], 'stock': 'overseas'}])
        self.assertEqual(self.printed(), {S_ROUTED: 500})
        route = self.app._pres_route_sku(
            {'jtw': 1000, 'tr': 0, 'dcw': 0, 'qa': 0, 'nj': 0, 'abfi': 0, 'committed': -2100, 'allocated': 0,
             'incoming': 2000},
            [{'units': 2000, 'arr': datetime(2026, 12, 15).date(), 'etd': None, 'fob_flag': False,
              'hidden': False, 'sup': False}],
            [o for o in ORDERS if o['style'] == S_ROUTED], [], [], NOW.date(), sku=S_ROUTED)
        self.assertEqual(route, {'wh': 600, 'os': 1500, 'os_visible': 1500})

    def test_sheet_equals_the_overseas_deck_admin_and_customer(self):
        bases = [S_SPLIT, S_ROUTED, S_WHCOVER, S_AS_OS, S_AS_WH, S_NJLOT, S_VAR]
        for cv in (False, True):
            with self.subTest(customer_view=cv):
                self.sheet([{'title': 'OS', 'skus': bases, 'stock': 'overseas'}], customer_view=cv)
                deck = self.deck(cv)
                self.assertEqual(self.printed(), {b: deck[b] for b in bases})

    # c. deduction assignments win over routing
    def test_deduction_assignments_are_respected(self):
        self.sheet([{'title': 'OS', 'skus': [S_AS_OS, S_AS_WH], 'stock': 'overseas'}])
        self.assertEqual(self.printed(), {S_AS_OS: 1400, S_AS_WH: 2000})

    # d. hidden-landing lots
    def test_customer_view_drops_an_nj_landing_lot_admin_counts_it(self):
        self.sheet([{'title': 'OS', 'skus': [S_NJLOT], 'stock': 'overseas'}], customer_view=True)
        self.assertEqual(self.printed(), {S_NJLOT: 2000})
        self.sheet([{'title': 'OS', 'skus': [S_NJLOT], 'stock': 'overseas'}], customer_view=False)
        self.assertEqual(self.printed(), {S_NJLOT: 3000})

    def test_a_variant_row_adds_to_its_base_like_the_deck(self):
        self.sheet([{'title': 'OS', 'skus': [S_VAR], 'stock': 'overseas'}])
        self.assertEqual(self.printed(), {S_VAR: 1400})

    # e. nothing left overseas: off the tab, and said so
    def test_style_with_no_overseas_left_is_named_on_a_curated_tab(self):
        out = self.sheet([{'title': 'OS', 'skus': [S_SPLIT, S_NONE, S_WHONLY], 'stock': 'overseas'}])
        self.assertEqual(list(self.items()), [S_SPLIT])
        e = self.entry(out, 'OS')
        self.assertEqual(e['styles'], 1)
        self.assertEqual(e['styles_without_overseas_availability'], [S_NONE, S_WHONLY])

    def test_style_with_no_overseas_left_is_counted_on_a_filter_tab(self):
        out = self.sheet([{'title': 'NA OS', 'brands': ['NAUTICA'], 'stock': 'overseas'}])
        got = self.printed()
        self.assertNotIn(S_NONE, got)
        self.assertEqual(got, {S_SPLIT: 468, S_ROUTED: 500, S_WHCOVER: 2000})
        e = self.entry(out, 'NA OS')
        self.assertEqual(e['styles'], 3)
        self.assertEqual(e['styles_without_overseas_availability'], 1)

    def test_a_tab_left_empty_is_skipped_with_a_reason(self):
        out = self.sheet([{'title': 'OS', 'skus': [S_NONE], 'stock': 'overseas'},
                          {'title': 'WH', 'skus': [S_NONE], 'stock': 'warehouse'}])
        self.assertEqual([t['tab_name'] for t in self.built['tabs']], ['WH'])
        e = self.entry(out, 'OS')
        self.assertTrue(e['skipped'])
        self.assertEqual(e['styles'], 0)
        self.assertIn('overseas', e['reason'])

    def test_every_tab_empty_is_an_error(self):
        out = self.sheet([{'title': 'OS', 'skus': [S_NONE, S_WHONLY], 'stock': 'overseas'}])
        self.assertIn('error', out)
        self.assertNotIn('tabs', self.built)

    # g. a routing failure is an error, never a sheet with total ATS
    def test_routing_failure_returns_an_error(self):
        def boom(*a, **k):
            raise RuntimeError('ledger exploded')
        self.app._pres_stock_cards = boom
        out = self.sheet([{'title': 'OS', 'skus': [S_SPLIT], 'stock': 'overseas'}])
        self.assertIn('error', out)
        self.assertIn('overseas', out['error'])
        self.assertNotIn('tabs', self.built)
        self.assertNotIn('body', self.built)

    def test_routing_runs_once_per_build(self):
        calls = []
        real = self._saved['_pres_stock_cards']

        def spy(*a, **k):
            calls.append(a)
            return real(*a, **k)
        self.app._pres_stock_cards = spy
        self.sheet([{'title': 'A', 'skus': [S_SPLIT], 'stock': 'overseas'},
                    {'title': 'B', 'skus': [S_ROUTED], 'stock': 'overseas'},
                    {'title': 'C', 'skus': [S_SPLIT], 'stock': 'warehouse'}])
        self.assertEqual(len(calls), 1)
        self.assertEqual(calls[0][0], 'overseas')

    def test_no_overseas_tab_never_runs_the_routing(self):
        def boom(*a, **k):
            raise AssertionError('must not run')
        self.app._pres_stock_cards = boom
        out = self.sheet([{'title': 'WH', 'skus': [S_SPLIT], 'stock': 'warehouse'}])
        self.assertNotIn('error', out)

    # row order: a filter tab sorts on the number it prints
    def test_filter_tab_rows_sort_by_the_overseas_number(self):
        # The filter orders by gross total ATS (S_SPLIT 7,200 with its warehouse),
        # which put a 468 overseas row above 2,000 and 500.
        self.sheet([{'title': 'NA OS', 'brands': ['NAUTICA'], 'stock': 'overseas'}])
        self.assertEqual([(s, it['total_ats']) for s, it in self.items().items()],
                         [(S_WHCOVER, 2000), (S_ROUTED, 500), (S_SPLIT, 468)])

    def test_filter_tab_sort_is_stable(self):
        # S_AS_OS and S_VAR both print 1,400: they keep the filter's order.
        self.sheet([{'title': 'DK OS', 'brands': ['DKNY'], 'stock': 'overseas'}])
        self.assertEqual(list(self.items()), [S_NJLOT, S_AS_WH, S_AS_OS, S_VAR])

    def test_curated_tab_keeps_the_callers_order(self):
        self.sheet([{'title': 'OS', 'skus': [S_SPLIT, S_ROUTED, S_WHCOVER], 'stock': 'overseas'}])
        self.assertEqual(list(self.items()), [S_SPLIT, S_ROUTED, S_WHCOVER])

    # min_units: applied again to the routed number on a filter tab
    def test_filter_tab_min_units_applies_to_the_routed_number(self):
        # S_ROUTED passes the filter on 2,000 gross incoming but only 500 is left.
        out = self.sheet([{'title': 'NA OS', 'brands': ['NAUTICA'], 'stock': 'overseas', 'min_units': 600}])
        self.assertEqual(self.printed(), {S_WHCOVER: 2000})
        e = self.entry(out, 'NA OS')
        self.assertEqual(e['styles'], 1)
        self.assertEqual(e['styles_below_min_units'], 1)
        self.assertNotIn('styles_without_overseas_availability', e)

    def test_tab_emptied_by_min_units_is_skipped_with_a_reason(self):
        out = self.sheet([{'title': 'OS', 'search': '932', 'stock': 'overseas', 'min_units': 600},
                          {'title': 'WH', 'skus': [S_ROUTED], 'stock': 'warehouse'}])
        self.assertEqual([t['tab_name'] for t in self.built['tabs']], ['WH'])
        e = self.entry(out, 'OS')
        self.assertTrue(e['skipped'])
        self.assertEqual(e['styles'], 0)
        self.assertEqual(e['styles_below_min_units'], 1)
        self.assertIn('600', e['reason'])

    def test_skus_tab_ignores_min_units(self):
        # skus overrides every other filter on a tab (the tool's contract).
        out = self.sheet([{'title': 'OS', 'skus': [S_ROUTED], 'stock': 'overseas', 'min_units': 600}])
        self.assertEqual(self.printed(), {S_ROUTED: 500})
        self.assertNotIn('styles_below_min_units', self.entry(out, 'OS'))

    # arrival window on a base style whose NJ lot sits inside it
    def test_customer_view_window_never_counts_a_hidden_lot(self):
        tab = {'title': 'OS', 'search': '936', 'stock': 'overseas', 'arrive_before': '2026-12-10'}
        self.sheet([tab])
        win = {'arrive_before': '2026-12-10'}
        # admin: only the NJ lot (1,000 of 3,000) is inside the window
        self.assertEqual(self.printed(), {S_NJLOT: 1000})
        self.assertEqual(self.deck(False, win)[S_NJLOT], 1000)
        # customer: the only lot inside the window is hidden, so nothing counts
        out = self.sheet([tab], customer_view=True)
        self.assertNotIn(S_NJLOT, self.deck(True, win))
        self.assertIn('error', out)
        e = self.entry(out, 'OS')
        self.assertEqual(e['styles_without_overseas_availability'], 1)
        # nothing is committed: the reason names the window, never committed demand
        self.assertIn('window', e['reason'])
        self.assertNotIn('committed', e['reason'])
        self.assertEqual(e['arrival_window'], {'arrive_before': '2026-12-10'})

    def test_summary_says_when_the_order_feed_is_down(self):
        tab = {'title': 'OS', 'skus': [S_SPLIT], 'stock': 'overseas'}
        note = next(e for e in self.sheet([tab])['tabs'] if 'overseas_ats' in e)
        self.assertIs(note['open_orders_feed_ok'], True)
        self.app._fetch_all_open_orders = lambda: ([], False)
        note = next(e for e in self.sheet([tab])['tabs'] if 'overseas_ats' in e)
        self.assertIs(note['open_orders_feed_ok'], False)

    # the mailbox: customer view + per-warehouse breakdown
    def test_mailbox_overseas_tab_keeps_its_styles_after_the_breakdown_pass(self):
        out = self.sheet([{'title': 'OS', 'skus': [S_SPLIT, S_ROUTED, S_NONE], 'stock': 'overseas'}],
                         customer_view=True, breakdown=True)
        self.assertNotIn('error', out)
        self.assertEqual(self.printed(), {S_SPLIT: 468, S_ROUTED: 500})
        e = self.entry(out, 'OS')
        self.assertEqual(e['styles'], 2)
        self.assertEqual(e['styles_without_overseas_availability'], [S_NONE])

    # the printed column itself
    def test_workbook_overseas_ats_cell(self):
        import openpyxl
        self.app.build_multi_brand_excel = self.real_excel
        out = self.sheet([{'title': 'OS', 'skus': [S_SPLIT, S_ROUTED], 'stock': 'overseas'}],
                         customer_view=True)
        self.assertNotIn('error', out)
        wb = openpyxl.load_workbook(io.BytesIO(self.built['body']))
        ws = wb['OS']
        head = [c.value for c in ws[1]]
        sku_col, os_col = head.index('SKU'), head.index('Overseas ATS')
        cells = {row[sku_col].value: row[os_col].value
                 for row in ws.iter_rows(min_row=2) if row[sku_col].value in (S_SPLIT, S_ROUTED)}
        self.assertEqual(cells, {S_SPLIT: 468, S_ROUTED: 500})


S_HIDBASE = 'TJDKZZ940SLS'    # base: one NJ lot only; its -V: one TR lot, no warehouse stock
S_NJONLY = 'TJDKZZ942SLS'     # one NJ lot, nothing else


class HiddenBaseVisibleVariant(_Base):
    """Customer view kept a hidden-landing gate keyed on the EXACT ledger style, so a
    base whose only lot lands NJ was dropped from an overseas tab although its -V
    has a visible lot the overseas deck shows (Oct 16 2026)."""

    EXTRA_FEED = [_item(S_HIDBASE, 'DKNY', inc=1000), _item(S_HIDBASE + '-V', 'DKNY', inc=400),
                  _item(S_NJONLY, 'DKNY', inc=700)]
    # The hidden lot lands FIRST, so a leak would show its date and production.
    EXTRA_LEDGER = [_lot(S_HIDBASE, 1000, '2026-11-20', wh='NJ', prod='ZZ26940'),
                    _lot(S_HIDBASE + '-V', 400, '2026-12-20', prod='ZZ26941'),
                    _lot(S_NJONLY, 700, '2026-11-25', wh='NJ', prod='ZZ26942')]

    def test_customer_view_prints_the_variant_lot_the_deck_shows(self):
        for tab in ({'title': 'OS', 'skus': [S_HIDBASE], 'stock': 'overseas'},
                    {'title': 'OS', 'search': '940', 'stock': 'overseas'}):
            with self.subTest(tab=tab):
                out = self.sheet([tab], customer_view=True)
                self.assertNotIn('error', out)
                self.assertEqual(self.printed(), {S_HIDBASE: 400})
                self.assertEqual(self.deck(True)[S_HIDBASE], 400)
                it = self.items()[S_HIDBASE]
                # dates and PO Ref # name the visible lot, never the NJ one
                self.assertEqual(it['arrival'], self.app._apo_fmt_date(datetime(2026, 12, 20).date()))
                self.assertEqual(it['po_ref'], 'ZZ26941')

    def test_admin_view_counts_both_lots(self):
        self.sheet([{'title': 'OS', 'skus': [S_HIDBASE], 'stock': 'overseas'}])
        self.assertEqual(self.printed(), {S_HIDBASE: 1400})
        self.assertEqual(self.deck(False)[S_HIDBASE], 1400)

    def test_hidden_only_style_is_named_and_never_printed(self):
        out = self.sheet([{'title': 'OS', 'skus': [S_HIDBASE, S_NJONLY], 'stock': 'overseas'}],
                         customer_view=True)
        self.assertEqual(self.printed(), {S_HIDBASE: 400})
        self.assertEqual(self.entry(out, 'OS')['styles_without_overseas_availability'], [S_NJONLY])

    def test_mailbox_breakdown_keeps_the_variant_lot_style(self):
        # The mailbox's breakdown pass kept an overseas row on its gross incoming, which
        # the exact-style hidden cut leaves at 0 here: the style vanished uncounted.
        out = self.sheet([{'title': 'OS', 'skus': [S_HIDBASE, S_SPLIT], 'stock': 'overseas'}],
                         customer_view=True, breakdown=True)
        self.assertNotIn('error', out)
        self.assertEqual(self.printed(), {S_HIDBASE: 400, S_SPLIT: 468})
        self.assertEqual(self.entry(out, 'OS')['styles'], 2)

    def test_warehouse_and_any_tabs_keep_the_gate(self):
        out = self.sheet([{'title': 'OS', 'skus': [S_HIDBASE], 'stock': 'overseas'},
                          {'title': 'WH', 'skus': [S_HIDBASE], 'stock': 'warehouse'},
                          {'title': 'ALL', 'skus': [S_HIDBASE]}], customer_view=True)
        self.assertEqual([t['tab_name'] for t in self.built['tabs']], ['OS'])
        for title in ('WH', 'ALL'):
            self.assertTrue(self.entry(out, title)['skipped'])
            self.assertEqual(self.entry(out, title)['styles'], 0)


S_WIN = 'TJDKZZ941SLS'        # 600 lands in November, 400 in January; 200 committed


class ArrivalWindows(_Base):
    """A filter tab's arrival window reaches the deck computation: only lots inside it
    count and the number scales by their share, the deck's windowed rule (Oct 16 2026)."""

    EXTRA_FEED = [_item(S_WIN, 'DKNY', inc=1000, committed=-200)]
    EXTRA_LEDGER = [_lot(S_WIN, 600, '2026-11-20', prod='ZZ26950'),
                    _lot(S_WIN, 400, '2027-01-20', prod='ZZ26951')]
    EXTRA_ORDERS = [_order(S_WIN, 200, '2027-02-10', '2027-02-20', '9100005')]

    def test_window_prints_the_decks_scaled_number(self):
        win = {'arrive_before': '2026-12-31'}
        for cv in (False, True):
            with self.subTest(customer_view=cv):
                out = self.sheet([dict(win, title='W', search='941', stock='overseas')], customer_view=cv)
                # 1,000 - 200 committed = 800 overseas; the November lot is 60% of it
                self.assertEqual(self.printed(), {S_WIN: 480})
                self.assertEqual(self.deck(cv, win)[S_WIN], 480)
                self.assertTrue(any('arrival_window' in e for e in out['tabs']))

    def test_no_window_prints_the_whole_overseas_number(self):
        out = self.sheet([{'title': 'W', 'search': '941', 'stock': 'overseas'}])
        self.assertEqual(self.printed(), {S_WIN: 800})
        self.assertFalse(any('arrival_window' in e for e in out['tabs']))

    def test_one_routing_run_per_distinct_window(self):
        calls = []
        real = self._saved['_pres_stock_cards']

        def spy(*a, **k):
            calls.append(a)
            return real(*a, **k)
        self.app._pres_stock_cards = spy
        self.sheet([{'title': 'A', 'search': '941', 'stock': 'overseas', 'arrive_before': '2026-12-31'},
                    # the same window spelled another way shares A's run
                    {'title': 'B', 'search': '941', 'stock': 'overseas', 'arrive_before': '12/31/2026'},
                    {'title': 'C', 'search': '941', 'stock': 'overseas'},
                    # a skus tab ignores the window, so it shares C's run
                    {'title': 'D', 'skus': [S_WIN], 'stock': 'overseas', 'arrive_before': '2026-12-31'},
                    {'title': 'F', 'search': '941', 'stock': 'overseas', 'arrive_after': '2027-01-01'},
                    {'title': 'E', 'skus': [S_WIN], 'stock': 'warehouse'}])
        self.assertEqual(len(calls), 3)
        self.assertEqual(sorted((sorted(c[3].items()) for c in calls), key=str),
                         sorted([[], [('arrive_before', '2026-12-31')], [('arrive_after', '2027-01-01')]], key=str))
        got = {t['tab_name']: {it['sku']: it['total_ats'] for it in t['items']} for t in self.built['tabs']}
        self.assertEqual(got['A'], {S_WIN: 480})
        self.assertEqual(got['B'], {S_WIN: 480})
        self.assertEqual(got['C'], {S_WIN: 800})
        self.assertEqual(got['D'], {S_WIN: 800})
        self.assertEqual(got['F'], {S_WIN: 320})   # the January lot is 40%
        self.assertEqual(self.deck(False, {'arrive_after': '2027-01-01'})[S_WIN], 320)


class OtherTabsUnchanged(_Base):
    """f. Warehouse and All tabs, and the per-warehouse columns, print exactly what
    they printed before the overseas fix (values pinned from the old code)."""

    BASES = [S_SPLIT, S_ROUTED, S_NONE, S_WHONLY, S_NJLOT]

    def test_warehouse_and_any_tabs_admin(self):
        self.sheet([{'title': 'WH', 'skus': self.BASES, 'stock': 'warehouse'},
                    {'title': 'ALL', 'skus': self.BASES}])
        pinned = {S_SPLIT: 7200, S_ROUTED: 900, S_NONE: 5000, S_WHONLY: 350, S_NJLOT: 3100}
        self.assertEqual(self.printed(0), pinned)
        self.assertEqual(self.printed(1), pinned)
        self.assertEqual([t['view_mode'] for t in self.built['tabs']], ['ats', 'all'])
        self.assertEqual(self.items(0)[S_ROUTED]['incoming'], 2000)
        self.assertEqual(self.items(0)[S_ROUTED]['total_warehouse'], 1000)

    def test_warehouse_and_any_tabs_customer(self):
        self.sheet([{'title': 'WH', 'skus': self.BASES, 'stock': 'warehouse'},
                    {'title': 'ALL', 'skus': self.BASES}], customer_view=True)
        pinned = {S_SPLIT: 7200, S_ROUTED: 900, S_NONE: 5000, S_WHONLY: 300, S_NJLOT: 2100}
        self.assertEqual(self.printed(0), pinned)
        self.assertEqual(self.printed(1), pinned)

    def test_filter_tabs_unchanged(self):
        self.sheet([{'title': 'WH', 'brands': ['NAUTICA'], 'stock': 'warehouse'},
                    {'title': 'ALL', 'brands': ['DKNY']}])
        self.assertEqual(self.printed(0), {S_SPLIT: 7200, S_NONE: 5000, S_WHCOVER: 2400, S_ROUTED: 900,
                                           S_WHONLY: 350})
        self.assertEqual(self.printed(1), {S_NJLOT: 3100, S_AS_OS: 2400, S_AS_WH: 2400, S_VAR: 1400})

    def test_mailbox_breakdown_columns_unchanged(self):
        out = self.sheet([{'title': 'WH', 'skus': self.BASES, 'stock': 'warehouse'},
                          {'title': 'ALL', 'skus': self.BASES}], customer_view=True, breakdown=True)
        self.assertNotIn('error', out)
        for ti in (0, 1):
            cols = {s: (it['wh_ats_jtw'], it['wh_ats_tr'], it['wh_ats_dcw'], it['wh_ats_qa'], it['wh_ats_total'])
                    for s, it in self.items(ti).items()}
            self.assertEqual(cols, {S_SPLIT: (6732, 0, 0, 0, 6732), S_ROUTED: (400, 0, 0, 0, 400),
                                    S_NONE: (5000, 0, 0, 0, 5000), S_WHONLY: (300, 0, 0, 0, 300),
                                    S_NJLOT: (100, 0, 0, 0, 100)})
            self.assertEqual(self.printed(ti), {S_SPLIT: 7200, S_ROUTED: 900, S_NONE: 5000, S_WHONLY: 300,
                                                S_NJLOT: 2100})


if __name__ == '__main__':
    unittest.main()
