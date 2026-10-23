"""A line sheet never prints a row below the minimum asked for (Oct 23 2026).

A colleague emailed the mailbox for Chaps and Geoffrey Beene solid dress shirts with a
36 piece minimum of what is available, and the sheet still carried rows of 24, 21
and 18. The filter checked the minimum against units ON HAND, before committed and
allocated come out (BUCHPT601SLS: 2,004 on hand, 1,980 allocated), and the
mailbox's per-warehouse pass printed what was left without checking it again. Now
min_units is checked against the number each tab prints:

- warehouse tabs on an emailed sheet: Warehouse ATS (JTW + TR + DCW + QA ATS);
- warehouse and overseas tabs elsewhere: the deck's number (already so for filter
  tabs since Oct 16, now for skus tabs too);
- all-stock tabs: Total ATS.

Everything here is synthetic and offline: made-up style numbers (serial 9xx,
fabric ZZ), a made-up feed, ledger and order book, S3 replaced by a recorder.
The shared fixtures come from test_line_sheet_overseas.py.
"""
import io
import os
import sys
import unittest

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
try:
    from tests.test_line_sheet_overseas import (  # noqa: E402
        S_AS_OS, S_AS_WH, S_NJLOT, S_ROUTED, S_SPLIT, _Base)
except ImportError:   # run from inside tests/
    from test_line_sheet_overseas import (  # noqa: E402
        S_AS_OS, S_AS_WH, S_NJLOT, S_ROUTED, S_SPLIT, _Base)

S_LOW = 'TJNAZZ981SLS'     # the reported shape: 2,004 at TR, 1,980 allocated, 24 left
S_OK = 'TJNAZZ982SLS'      # 2,004 at TR, 504 allocated, 1,500 left
S_EDGE = 'TJNAZZ983SLS'    # 100 at JTW, 64 committed, exactly 36 left
S_TWO = 'TJNAZZ984SLS'     # 300 at JTW and 300 at TR, 500 committed: 100 left, split inferred


def _row(sku, jtw=0, tr=0, committed=0, allocated=0):
    return {'sku': sku, 'brand_abbr': 'NAUTICA', 'brand': 'NAUTICA', 'jtw': jtw, 'tr': tr, 'dcw': 0, 'qa': 0,
            'nj': 0, 'abfi': 0, 'incoming': 0, 'committed': committed, 'allocated': allocated,
            'total_ats': jtw + tr + committed + allocated}


class _MinBase(_Base):
    EXTRA_FEED = [_row(S_LOW, tr=2004, allocated=-1980), _row(S_OK, tr=2004, allocated=-504),
                  _row(S_EDGE, jtw=100, committed=-64), _row(S_TWO, jtw=300, tr=300, committed=-500)]

    def mail(self, tabs):
        """What the mailbox builds: customer view with the per-warehouse columns."""
        return self.sheet(tabs, customer_view=True, breakdown=True)

    def wh_printed(self, tab_index=0):
        return {s: it['wh_ats_total'] for s, it in self.items(tab_index).items()}


class MailboxWarehouseTabs(_MinBase):

    def test_the_emailed_request_drops_every_row_below_the_minimum(self):
        out = self.mail([{'title': 'WH', 'search': '98', 'stock': 'warehouse', 'min_units': 36}])
        self.assertNotIn('error', out)
        self.assertEqual(self.wh_printed(), {S_OK: 1500, S_TWO: 100, S_EDGE: 36})
        e = self.entry(out, 'WH')
        self.assertEqual(e['styles'], 3)
        self.assertEqual(e['styles_below_min_units'], 1)          # a count on a filter tab
        bd = next(x for x in out['tabs'] if x.get('warehouse_breakdown'))
        self.assertEqual(bd['styles_dropped_no_warehouse_stock'], 0)   # below the minimum is not "no stock"

    def test_without_a_minimum_the_sheet_is_unchanged(self):
        out = self.mail([{'title': 'WH', 'search': '98', 'stock': 'warehouse'}])
        self.assertEqual(self.wh_printed(), {S_LOW: 24, S_OK: 1500, S_TWO: 100, S_EDGE: 36})
        self.assertNotIn('styles_below_min_units', self.entry(out, 'WH'))

    def test_curated_tab_names_the_styles_below_the_minimum(self):
        out = self.mail([{'title': 'WH', 'skus': [S_LOW, S_OK, S_EDGE], 'stock': 'warehouse', 'min_units': 36}])
        self.assertEqual(list(self.items()), [S_OK, S_EDGE])      # the caller's order, minus the low row
        self.assertEqual(self.entry(out, 'WH')['styles_below_min_units'], [S_LOW])

    def test_a_tab_emptied_by_the_minimum_is_skipped_with_a_reason(self):
        out = self.mail([{'title': 'LOW', 'skus': [S_LOW], 'stock': 'warehouse', 'min_units': 36},
                         {'title': 'OK', 'skus': [S_OK], 'stock': 'warehouse'}])
        self.assertEqual([t['tab_name'] for t in self.built['tabs']], ['OK'])
        e = self.entry(out, 'LOW')
        self.assertTrue(e['skipped'])
        self.assertEqual(e['styles'], 0)
        self.assertEqual(e['styles_below_min_units'], [S_LOW])
        self.assertIn('36', e['reason'])
        self.assertIn('warehouse', e['reason'])

    def test_every_tab_emptied_by_the_minimum_is_an_error_that_says_so(self):
        out = self.mail([{'title': 'LOW', 'skus': [S_LOW], 'stock': 'warehouse', 'min_units': 36}])
        self.assertIn('error', out)
        self.assertIn('minimum', out['error'])
        self.assertEqual(self.entry(out, 'LOW')['styles_below_min_units'], [S_LOW])
        self.assertNotIn('body', self.built)

    def test_inferred_list_names_only_styles_in_the_file(self):
        out = self.mail([{'title': 'WH', 'search': '98', 'stock': 'warehouse'}])
        bd = next(x for x in out['tabs'] if x.get('warehouse_breakdown'))
        self.assertEqual(bd['styles_where_the_warehouse_was_inferred'], [S_TWO])
        self.assertEqual(bd['styles_split_exactly'], 3)
        out = self.mail([{'title': 'WH', 'search': '98', 'stock': 'warehouse', 'min_units': 150}])
        self.assertEqual(self.wh_printed(), {S_OK: 1500})
        bd = next(x for x in out['tabs'] if x.get('warehouse_breakdown'))
        self.assertEqual(bd['styles_where_the_warehouse_was_inferred'], [])
        self.assertEqual(bd['styles_split_exactly'], 1)

    def test_workbook_warehouse_ats_cells_never_below_the_minimum(self):
        import openpyxl
        self.app.build_multi_brand_excel = self.real_excel
        out = self.mail([{'title': 'WH', 'search': '98', 'stock': 'warehouse', 'min_units': 36}])
        self.assertNotIn('error', out)
        ws = openpyxl.load_workbook(io.BytesIO(self.built['body']))['WH']
        head = [c.value for c in ws[1]]
        sku_col, col = head.index('SKU'), head.index('Warehouse ATS')
        cells = {row[sku_col].value: row[col].value for row in ws.iter_rows(min_row=2)
                 if str(row[sku_col].value or '').startswith('TJNAZZ98')}
        self.assertEqual(cells, {S_OK: 1500, S_TWO: 100, S_EDGE: 36})

    def test_mailbox_overseas_tab_honors_the_minimum(self):
        out = self.mail([{'title': 'OS', 'skus': [S_SPLIT, S_ROUTED], 'stock': 'overseas', 'min_units': 480}])
        self.assertEqual(self.printed(), {S_ROUTED: 500})
        self.assertEqual(self.entry(out, 'OS')['styles_below_min_units'], [S_SPLIT])


class AllStockTabs(_MinBase):
    """An all-stock tab prints Total ATS; its filter checked the minimum on warehouse
    plus incoming before committed came out (S_AS_OS: 3,000 gross, 2,400 ATS)."""

    def test_filter_tab_admin(self):
        out = self.sheet([{'title': 'ALL', 'brands': ['DKNY'], 'min_units': 2500}])
        self.assertEqual(self.printed(), {S_NJLOT: 3100})
        e = self.entry(out, 'ALL')
        self.assertEqual(e['styles'], 1)
        self.assertEqual(e['styles_below_min_units'], 2)

    def test_filter_tab_customer_view_counts_what_the_customer_sees(self):
        # S_NJLOT's NJ lot is admin-only: the customer sheet prints 2,100, below 2,500.
        out = self.sheet([{'title': 'ALL', 'brands': ['DKNY'], 'min_units': 2500}], customer_view=True)
        self.assertIn('error', out)
        self.assertIn('minimum', out['error'])
        e = self.entry(out, 'ALL')
        self.assertTrue(e['skipped'])
        self.assertEqual(e['styles_below_min_units'], 3)
        self.assertIn('2,500', e['reason'])

    def test_curated_tab_names_the_styles_below_the_minimum(self):
        out = self.sheet([{'title': 'ALL', 'skus': [S_NJLOT, S_AS_OS], 'min_units': 2500}])
        self.assertEqual(self.printed(), {S_NJLOT: 3100})
        self.assertEqual(self.entry(out, 'ALL')['styles_below_min_units'], [S_AS_OS])

    def test_mailbox_all_tab(self):
        out = self.mail([{'title': 'ALL', 'skus': [S_NJLOT, S_AS_OS, S_SPLIT], 'min_units': 2200}])
        self.assertEqual(self.printed(), {S_AS_OS: 2400, S_SPLIT: 7200})
        self.assertEqual(self.entry(out, 'ALL')['styles_below_min_units'], [S_NJLOT])

    def test_without_a_minimum_nothing_changes(self):
        out = self.sheet([{'title': 'ALL', 'skus': [S_NJLOT, S_AS_OS, S_AS_WH]}])
        self.assertEqual(self.printed(), {S_NJLOT: 3100, S_AS_OS: 2400, S_AS_WH: 2400})
        self.assertNotIn('styles_below_min_units', self.entry(out, 'ALL'))

    def test_unreadable_minimum_is_no_minimum(self):
        out = self.sheet([{'title': 'ALL', 'skus': [S_NJLOT, S_AS_OS], 'min_units': 'lots'}])
        self.assertEqual(self.printed(), {S_NJLOT: 3100, S_AS_OS: 2400})
        self.assertNotIn('styles_below_min_units', self.entry(out, 'ALL'))


class ToolText(_MinBase):
    """The model is told the minimum goes on every tab and what query_inventory's
    own min_units means."""

    def _tool(self, name):
        return next(t for t in self.app._AI_AGENT_TOOLS if t['name'] == name)

    def test_line_sheet_tool_says_min_units_applies_to_every_tab(self):
        d = self._tool('build_line_sheet')['description']
        self.assertIn('except min_units', d)
        self.assertIn('EVERY tab', d)
        self.assertIn('styles_below_min_units', d)

    def test_query_inventory_says_its_minimum_is_not_an_ats_minimum(self):
        self.assertIn('not an ATS minimum', self._tool('query_inventory')['description'])

    def test_every_surface_gets_the_minimums_rule(self):
        core = self.app._AI_AGENT_TOOL_GUIDANCE_CORE
        self.assertIn('MINIMUMS', core)
        self.assertIn('min_units', core)
        self.assertIn('MINIMUMS', self.app._AI_AGENT_TOOL_GUIDANCE)

    def test_no_em_dash_in_the_new_text(self):
        for text in (self._tool('build_line_sheet')['description'], self._tool('query_inventory')['description']):
            seg = text[text.find('min_units'):]
            self.assertNotIn(chr(0x2014), seg[:600])
        core = self.app._AI_AGENT_TOOL_GUIDANCE_CORE
        rule = core[core.find('MINIMUMS'):]
        self.assertNotIn(chr(0x2014), rule[:rule.find('\n')])


if __name__ == '__main__':
    unittest.main()
