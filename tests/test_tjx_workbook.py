"""TJX-layout workbook rendering (Oct 10 2026): brand bands with borders, zebra restart,
images staying on their rows, the size-grid anchor, and the Warehouse / Overseas prefix
filters (TJ, TM, RO, RM against TJ, TM). Synthetic rows only, offline.
Run from tests/: python -B -m unittest test_tjx_workbook
"""
import io
import unittest

import openpyxl

from _support import load_app


def row(sku, brand_key, brand_full, ats, **extra):
    d = {'sku': sku, 'brand_abbr': brand_key, 'brand_full': brand_full, 'brand': brand_key,
         'color': 'White Solid', 'fit': 'Slim Fit', 'fabrication': 'Synthetic Fabric',
         'warehouse': 'TR', 'total_ats': ats, 'incoming': 0,
         'tjx_color_family': 'White', 'tjx_new_fabric': '', 'tjx_source': 'Warehouse'}
    d.update(extra)
    return d


WH = [row('TJAAXX001SLS', 'AAA', 'Brand A', 900), row('TJAAXX002SLS', 'AAA', 'Brand A', 500),
      row('TMAAXX003SLS', 'AAA', 'Brand A', 100),
      row('ROBBXX010SLS', 'BBB', 'Brand B', 800), row('RMBBXX011SLS', 'BBB', 'Brand B', 40),
      row('TJCCXX020SLS', 'CCC', 'Brand C', 36)]
OS = [row('TJAAXX001SLS', 'AAA', 'Brand A', 300, po_ref='ZZ26001', production='ZZ26001',
          arrival='Nov 3, 2026', tjx_source='Overseas', incoming=300),
      row('TJCCXX020SLS', 'CCC', 'Brand C', 72, po_ref='ZZ26002', production='ZZ26002',
          arrival='TBD', tjx_source='Overseas', incoming=72)]
WH_COLS = 11          # IMAGE .. Total ATS


def tabs():
    return [{'brand_name': 'Warehouse', 'tab_name': 'Warehouse', 'view_mode': 'ats', 'flow_mode': False,
             'keep_order': True, 'items': [dict(r) for r in WH]},
            {'brand_name': 'Overseas', 'tab_name': 'Overseas', 'view_mode': 'incoming', 'flow_mode': True,
             'keep_order': True, 'items': [dict(r) for r in OS]}]


def fake_image():
    from PIL import Image
    b = io.BytesIO()
    Image.new('RGB', (150, 150), (200, 30, 30)).save(b, 'PNG')
    return {'image_data': io.BytesIO(b.getvalue()), 'x_scale': 1, 'y_scale': 1,
            'x_offset': 0, 'y_offset': 0, 'url': 'https://photos.test/x.png'}


def merged_above(ws, last_row):
    return sorted(str(m) for m in ws.merged_cells.ranges if m.min_row <= last_row)


class Base(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.app = load_app()

    def setUp(self):
        a = self.app
        self._saved = {k: getattr(a, k) for k in ('download_images_for_items', '_landing_index')}
        self.images = {}
        a.download_images_for_items = lambda items, url, use_cache=True: dict(self.images)
        a._landing_index = lambda: None

    def tearDown(self):
        for k, v in self._saved.items():
            setattr(self.app, k, v)

    def render(self, tjx_layout=True, images=None, tab_list=None):
        self.images = images or {}
        data = self.app.build_multi_brand_excel(tab_list or tabs(), 'http://photos.test/', catalog_mode=True,
                                                view_mode='ats', flow_mode=False, prepack_defaults=None,
                                                tjx_layout=tjx_layout)
        return openpyxl.load_workbook(io.BytesIO(data))


class BrandBands(Base):
    def test_warehouse_tab_is_broken_out_by_brand(self):
        ws = self.render()['Warehouse']
        # header 1; band 2; Brand A rows 3-5; band 6; Brand B rows 7-8; band 9; Brand C row 10
        self.assertEqual(merged_above(ws, 10), ['A2:K2', 'A6:K6', 'A9:K9'])
        self.assertEqual([ws['A2'].value, ws['A6'].value, ws['A9'].value], ['Brand A', 'Brand B', 'Brand C'])
        self.assertEqual([ws.cell(r, 2).value for r in (3, 4, 5, 7, 8, 10)],
                         ['TJAAXX001SLS', 'TJAAXX002SLS', 'TMAAXX003SLS', 'ROBBXX010SLS', 'RMBBXX011SLS',
                          'TJCCXX020SLS'])
        self.assertEqual(ws.cell(10, WH_COLS).value, 36)                  # Total ATS stays the last column

    def test_band_has_borders_and_a_title_look(self):
        ws = self.render()['Warehouse']
        band = ws['A2']
        self.assertTrue(band.font.bold)
        self.assertEqual(band.border.top.style, 'medium')
        self.assertEqual(band.border.bottom.style, 'medium')
        self.assertEqual(band.fill.fgColor.rgb[-6:], 'CBD5E1')
        self.assertEqual(ws.row_dimensions[2].height, 24)
        self.assertEqual(ws['B3'].border.top.style, 'thin')             # data rows unchanged

    def test_zebra_restarts_under_every_band(self):
        ws = self.render()['Warehouse']
        fills = {r: ws.cell(r, 2).fill.fgColor.rgb[-6:] for r in (3, 4, 5, 7, 8, 10)}
        self.assertEqual(fills, {3: 'FFFFFF', 4: 'F0F4F8', 5: 'FFFFFF', 7: 'FFFFFF', 8: 'F0F4F8', 10: 'FFFFFF'})

    def test_overseas_tab_gets_bands_too(self):
        ws = self.render()['Overseas']
        self.assertEqual(merged_above(ws, 5), ['A2:N2', 'A4:N4'])
        self.assertEqual([ws['A2'].value, ws['B3'].value, ws['A4'].value, ws['B5'].value],
                         ['Brand A', 'TJAAXX001SLS', 'Brand C', 'TJCCXX020SLS'])

    def test_images_and_placeholders_stay_on_their_rows(self):
        ws = self.render(images={0: fake_image(), 3: fake_image()})['Warehouse']
        anchors = sorted(img.anchor._from.row for img in ws._images)    # 0-based sheet rows
        self.assertEqual(anchors, [2, 6])                                # items 0 and 3 -> rows 3 and 7
        self.assertEqual([ws.cell(r, 1).value for r in (4, 5, 8, 10)], ['No Image'] * 4)
        self.assertIsNone(ws['A3'].value)                                # image cell, no placeholder text
        self.assertEqual(ws['A2'].value, 'Brand A')                      # never a placeholder on a band

    def test_size_grids_start_under_the_last_row(self):
        ws = self.render()['Warehouse']
        self.assertTrue(all(c.value is None for c in ws[11]))           # one blank row
        self.assertGreaterEqual(ws.max_row, 12)
        self.assertTrue(any(c.value for c in ws[12]))                   # the size-scale grid

    def test_a_brand_split_by_the_caller_is_still_one_block(self):
        t = tabs()
        t[0]['items'] = [dict(WH[0]), dict(WH[3]), dict(WH[1]), dict(WH[5]), dict(WH[4]), dict(WH[2])]
        ws = self.render(tab_list=t)['Warehouse']
        self.assertEqual(merged_above(ws, 10), ['A2:K2', 'A6:K6', 'A9:K9'])
        self.assertEqual([ws.cell(r, 2).value for r in (3, 4, 5, 7, 8, 10)],
                         ['TJAAXX001SLS', 'TJAAXX002SLS', 'TMAAXX003SLS', 'ROBBXX010SLS', 'RMBBXX011SLS',
                          'TJCCXX020SLS'])

    def test_other_customer_sheets_are_untouched(self):
        ws = self.render(tjx_layout=False)['Warehouse']
        self.assertEqual(merged_above(ws, 7), [])
        self.assertEqual([ws.cell(r, 2).value for r in range(2, 8)],
                         ['TJAAXX001SLS', 'TJAAXX002SLS', 'TMAAXX003SLS', 'ROBBXX010SLS', 'RMBBXX011SLS',
                          'TJCCXX020SLS'])
        self.assertTrue(all(c.value is None for c in ws[8]))
        self.assertEqual(ws['B2'].fill.fgColor.rgb[-6:], 'FFFFFF')
        self.assertEqual(ws['B3'].fill.fgColor.rgb[-6:], 'F0F4F8')


class PrefixFilters(Base):
    def test_warehouse_takes_tjx_and_ross_overseas_tjx_only(self):
        a = self.app
        self.assertEqual(a._TJX_ATS_PREFIXES, ('TJ', 'TM'))
        self.assertEqual(a._TJX_ATS_WH_PREFIXES, ('TJ', 'TM', 'RO', 'RM'))
        wh, ov = a._tjx_ats_sku_filter(a._TJX_ATS_WH_PREFIXES), a._tjx_ats_sku_filter(a._TJX_ATS_PREFIXES)
        for sku in ('TJAAXX001SLS', 'TMAAXX003SLS', 'ROBBXX010SLS', 'RMBBXX011SLS', 'robbxx012sls'):
            self.assertTrue(wh(sku), sku)
        for sku in ('ROBBXX010SLS', 'RMBBXX011SLS', 'BUAAXX001SLS', ''):
            self.assertFalse(ov(sku), sku)
        self.assertTrue(ov('TJAAXX001SLS') and ov('TMAAXX003SLS'))
        self.assertFalse(wh('BUAAXX001SLS'))
        self.assertTrue(wh('ROBBXX010SLS-V'))                            # a variant is a style
        self.assertFalse(wh('ROBBXX010SLS-M'))                           # a by-size row is hidden
        self.assertFalse(wh('TJAAXX001SLS-1515.53233'))


if __name__ == '__main__':
    unittest.main()
