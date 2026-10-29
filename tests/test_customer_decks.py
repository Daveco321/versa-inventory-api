"""One customer's booked, on-order and shipped goods as decks and line sheets
(David, Oct 29 2026: "train the email bot to create line sheets and presentations
that showcase current allocations for a specific customer, by brand if requested").

Covers the APO customer-name match, substitutes held back, the booked-line supply
rule (warehouse first, then production by arrival, and the production 'twin' when the
allocated style number and the production style number differ), the server-only
_admin_ok flag that keeps shipped history admin only, open_orders_lookup's brand
filter, build_allocation_sheet, build_sales_sheet's Booked tab and brand filter, and
the renderer's per-section options.

All synthetic: customer names and codes start with Z, styles are Z-prefixed. No
network: the open-orders service, S3, photos, logos and annotations are fakes.
Run: python -B -m unittest discover -s tests -t ."""
import io
import json
import unittest

import openpyxl

from tests._support import AppHarness, FakeOpenOrders, load_app, synthetic_data

EM, EN = chr(0x2014), chr(0x2013)

APO = [
    # ZALPHA (open-orders code ZALPHA) books two VD styles and one Nautica style
    {'customer': 'ZALPHA', 'po': 'ZALPHA - VD SHAKETS SPRING ZP26001', 'qty': 600, 'style': 'ZZVDSG041SSS'},
    {'customer': 'Zalpha ', 'po': 'ZALPHA - VD SHAKETS SPRING ZP26001', 'qty': 400, 'style': 'ZZVDSN055SSP'},
    {'customer': 'ZALPHA', 'po': 'ZALPHA - NAUTICA WOVENS', 'qty': 300, 'style': 'ZZNAPO201SLS'},
    # a substitute: never shown
    {'customer': 'ZALPHA', 'po': 'ZALPHA SUB PO#77002001', 'qty': 120, 'style': 'ZZVDCO028SLS'},
    # two customers sharing a word, and a zero row
    {'customer': 'ZBETA STORES', 'po': 'ZBETA - BOOK', 'qty': 50, 'style': 'ZZNAPO201SLS'},
    {'customer': 'ZBETA OUTLET', 'po': 'ZBETA - BOOK', 'qty': 70, 'style': 'ZZNAPO201SLS'},
    {'customer': 'ZGAMMA', 'po': 'ZGAMMA - NOTHING', 'qty': 0, 'style': 'ZZNAPO201SLS'},
]
LEDGER = [
    # ZZVDSN055SSP is on its own production; ZZVDSG041SSS has none, but the
    # production its allocation names carries ZZVDSL041SSS (another fabric code)
    {'style': 'ZZVDSN055SSP', 'production': 'ZP26001', 'units': 400, 'etd': '2027-01-25', 'arrival': None,
     'warehouse': 'JTW', 'poName': 'ZALPHA VD'},
    {'style': 'ZZVDSL041SSS', 'production': 'ZP26001', 'units': 600, 'etd': '2027-01-25', 'arrival': '2027-03-01',
     'warehouse': 'JTW', 'poName': 'ZALPHA VD'},
]
INVENTORY = [
    # the Nautica style is in the warehouse; VD SN055 is incoming and allocated
    {'sku': 'ZZNAPO201SLS', 'brand_abbr': 'NAUTICA', 'jtw': 1000, 'tr': 0, 'dcw': 0, 'qa': 0, 'nj': 0, 'abfi': 0,
     'incoming': 0, 'committed': 0, 'allocated': -420},
    {'sku': 'ZZVDSN055SSP', 'brand_abbr': 'VD', 'jtw': 0, 'tr': 0, 'dcw': 0, 'qa': 0, 'nj': 0, 'abfi': 0,
     'incoming': 400, 'committed': 0, 'allocated': -400},
    {'sku': 'ZZVDSL041SSS', 'brand_abbr': 'VD', 'jtw': 0, 'tr': 0, 'dcw': 0, 'qa': 0, 'nj': 0, 'abfi': 0,
     'incoming': 600, 'committed': 0, 'allocated': 0},
]
VD_OPEN = {'customer': 'ZALPHA', 'customerFull': 'Alpha Test Stores', 'orderNo': '77002002',
           'style': 'ZZVDSL034SLC', 'openQty': 900, 'pickQty': 0, 'openValue': 18000.0, 'pickValue': 0.0,
           'startDate': '2026-11-24T00:00:00', 'cancelDate': '2026-11-29T00:00:00', 'isPipeline': False}


class Base(unittest.TestCase):
    mode = 'new'

    def setUp(self):
        self.app = a = load_app()
        self.data = synthetic_data()
        self.data['open_orders'].append(dict(VD_OPEN))
        self.fake = FakeOpenOrders(self.data, mode=self.mode)
        self.h = AppHarness(self.fake)
        self.h.__enter__()
        self.addCleanup(self.h.__exit__, None, None, None)
        self.rendered = []
        real_render = a._pres_render_pdf

        def render(groups, headline, date_label, density, orders_mode, **kw):
            self.rendered.append({'groups': groups, 'headline': headline, 'cover': kw.get('cover')})
            return real_render(groups, headline, date_label, density, orders_mode, **kw)
        for name, value in (('_apo_data', [dict(r) for r in APO]),
                            ('_ledger_rows', lambda: [dict(r) for r in LEDGER]),
                            ('load_apo_from_dropbox', lambda: False),
                            ('_fetch_sub_annotations', lambda: ({}, True)),
                            ('_pres_photo_jpeg', lambda sku, brand: None),
                            ('_pres_logo_png', lambda key: None),
                            ('_apo_color_map', lambda: {}),
                            ('_style_overrides', {}),
                            ('_pres_render_pdf', render)):
            saved = getattr(a, name)
            setattr(a, name, value)
            self.addCleanup(setattr, a, name, saved)
        saved_items = a._inventory.get('items')
        a._inventory['items'] = [dict(r) for r in INVENTORY]
        self.addCleanup(a._inventory.__setitem__, 'items', saved_items)

    def deck(self, **params):
        return self.app._ai_tool_build_presentation(params)


# ── which customer ─────────────────────────────────────────────────────────

class CustomerMatch(Base):
    def test_case_and_spacing_variants_are_one_customer(self):
        names, err = self.app._apo_customer_resolve('zalpha')
        self.assertIsNone(err)
        self.assertEqual({'ZALPHA'}, names)

    def test_a_shared_word_asks_which_customer(self):
        names, err = self.app._apo_customer_resolve('ZBETA')
        self.assertEqual(set(), names)
        self.assertEqual(['ZBETA OUTLET', 'ZBETA STORES'], err['customers_matched'])
        names, err = self.app._apo_customer_resolve('ZBETA', combine=True)
        self.assertEqual({'ZBETA OUTLET', 'ZBETA STORES'}, names)

    def test_an_unknown_customer_lists_who_has_allocations_and_never_swaps(self):
        _names, err = self.app._apo_customer_resolve('Zomega')
        self.assertIn("no open allocations for 'Zomega'", err['error'])
        self.assertIn('ZALPHA', err['allocation_customers'])
        self.assertNotIn('ZGAMMA', err['allocation_customers'], 'a customer with only zero rows has nothing open')
        self.assertIn('never substitute', err['hint'])

    def test_known_aliases(self):
        self.app._apo_data.append({'customer': 'MACY BACKSTAGE', 'po': 'X', 'qty': 5, 'style': 'ZZNAPO201SLS'})
        self.app._apo_data.append({'customer': 'MACY', 'po': 'Y', 'qty': 5, 'style': 'ZZNAPO201SLS'})
        self.assertEqual({'MACY BACKSTAGE'}, self.app._apo_customer_resolve("Macy's Backstage")[0])
        self.assertEqual({'MACY BACKSTAGE'}, self.app._apo_customer_resolve('MACY1')[0])
        self.assertEqual({'MACY'}, self.app._apo_customer_resolve("Macy's")[0])


# ── booked lines ───────────────────────────────────────────────────────────

class BookedLines(Base):
    def rows(self, **spec):
        res = self.app._customer_alloc_rows('Zalpha', self.app._pres_filter_spec(spec))
        self.assertNotIn('error', res)
        return res, {r['style']: r for r in res['rows']}

    def test_substitutes_are_held_back_and_counted_per_brand(self):
        res, rows = self.rows()
        self.assertNotIn('ZZVDCO028SLS', rows)
        self.assertEqual((1, 120), (res['held']['lines'], res['held']['units']))
        res, _ = self.rows(brands=['NAUTICA'])
        self.assertEqual(0, res['held']['lines'], 'a Nautica deck holds back no VD substitute')

    def test_brand_filter(self):
        _res, rows = self.rows(brands=['Von Dutch'])
        self.assertEqual({'ZZVDSG041SSS', 'ZZVDSN055SSP'}, set(rows))

    def test_supply_warehouse_first_then_production_by_arrival(self):
        _res, rows = self.rows()
        self.assertEqual('In warehouse', rows['ZZNAPO201SLS']['supply'])
        self.assertEqual('JTW', rows['ZZNAPO201SLS']['warehouse'])
        # ETD Jan 25 + the transit rule; never TBD for a dated lot
        self.assertRegex(rows['ZZVDSN055SSP']['supply'], r'^(Mar|Apr) \d+, 2027$')
        self.assertEqual('ZP26001', rows['ZZVDSN055SSP']['production'])

    def test_production_twin_gives_the_arrival_and_is_reported(self):
        res, rows = self.rows()
        r = rows['ZZVDSG041SSS']
        self.assertEqual('Mar 1, 2027', r['supply'], 'the twin lot arrival, with no marker in the file')
        self.assertEqual('ZZVDSL041SSS', r['twin'])
        m = res['mismatches']
        self.assertEqual(1, len(m))
        self.assertEqual(('ZZVDSG041SSS', 'ZZVDSL041SSS', 'ZP26001', 600),
                         (m[0]['allocated_style'], m[0]['production_style'], m[0]['production'], m[0]['units']))

    def test_no_twin_without_a_named_production(self):
        for a in self.app._apo_data:
            if a['style'] == 'ZZVDSG041SSS':
                a['po'] = 'ZALPHA - VD SHAKETS SPRING'
        res, rows = self.rows()
        self.assertEqual('Not on production', rows['ZZVDSG041SSS']['supply'])
        self.assertEqual([], res['mismatches'])

    def test_allocation_names_lose_the_customer_prefix(self):
        _res, rows = self.rows()
        self.assertEqual('VD SHAKETS SPRING ZP26001', rows['ZZVDSG041SSS']['short'])
        self.assertEqual('NAUTICA WOVENS', rows['ZZNAPO201SLS']['short'])


# ── decks ──────────────────────────────────────────────────────────────────

class Decks(Base):
    def test_allocations_deck(self):
        out = self.deck(source='allocations', customer='Zalpha', brands=['VD'], _admin_ok=False)
        self.assertNotIn('error', out)
        self.assertEqual({'Booked': {'styles': 2, 'units': 1000, 'allocations': 1}}, out['by_section'])
        self.assertTrue(out['download_url'].endswith('.pdf'))
        self.assertIn('Zalpha Von Dutch Booked Presentation', self.h.s3.puts[-1]['Key'])
        g = self.rendered[-1]['groups']
        self.assertEqual(['Von Dutch'], [x['name'] for x in g])
        o = g[0]['opts']
        self.assertEqual(('Allocation', 'Arrives', 'Total booked', False),
                         (o['po_hdr'], o['date_hdr'], o['total_label'], o['show_cost']))
        self.assertIsNone(self.rendered[-1]['cover'])
        self.assertEqual(120, out['held_back_substitutes']['units'])
        self.assertEqual(1, len(out['style_mismatches']))
        self.assertTrue(self.h.s3.puts[-1]['Body'][:4] == b'%PDF')

    def test_shipped_is_admin_only(self):
        out = self.deck(source='shipped', customer='ZALPHA')
        self.assertTrue(out.get('admin_only'), 'no _admin_ok = not an admin')
        out = self.deck(source='shipped', customer='ZALPHA', _admin_ok=False)
        self.assertIn('admins only', out['error'])
        self.assertEqual([], self.h.s3.puts)

    def test_shipped_deck_for_an_admin(self):
        out = self.deck(source='shipped', customer='ZALPHA', _admin_ok=True)
        self.assertNotIn('error', out)
        self.assertIn('history_label', out)
        cards = [cd for g in self.rendered[-1]['groups'] for cd in g['cards']]
        whens = [ln['when'] for cd in cards for ln in cd['lines']]
        self.assertTrue(any(w.endswith(' est.') for w in whens), 'shipments after the cut-off are estimates')
        self.assertTrue(any(not w.endswith(' est.') for w in whens))
        for cd in cards:
            ds = [ln['date'] for ln in cd['lines']]
            self.assertEqual(sorted(ds, reverse=True), ds, 'newest first')
        o = self.rendered[-1]['groups'][0]['opts']
        self.assertEqual(('Shipped', 'Total shipped'), (o['date_hdr'], o['total_label']))
        self.assertIn('Invoices through', o['footer_note'])

    def test_account_deck_for_an_admin_has_every_section_and_a_summary_page(self):
        out = self.deck(source='account', customer='ZALPHA', _admin_ok=True)
        self.assertNotIn('error', out)
        self.assertEqual(['allocations', 'open_orders', 'shipped'], out['sections'])
        self.assertIsNotNone(self.rendered[-1]['cover'])
        names = [g['name'] for g in self.rendered[-1]['groups']]
        self.assertTrue(names[0].endswith('· Booked'))
        self.assertTrue(any(n.endswith('· On Order') for n in names))
        self.assertTrue(names[-1].endswith('· Shipped'))
        self.assertEqual(out['pages'], 1 + sum(g['pages'] for g in out['groups']))

    def test_account_deck_for_staff_leaves_shipped_out_and_says_so(self):
        out = self.deck(source='account', customer='ZALPHA', _admin_ok=False)
        self.assertNotIn('error', out)
        self.assertNotIn('shipped', out['sections'])
        self.assertTrue(any('admin only' in n for n in out['notes']))
        self.assertNotIn('history_label', out)

    def test_account_brand_filter_reaches_every_section(self):
        out = self.deck(source='account', customer='ZALPHA', brands=['VD'], _admin_ok=True)
        self.assertNotIn('error', out)
        skus = {cd['sku'] for g in self.rendered[-1]['groups'] for cd in g['cards']}
        self.assertEqual({'ZZVDSG041SSS', 'ZZVDSN055SSP', 'ZZVDSL034SLC'}, skus)
        self.assertIn('Von Dutch', self.rendered[-1]['headline'])

    def test_sections_param(self):
        out = self.deck(source='account', customer='ZALPHA', sections=['on order', 'booked'], _admin_ok=True)
        self.assertEqual(['allocations', 'open_orders'], out['sections'], 'the deck keeps its own order')

    def test_unknown_source_names_the_new_ones(self):
        out = self.deck(source='nonsense', customer='ZALPHA')
        for s in ('allocations', 'shipped', 'account'):
            self.assertIn(s, out['error'])


# ── the server, not the model, says who is an admin ─────────────────────────

class AdminFlag(unittest.TestCase):
    def setUp(self):
        self.app = load_app()

    def test_model_input_can_never_carry_the_flag(self):
        f = self.app._ai_agent_tool_input
        self.assertEqual({'source': 'shipped', '_admin_ok': False},
                         f('build_presentation', {'source': 'shipped', '_admin_ok': True}, False))
        self.assertEqual({'source': 'shipped', '_admin_ok': True},
                         f('build_presentation', {'source': 'shipped'}, True))
        self.assertEqual({'brand': 'VD'}, f('query_inventory', {'brand': 'VD', '_admin_ok': True}, True))

    def test_agent_loop_passes_the_servers_flag(self):
        seen = []

        class Block:
            def __init__(self, **kw):
                self.__dict__.update(kw)

        class Resp:
            def __init__(self, content, stop='tool_use'):
                self.content, self.stop_reason, self.usage = content, stop, None

        replies = [Resp([Block(type='tool_use', id='t1', name='build_presentation',
                               input={'source': 'shipped', 'customer': 'ZALPHA', '_admin_ok': True})]),
                   Resp([Block(type='text', text='done')], stop='end_turn')]

        class Client:
            def with_options(self, **kw):
                return self

            class messages:
                @staticmethod
                def create(**kw):
                    return replies.pop(0)
        self.app._ai_agent_run(Client(), [{'role': 'user', 'content': 'x'}], [], [], 'm', 100, False,
                               tool_fns={'build_presentation': lambda p: seen.append(p) or {'ok': 1}})
        self.assertEqual([{'source': 'shipped', 'customer': 'ZALPHA', '_admin_ok': False}], seen)

    def test_mcp_route_is_the_full_table(self):
        seen = []
        saved = self.app._AI_AGENT_TOOL_FNS['build_presentation']
        self.app._AI_AGENT_TOOL_FNS['build_presentation'] = lambda p: seen.append(p) or {'ok': 1}
        try:
            with self.app.app.test_request_context():
                self.app._mcp_handle({'jsonrpc': '2.0', 'id': 1, 'method': 'tools/call',
                                      'params': {'name': 'build_presentation',
                                                 'arguments': {'source': 'account', '_admin_ok': False}}})
        finally:
            self.app._AI_AGENT_TOOL_FNS['build_presentation'] = saved
        self.assertEqual([{'source': 'account', '_admin_ok': True}], seen)


# ── open_orders_lookup by brand ────────────────────────────────────────────

class OpenOrdersBrand(Base):
    def test_brand_filter_and_brand_on_rows(self):
        out = self.app._ai_tool_open_orders({'customer': 'Alpha', 'brands': ['Von Dutch']})
        self.assertEqual(1, out['matched_orders'])
        self.assertEqual('ZZVDSL034SLC', out['rows'][0]['style'])
        self.assertEqual('Von Dutch', out['rows'][0]['brand'])

    def test_a_style_prefix_matches_every_style_that_starts_with_it(self):
        out = self.app._ai_tool_open_orders({'style': 'ZZVD'})
        self.assertEqual(['ZZVDSL034SLC'], [r['style'] for r in out['rows']])
        out = self.app._ai_tool_open_orders({'style': 'ZZVDSL034SLC'})
        self.assertEqual(1, out['matched_orders'])

    def test_several_brands_are_rolled_up(self):
        out = self.app._ai_tool_open_orders({'customer': 'Alpha'})
        self.assertIn('Von Dutch', out['units_by_brand'])


# ── line sheets ────────────────────────────────────────────────────────────

class AllocationSheet(Base):
    def test_a_tab_per_brand_with_the_allocation_columns(self):
        out = self.app._ai_tool_build_allocation_sheet({'customer': 'Zalpha'})
        self.assertNotIn('error', out)
        self.assertEqual(['Von Dutch', 'Nautica'], [t['tab'] for t in out['tabs']], 'biggest first')
        wb = openpyxl.load_workbook(io.BytesIO(self.h.s3.puts[-1]['Body']))
        self.assertEqual(['Von Dutch', 'Nautica'], wb.sheetnames)
        ws = wb['Von Dutch']
        head = [c.value for c in ws[1]]
        for col in ('SKU', 'Allocation', 'Units Allocated', 'Production #', 'Arrival', 'Warehouse'):
            self.assertIn(col, head)
        cells = [str(c.value) for row in wb.worksheets for r in row.iter_rows() for c in r if c.value is not None]
        self.assertFalse(any('ZZVDCO028SLS' in c for c in cells), 'no substitute in the file')
        self.assertFalse(any('ZZVDSL041SSS' in c for c in cells), 'the twin never shows in the file')
        self.assertFalse(any('SUB' in c.split() for c in cells))
        self.assertEqual(120, out['held_back_substitutes']['units'])
        self.assertEqual('ZZVDSL041SSS', out['style_mismatches'][0]['production_style'])

    def test_brand_filter(self):
        out = self.app._ai_tool_build_allocation_sheet({'customer': 'Zalpha', 'brands': ['NAUTICA']})
        self.assertEqual(['Nautica'], [t['tab'] for t in out['tabs']])
        self.assertEqual(300, out['units'])

    def test_unknown_customer(self):
        out = self.app._ai_tool_build_allocation_sheet({'customer': 'Zomega'})
        self.assertIn('allocation_customers', out)


class SalesSheetBooked(Base):
    def build(self, **params):
        out = self.app._ai_tool_build_sales_sheet(params)
        if 'error' in out:
            return out, None
        return out, openpyxl.load_workbook(io.BytesIO(self.h.s3.puts[-1]['Body']))

    def test_booked_tab_on_request_only(self):
        out, wb = self.build(customer='ZALPHA')
        self.assertNotIn('Booked', wb.sheetnames, 'the default tabs are unchanged')
        out, wb = self.build(customer='ZALPHA', include=['allocated', 'on_order'])
        self.assertIn('Booked', wb.sheetnames)
        ws = wb['Booked']
        self.assertEqual(['Image', 'Style', 'Color', 'Allocation', 'Qty', 'Supply'], [c.value for c in ws[1]])
        supply = {r[1].value: r[5].value for r in ws.iter_rows(min_row=2)}
        self.assertEqual('In warehouse', supply['ZZNAPO201SLS'])
        self.assertTrue(supply['ZZVDSN055SSP'].startswith('Arrives '))
        summary = [str(c.value) for r in wb['Summary'].iter_rows() for c in r if c.value is not None]
        self.assertFalse(any('substitut' in s.lower() for s in summary), 'the file never mentions substitutes')
        self.assertFalse(any('ZZVDSL041SSS' in s for s in summary))
        self.assertEqual(120, out['held_back_substitutes']['units'])

    def test_brands_narrow_every_tab(self):
        out, wb = self.build(customer='ZALPHA', brands=['VD'], include=['allocated', 'on_order', 'invoiced'])
        self.assertEqual(['VD'], out['brands'])
        self.assertEqual({'Summary', 'On Order', 'Booked'}, set(wb.sheetnames),
                         'the synthetic invoices carry no VD style')
        styles = {r[1].value for r in wb['On Order'].iter_rows(min_row=2)}
        self.assertEqual({'ZZVDSL034SLC'}, styles)

    def test_brand_with_nothing_says_so(self):
        out, _wb = self.build(customer='ZALPHA', brands=['CHAPS'])
        self.assertIn('for those brands', out['error'])


# ── renderer ───────────────────────────────────────────────────────────────

class RendererOptions(unittest.TestCase):
    def setUp(self):
        self.app = load_app()
        try:
            import pymupdf
            self.pdf = pymupdf
        except ImportError:
            self.skipTest('pymupdf not installed')

    def text(self, groups, **kw):
        pdf, n = self.app._pres_render_pdf(groups, 'Versa Group · Z', 'Oct 29, 2026', 8, True, **kw)
        doc = self.pdf.open(stream=pdf, filetype='pdf')
        return n, [p.get_text() for p in doc]

    def card(self, lines):
        return {'sku': 'ZZNAPO201SLS', 'base': 'ZZNAPO201SLS', 'brand': 'NAUTICA', 'color': 'Navy', 'fab': 'Poplin',
                'fit': 'Slim Fit', 'number': sum(l['units'] for l in lines), 'lines': lines}

    def test_defaults_are_unchanged(self):
        g = [{'key': 'NAUTICA', 'name': 'Nautica', 'logo': None,
              'cards': [self.card([{'po': '111', 'units': 5, 'cost': 2.0, 'start': '2026-11-01', 'cancel': '2026-11-05'},
                                   {'po': '222', 'units': 7, 'cost': 2.0, 'start': '2026-12-01', 'cancel': '2026-12-05'}])]}]
        n, pages = self.text(g)
        self.assertEqual(1, n)
        for s in ('PO #', 'Ship window', 'Cost', 'Total on order: 12 units', '2 POs', 'Page 1 of 1', 'Nov 1-5, 2026'):
            self.assertIn(s, pages[0])

    def test_section_options_and_a_cover_page(self):
        drawn = []
        g = [{'key': 'NAUTICA', 'name': 'Nautica · Booked', 'logo': None,
              'opts': {'po_hdr': 'Allocation', 'date_hdr': 'Arrives', 'total_label': 'Total booked',
                       'show_cost': False, 'count_word': 'allocation', 'footer_note': 'Booked = no PO yet.'},
              'cards': [self.card([{'po': 'BOOK A', 'units': 5, 'cost': 0, 'when': 'Mar 1, 2027'},
                                   {'po': 'BOOK B', 'units': 7, 'cost': 0, 'when': 'In warehouse'}])]}]
        n, pages = self.text(g, cover=lambda c, w, h, total: drawn.append(total) or c.drawString(40, 40, 'COVER'))
        self.assertEqual([2], drawn)
        self.assertEqual(2, n)
        self.assertIn('COVER', pages[0])
        p = pages[1]
        for s in ('Allocation', 'Arrives', 'Total booked: 12 units', '2 allocations', 'Mar 1, 2027',
                  'In warehouse', 'Booked = no PO yet.', 'Page 2 of 2'):
            self.assertIn(s, p)
        self.assertNotIn('Cost', p)
        self.assertNotIn('Ship window', p)


# ── tool text ──────────────────────────────────────────────────────────────

class ToolText(unittest.TestCase):
    def setUp(self):
        self.app = load_app()

    def test_schemas_and_guidance(self):
        tools = {t['name']: t for t in self.app._AI_AGENT_TOOLS}
        self.assertIn('build_allocation_sheet', tools)
        self.assertIn('build_allocation_sheet', self.app._AI_AGENT_TOOL_FNS)
        self.assertNotIn('build_allocation_sheet', self.app._AI_AGENT_ADMIN_TOOLS, 'booked goods are staff data')
        src = tools['build_presentation']['input_schema']['properties']['source']['enum']
        for s in ('allocations', 'shipped', 'account'):
            self.assertIn(s, src)
        self.assertIn('brands', tools['open_orders_lookup']['input_schema']['properties'])
        self.assertIn('allocated', tools['build_sales_sheet']['input_schema']['properties']['include']['items']['enum'])
        core = self.app._AI_AGENT_TOOL_GUIDANCE_CORE
        self.assertIn("source:'account', customer:'Ross', brands:['VD']", core)
        self.assertIn('Never mention substitute allocations', core)
        for t in ('build_presentation', 'build_allocation_sheet', 'open_orders_lookup', 'build_sales_sheet'):
            text = json.dumps(tools[t], ensure_ascii=False)
            self.assertNotIn(EM, text, t)
            self.assertNotIn(EN, text, t)
        para = next(ln for ln in core.splitlines() if ln.startswith("ONE CUSTOMER'S BOOKED"))
        self.assertNotIn(EM, para)
        self.assertNotIn(EN, para)


if __name__ == '__main__':
    unittest.main()
