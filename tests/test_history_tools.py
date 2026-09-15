"""Past Orders AI tools against stubbed old and new open-orders answers
(INV-02 to INV-08). Run: python -B -m unittest discover -s tests -t .

All data is synthetic (customer codes start with Z). No network: the
open-orders service, S3, photos and the open-order feed are all fakes."""
import inspect
import io
import json
import unittest

import openpyxl

from tests._support import AppHarness, FakeOpenOrders, load_app, synthetic_data

EM, EN = chr(0x2014), chr(0x2013)
PO_HISTORY = ('GET', '/api/po-history')
SALES = ('GET', '/api/sales-history')
LOOKUP = ('POST', '/api/sales-history/po-lookup')


def _strings(x):
    if isinstance(x, str):
        yield x
    elif isinstance(x, dict):
        for k, v in x.items():
            yield str(k)
            yield from _strings(v)
    elif isinstance(x, (list, tuple)):
        for v in x:
            yield from _strings(v)


class Base(unittest.TestCase):
    mode = 'old'

    def setUp(self):
        self.app = load_app()
        self.data = synthetic_data()
        self.fake = FakeOpenOrders(self.data, mode=self.mode)
        self.h = AppHarness(self.fake)
        self.h.__enter__()

    def tearDown(self):
        self.h.__exit__(None, None, None)

    def assertPlain(self, out):
        text = ' '.join(_strings(out))
        low = text.lower()
        self.assertNotIn('cancelled', low)
        self.assertNotIn('pulled', low)
        self.assertNotIn(EM, text)
        self.assertNotIn(EN, text)

    def assertMeta(self, out):
        self.assertIn('invoices_through', out)
        self.assertIn('history_label', out)
        self.assertTrue(out['history_label'])

    def rows_by_po(self, out):
        return {r['order_no']: r for r in out['rows']}


# ── INV-02: the proxy cache never keeps a bad moment ─────────────────────

class CacheRules(Base):
    mode = 'new'

    def expire(self):
        c = self.app._po_hist_proxy_cache
        c['summary_at'] = 0.0
        for k in list(c['accounts_at']):
            c['accounts_at'][k] = 0.0

    def test_ready_answer_is_cached(self):
        self.app._fetch_po_history()
        self.app._fetch_po_history()
        self.assertEqual(self.fake.count(*PO_HISTORY), 1)

    def test_not_ready_is_not_cached(self):
        self.fake.po_history_override[None] = {'ready': False, 'building': True,
                                               'progress': {'done': 3, 'total': 105}}
        d1, ok1 = self.app._fetch_po_history()
        d2, ok2 = self.app._fetch_po_history()
        self.assertFalse(ok1)
        self.assertEqual(self.fake.count(*PO_HISTORY), 2)
        out = self.app._ai_tool_past_orders({})
        self.assertTrue(out.get('building'))
        self.assertPlain(out)

    def test_view_not_ready_is_not_cached(self):
        s = self.fake.po_history()
        s['history']['viewReady'] = False
        self.fake.po_history_override[None] = s
        self.app._fetch_po_history()
        self.app._fetch_po_history()
        self.assertEqual(self.fake.count(*PO_HISTORY), 2)

    def test_degraded_is_not_cached(self):
        s = self.fake.po_history()
        s['history']['flags']['degraded'] = True
        self.fake.po_history_override[None] = s
        self.app._fetch_po_history()
        self.app._fetch_po_history()
        self.assertEqual(self.fake.count(*PO_HISTORY), 2)
        out = self.app._ai_tool_past_orders({})
        self.assertIn('degraded_note', out)

    def test_empty_list_after_full_keeps_last_good(self):
        good, _ = self.app._fetch_po_history()
        self.expire()
        empty = self.fake.po_history()
        empty['accounts'] = []
        self.fake.po_history_override[None] = empty
        got, ok = self.app._fetch_po_history()
        self.assertFalse(ok)
        self.assertTrue(got['accounts'])
        self.assertIs(self.app._po_hist_proxy_cache['summary'], good)
        # the same rule for one account's PO list
        acct, _ = self.app._fetch_po_history('ZALPHA')
        self.expire()
        e2 = self.fake.po_history('ZALPHA')
        e2['pos'] = []
        self.fake.po_history_override['ZALPHA'] = e2
        got2, ok2 = self.app._fetch_po_history('ZALPHA')
        self.assertFalse(ok2)
        self.assertEqual(len(got2['pos']), len(acct['pos']))

    def test_building_serves_last_good(self):
        good, _ = self.app._fetch_po_history()
        self.expire()
        self.fake.po_history_override[None] = {'ready': False, 'building': True,
                                               'progress': {'done': 3, 'total': 105}}
        got, ok = self.app._fetch_po_history()
        self.assertFalse(ok)
        self.assertIs(got, good)
        out = self.app._ai_tool_past_orders({})
        self.assertNotIn('building', out)
        self.assertIn('stale_note', out)
        self.assertTrue(out['customers'])

    def test_empty_first_answer_is_not_cached(self):
        empty = self.fake.po_history()
        empty['accounts'] = []
        self.fake.po_history_override[None] = empty
        self.app._fetch_po_history()
        self.app._fetch_po_history()
        self.assertEqual(self.fake.count(*PO_HISTORY), 2)
        out = self.app._ai_tool_past_orders({})
        self.assertIn('no customers', out['error'])
        self.assertMeta(out)
        self.assertPlain(out)

    def test_error_serves_last_good_and_backs_off(self):
        self.app._fetch_po_history()
        self.expire()
        self.fake.fail[PO_HISTORY] = 500
        got, ok = self.app._fetch_po_history()
        self.assertFalse(ok)
        self.assertTrue(got['accounts'])
        self.assertIn('HTTP 500', self.app._po_hist_error())
        n = self.fake.count(*PO_HISTORY)
        self.app._fetch_po_history()      # inside the back-off: no new call
        self.assertEqual(self.fake.count(*PO_HISTORY), n)
        out = self.app._ai_tool_past_orders({})
        self.assertIn('stale_note', out)

    def test_error_with_no_copy_is_an_error(self):
        self.fake.fail[PO_HISTORY] = 401
        got, ok = self.app._fetch_po_history()
        self.assertIsNone(got)
        out = self.app._ai_tool_past_orders({})
        self.assertIn('HTTP 401', out['error'])
        self.assertPlain(out)

    def test_new_view_drops_account_copies(self):
        self.app._fetch_po_history()
        self.app._fetch_po_history('ZALPHA')
        self.assertIn('ZALPHA', self.app._po_hist_proxy_cache['accounts'])
        self.expire()
        s = self.fake.po_history()
        s['history']['viewKey'] = 'view-2'
        self.fake.po_history_override[None] = s
        self.app._fetch_po_history()
        self.assertNotIn('ZALPHA', self.app._po_hist_proxy_cache['accounts'])


# ── INV-03: one row, both server versions ────────────────────────────────

class RowMapping(Base):
    def test_old_rows(self):
        rows = {p['orderNo']: p for p in self.fake.rows('ZALPHA')}
        st = self.app._past_po_state(rows['77001004'], inv_checked=True)
        self.assertEqual(st['status'], 'shipped_awaiting_invoice')
        self.assertTrue(st['counted'])
        self.assertEqual((st['units'], st['units_basis']), (500, 'peak booked (estimate)'))
        st = self.app._past_po_state(rows['77001005'], inv_checked=True)
        self.assertEqual(st['status'], 'left_before_window_no_invoice')
        self.assertFalse(st['counted'])
        inv = {'qty': 400, 'value': 800.0, 'firstInv': '2026-08-10', 'lastInv': '2026-08-10',
               'invoices': 1, 'lines': [{'style': 'ZZTEST003SLS', 'qty': 400, 'value': 800.0}], 'pos': ['77001003']}
        st = self.app._past_po_state(rows['77001003'], inv=inv, inv_checked=True)
        self.assertEqual((st['status'], st['bucket'], st['units'], st['units_basis']),
                         ('invoiced', 'invoiced', 400, 'invoice'))
        row = self.app._past_po_row(rows['77001003'], state=st)
        self.assertEqual(row['invoice']['qty'], 400)
        self.assertEqual(row['lines_basis'], 'invoice')
        self.assertEqual(row['lines'][0]['price'], 2.0)
        for p in rows.values():
            self.assertPlain(self.app._past_po_row(p))

    def test_new_rows(self):
        self.fake.mode = 'new'
        rows = {p['orderNo']: p for p in self.fake.rows('ZALPHA')}
        st = self.app._past_po_state(rows['77001004'], final=True, estimate_rule='pick')
        self.assertEqual((st['bucket'], st['units'], st['units_basis']), ('awaiting_invoice', 480, 'estimate'))
        st = self.app._past_po_state(rows['77001004'], final=True, estimate_rule='peak')
        self.assertEqual(st['units_basis'], 'peak booked (estimate)')
        row = self.app._past_po_row(rows['66-000001'], state=self.app._past_po_state(rows['66-000001'], final=True))
        self.assertEqual(row['status'], 'invoiced')
        self.assertEqual(row['units'], 600)
        self.assertEqual(row['after_cutoff_estimate'], {'units': 100, 'value': 250.0, 'ship_date_est': '2026-09-03'})
        row = self.app._past_po_row(rows['77001008'], state=self.app._past_po_state(rows['77001008'], final=True))
        self.assertEqual(row['moved_to'], ['77001009'])
        self.assertFalse(row['counted'])
        self.assertEqual(row['status_label'], 'Moved to PO 77001009')
        row = self.app._past_po_row(rows['77001004'], state=self.app._past_po_state(rows['77001004'], final=True))
        self.assertEqual(row['ship_date_est'], '2026-09-01')
        self.assertEqual(row['left_book_range'], ['2026-09-03', '2026-09-03'])
        for p in rows.values():
            self.assertPlain(self.app._past_po_row(p, state=self.app._past_po_state(p, final=True)))


# ── INV-04: past_orders_lookup ───────────────────────────────────────────

class PastOrdersOld(Base):
    def test_summary(self):
        out = self.app._ai_tool_past_orders({})
        self.assertMeta(out)
        self.assertEqual(out['invoices_through'], {'wholesale': '2026-08-21', 'dropship': None})
        self.assertEqual(out['history_label'], 'Invoices through Aug 21, 2026. Dropship end date not reported.')
        # Weekdays Aug 24 to Sep 14: today (Sep 15) is not counted, like the desktop page.
        self.assertEqual(out['invoices_behind'], 'Invoices are 16 business days behind.')
        self.assertIn('dropship_note', out)
        self.assertEqual(out['totals']['left_before_window_pos'], 2)
        self.assertNotIn('cancelled', json.dumps(out['totals']))
        self.assertTrue(all('units_basis' in c for c in out['customers']))
        self.assertPlain(out)

    def test_customer_counts_only_what_shipped(self):
        out = self.app._ai_tool_past_orders({'customer': 'ZALPHA'})
        self.assertMeta(out)
        # One key-aware lookup, scoped to the account, both spellings of the padded PO.
        self.assertEqual(self.fake.count(*LOOKUP), 1)
        body = [c for c in self.fake.calls if c[:2] == LOOKUP][0][3]
        self.assertEqual(body['account'], 'ZALPHA')
        self.assertIn('0077001001', body['pos'])
        self.assertIn('77001001', body['pos'])
        rows = self.rows_by_po(out)
        self.assertEqual(rows['0077001001']['status'], 'invoiced')          # padded PO, counted once
        self.assertEqual(rows['0077001001']['units'], 1000)
        self.assertEqual(rows['77001003']['status'], 'invoiced')            # early shipment, invoiced
        self.assertEqual(rows['77001005']['status'], 'left_before_window_no_invoice')
        self.assertFalse(rows['77001005']['counted'])
        self.assertEqual(rows['77001004']['status'], 'shipped_awaiting_invoice')
        bd = out['by_status']
        self.assertEqual((bd['invoiced']['pos'], bd['invoiced']['units'], bd['invoiced']['value']), (4, 2600, 6800.0))
        self.assertEqual((bd['awaiting_invoice']['units'], bd['awaiting_invoice']['value']), (1170, 3460.0))
        self.assertEqual((bd['not_counted']['pos'], bd['not_counted']['value']), (1, 400.0))
        self.assertEqual((out['total_units'], out['total_value']), (3770, 10260.0))
        self.assertEqual(out['awaiting_invoice']['units'], 1170)
        life = out['invoiced_lifetime']
        self.assertEqual((life['units'], life['value']), (4840, 11280.0))
        self.assertIn('renumbered', out['server_note'])
        self.assertPlain(out)

    def test_status_filters_and_old_values(self):
        out = self.app._ai_tool_past_orders({'customer': 'ZALPHA', 'status': 'cancelled'})
        self.assertEqual(sorted(self.rows_by_po(out)), ['77001005'])
        self.assertEqual(out['total_units'], 0)
        out = self.app._ai_tool_past_orders({'customer': 'ZALPHA', 'status': 'shipped'})
        self.assertEqual(out['matched_pos'], 8)
        out = self.app._ai_tool_past_orders({'customer': 'ZALPHA', 'status': 'invoiced'})
        self.assertEqual((out['matched_pos'], out['total_units']), (4, 2600))
        out = self.app._ai_tool_past_orders({'customer': 'ZALPHA', 'status': 'awaiting_invoice'})
        self.assertEqual(out['total_units'], 1170)
        out = self.app._ai_tool_past_orders({'customer': 'ZALPHA', 'status': 'bogus'})
        self.assertIn('unknown status', out['error'])
        self.assertMeta(out)

    def test_lookup_failure_is_said_not_hidden(self):
        self.fake.fail[LOOKUP] = 500
        out = self.app._ai_tool_past_orders({'customer': 'ZALPHA'})
        self.assertIsNone(out['awaiting_invoice'])
        self.assertIn('HTTP 500', out['invoice_check_note'])
        self.assertIn('Not checked against invoices', self.rows_by_po(out)['77001004']['status_label'])
        self.assertPlain(out)

    def test_style_search(self):
        out = self.app._ai_tool_past_orders({'style': 'ZZTEST001'})
        self.assertMeta(out)
        self.assertEqual(self.fake.count(*LOOKUP), 1)
        self.assertEqual(out['total_units'], 1000 + 500 + 120)
        self.assertPlain(out)

    def test_unknown_customer_points_to_dropship(self):
        out = self.app._ai_tool_past_orders({'customer': 'ZGAMDROP'})
        self.assertIn('Dropship', out['error'])
        self.assertMeta(out)


class PastOrdersNew(Base):
    mode = 'new'

    def test_summary(self):
        out = self.app._ai_tool_past_orders({})
        self.assertEqual(out['invoices_through'], {'wholesale': '2026-08-21', 'dropship': '2026-07-31'})
        self.assertEqual(out['history_label'], self.data['history']['label'])
        alpha = [c for c in out['customers'] if c['customer'] == 'ZALPHA'][0]
        self.assertEqual(alpha['invoiced']['units'], 2600)
        self.assertEqual(alpha['awaiting_invoice']['units'], 700)
        self.assertIn('invoiced', out['totals'])
        self.assertPlain(out)

    def test_customer_uses_server_numbers(self):
        out = self.app._ai_tool_past_orders({'customer': 'Alpha Test'})
        self.assertEqual(self.fake.count(*LOOKUP), 0)
        self.assertEqual((out['total_units'], out['total_value']), (3300, 8850.0))
        bd = out['by_status']
        self.assertEqual(bd['invoiced']['units'] + bd['awaiting_invoice']['units'], out['total_units'])
        self.assertEqual(out['awaiting_invoice']['units'], 700)
        self.assertEqual((bd['pending']['pos'], bd['not_counted']['pos']), (1, 2))
        rows = self.rows_by_po(out)
        self.assertEqual(rows['10-424242']['status'], 'left_pending')
        self.assertFalse(rows['10-424242']['counted'])
        self.assertEqual(rows['66-000001']['after_cutoff_estimate']['units'], 100)
        self.assertNotIn('server_note', out)
        self.assertPlain(out)

    def test_filtered_totals_follow_the_filter(self):
        out = self.app._ai_tool_past_orders({'customer': 'ZALPHA', 'status': 'awaiting_invoice'})
        self.assertEqual(out['total_units'], 700)       # 480 + 120 + the 100 remainder
        out = self.app._ai_tool_past_orders({'customer': 'ZALPHA', 'status': 'invoiced'})
        self.assertEqual(out['total_units'], 2600)
        out = self.app._ai_tool_past_orders({'customer': 'ZALPHA', 'status': 'pending'})
        self.assertEqual((out['matched_pos'], out['total_units']), (1, 0))

    def test_view_not_ready_falls_back_to_key_lookup(self):
        acct = self.fake.po_history('ZALPHA')
        acct['history']['viewReady'] = False
        for p in acct['pos']:      # a cold server knows only book statuses
            if p.get('bucket') == 'invoiced':
                p.update(status='shipped_awaiting_invoice', bucket='awaiting', statusLabel='Left the book')
                p.pop('invoice', None)
                p.pop('afterCutoffEstimate', None)
        self.fake.po_history_override['ZALPHA'] = acct
        out = self.app._ai_tool_past_orders({'customer': 'ZALPHA'})
        self.assertEqual(self.fake.count(*LOOKUP), 1)
        rows = self.rows_by_po(out)
        self.assertEqual(rows['0077001001']['status'], 'invoiced')
        self.assertEqual(rows['77001008']['status'], 'moved')      # the book status still stands
        self.assertIn('invoice_check_note', out)
        self.assertPlain(out)

    def test_error_answers_carry_the_cutoff(self):
        self.app._fetch_sales_summary()          # a good invoice summary is in memory
        n = self.fake.count(*SALES)
        self.fake.fail[PO_HISTORY] = 500
        out = self.app._ai_tool_past_orders({'customer': 'ZALPHA'})
        self.assertIn('HTTP 500', out['error'])
        self.assertEqual(out['invoices_through'], {'wholesale': '2026-08-21', 'dropship': '2026-07-31'})
        self.assertEqual(out['history_label'], self.data['history']['label'])
        self.assertIn('dropship_note', out)
        self.assertEqual(self.fake.count(*SALES), n)   # the error answer made no extra call
        self.assertPlain(out)

    def test_customer_total_is_invoiced_plus_awaiting(self):
        s = self.fake.po_history()
        for a in s['accounts']:
            if a['customer'] == 'ZALPHA':
                a['units'], a['value'] = 1, 1.0      # a server total that disagrees with its parts
        self.fake.po_history_override[None] = s
        out = self.app._ai_tool_past_orders({'customer': 'ZALPHA'})
        bd = out['by_status']
        self.assertEqual(out['total_units'], bd['invoiced']['units'] + bd['awaiting_invoice']['units'])
        self.assertEqual((out['total_units'], out['total_value']), (3300, 8850.0))

    def test_style_lines_without_invoice_lines_are_flagged(self):
        self.fake.fail[LOOKUP] = 500
        out = self.app._ai_tool_past_orders({'customer': 'ZALPHA', 'style': 'ZZTEST001'})
        self.assertIn('booked sizes', out['invoice_check_note'])
        self.assertPlain(out)


class LookupScoping(Base):
    mode = 'new'

    def test_another_customers_invoice_is_dropped(self):
        self.fake.po_lookup = lambda body: {'ready': True, 'found': {
            '55500001': {'qty': 77, 'value': 154.0, 'firstInv': '2026-08-01', 'lastInv': '2026-08-01',
                         'invoices': 1, 'lines': [], 'customer': 'ZOMEGA', 'poKey': '55500001',
                         'matchedBy': 'exact', 'pos': ['55500001']}}}
        m, err = self.app._po_invoice_lookup([('ZALPHA', '55500001')])
        self.assertIsNone(err)
        self.assertEqual(m, {})

    def test_new_server_entries_are_not_summed_twice(self):
        e = {'qty': 1000, 'value': 3000.0, 'firstInv': '2026-07-09', 'lastInv': '2026-07-09', 'invoices': 1,
             'lines': [], 'customer': 'ZALPHA', 'poKey': '77001001', 'matchedBy': 'key', 'pos': ['77001001']}
        self.fake.po_lookup = lambda body: {'ready': True, 'found': {'0077001001': e, '77001001': e}}
        m, _err = self.app._po_invoice_lookup([('ZALPHA', '0077001001')])
        self.assertEqual(m[('ZALPHA', '0077001001')]['qty'], 1000)

    def test_never_joining_kinds_are_not_sent(self):
        m, err = self.app._po_invoice_lookup([('ZALPHA', ''), ('ZALPHA', '6.00E+11'), ('ZALPHA', 'SAMPLE')])
        self.assertEqual((m, err), ({}, None))
        self.assertEqual(self.fake.count(*LOOKUP), 0)


# ── INV-05: customer resolution ──────────────────────────────────────────

class Resolver(Base):
    def test_http_errors_are_errors(self):
        for status in (401, 500):
            self.app.__dict__['_sales_sum_cache'].update({'data': None, 'at': 0.0, 'fail_until': 0.0})
            self.fake.fail[SALES] = status
            code, err = self.app._sales_resolve_customer('Nobody Such')
            self.assertIsNone(code)
            self.assertIn(f'HTTP {status}', err['error'])

    def test_archive_only_customer_resolves(self):
        self.assertEqual(self.app._sales_resolve_customer('ZBETA1'), ('ZBETA1', None))
        self.assertEqual(self.app._sales_resolve_customer('Beta Test'), ('ZBETA1', None))
        self.fake.fail[SALES] = 500      # invoices down: the archive still resolves it
        self.app._sales_sum_cache.update({'data': None, 'at': 0.0, 'fail_until': 0.0})
        self.assertEqual(self.app._sales_resolve_customer('Beta Test Group'), ('ZBETA1', None))

    def test_families_stay_apart_and_names_can_be_ambiguous(self):
        self.assertEqual(self.app._sales_resolve_customer('ZBETA'), ('ZBETA', None))
        code, err = self.app._sales_resolve_customer('Test')
        self.assertIsNone(code)
        self.assertEqual(sorted(a['customer'] for a in err['ambiguous_customer']), ['ZALPHA', 'ZBETA1'])


# ── INV-06: sales_history_lookup ─────────────────────────────────────────

class SalesHistoryOld(Base):
    def test_every_http_failure_is_explicit(self):
        cases = [({}, SALES), ({'customer': 'ZALPHA'}, SALES), ({'style': 'ZZCAP'}, SALES)]
        for status in (401, 403, 500, 503):
            for params, route in cases:
                self.app._sales_sum_cache.update({'data': None, 'at': 0.0, 'fail_until': 0.0})
                self.fake.fail[route] = status
                out = self.app._ai_tool_sales_history(params)
                self.assertIn('error', out, (status, params))
                if status == 503:
                    self.assertIn('temporarily unavailable', out['error'])
                else:
                    self.assertIn(f'HTTP {status}', out['error'])
                self.assertNotIn('not available yet', out['error'])

    def test_summary_cutoff_from_source(self):
        out = self.app._ai_tool_sales_history({})
        self.assertEqual(out['invoices_through'], {'wholesale': '2026-08-21', 'dropship': None})
        self.assertNotIn('awaiting_invoice', out)
        self.assertPlain(out)

    def test_style_cap_with_customer_says_so(self):
        out = self.app._ai_tool_sales_history({'style': 'ZZCAP001', 'customer': 'ZALPHA'})
        self.assertEqual((out['matched_rows'], out['matched_qty']), (200, 1400))
        self.assertIn('newest 500 of 620', out['note'])
        self.assertIs(out['totals_complete'], False)
        self.assertTrue(out['truncated'])
        self.assertMeta(out)

    def test_style_without_customer_uses_full_totals(self):
        out = self.app._ai_tool_sales_history({'style': 'ZZCAP001'})
        self.assertEqual((out['matched_rows'], out['matched_qty']), (620, 5240))
        self.assertEqual(out['note'], 'Showing the newest 500 of 620 invoice lines. Totals use all 620.')

    def test_archive_only_customer(self):
        out = self.app._ai_tool_sales_history({'customer': 'ZBETA1'})
        self.assertIn('note2', out)
        self.assertMeta(out)


class SalesHistoryNew(Base):
    mode = 'new'

    def test_summary_has_both_dates_and_awaiting(self):
        out = self.app._ai_tool_sales_history({})
        self.assertEqual(out['invoices_through'], {'wholesale': '2026-08-21', 'dropship': '2026-07-31'})
        self.assertEqual(out['awaiting_invoice']['basis'], 'estimate')
        self.assertPlain(out)

    def test_summary_awaiting_by_customer(self):
        orig = self.fake.sales_summary

        def with_custs():
            s = orig()
            s['pending']['customers'] = {'ZBETA1': {'pos': 1, 'units': 10, 'value': 5.0},
                                         'ZALPHA': {'pos': 3, 'units': 700, 'value': 2050.0}}
            return s
        self.fake.sales_summary = with_custs
        out = self.app._ai_tool_sales_history({})
        self.assertEqual([c['customer'] for c in out['awaiting_invoice']['by_customer']], ['ZALPHA', 'ZBETA1'])
        self.assertEqual(out['awaiting_invoice']['basis'], 'estimate')

    def test_summary_awaiting_not_ready_is_said(self):
        orig = self.fake.sales_summary

        def not_ready():
            s = orig()
            s['pending'] = {'ready': False}
            return s
        self.fake.sales_summary = not_ready
        out = self.app._ai_tool_sales_history({})
        self.assertNotIn('awaiting_invoice', out)
        self.assertIn('still loading', out['awaiting_note'])
        self.assertPlain(out)

    def test_by_div_without_history_block(self):
        s = self.fake.sales_summary()
        s.pop('history')
        info = self.app._history_info(sales_summary=s)
        self.assertEqual(info['invoices_through'], {'wholesale': '2026-08-21', 'dropship': '2026-07-20'})
        self.assertEqual(info['history_label'], 'Invoices through Aug 21, 2026. Dropship invoices through '
                                                 'Jul 20, 2026. Shipments after that are estimates from the order book.')

    def test_style_paging_gives_full_customer_totals(self):
        self.fake.max_limit = 250
        orig = self.fake.sales

        def capped(params):
            params = dict(params)
            params['limit'] = min(int(params.get('limit') or 500), self.fake.max_limit)
            return orig(params)
        self.fake.sales = capped
        out = self.app._ai_tool_sales_history({'style': 'ZZCAP001', 'customer': 'ZALPHA'})
        self.assertEqual((out['matched_rows'], out['matched_qty']), (320, 2240))
        self.assertNotIn('note', out)
        self.assertIs(out['totals_complete'], True)
        offs = [c[2].get('offset') for c in self.fake.calls if c[:2] == SALES and c[2].get('style')]
        self.assertEqual(offs, [None, 250, 500])

    def test_account_and_range_awaiting_sections(self):
        out = self.app._ai_tool_sales_history({'customer': 'ZALPHA'})
        aw = out['awaiting_invoice']
        self.assertEqual((aw['pos'], aw['units'], aw['basis']), (3, 700, 'estimate'))
        out = self.app._ai_tool_sales_history({'customer': 'ZALPHA', 'month': '2026-09'})
        self.assertEqual(sorted(r['po'] for r in out['awaiting_invoice']['rows']),
                         ['66-000001', '77001004', '77001010'])
        self.assertEqual(out['awaiting_invoice']['units'], 700)
        self.assertPlain(out)


# ── INV-07: build_sales_sheet ────────────────────────────────────────────

class SheetBase(Base):
    def build(self, **params):
        out = self.app._ai_tool_build_sales_sheet(params)
        if 'error' in out:
            return out, None
        body = self.h.s3.puts[-1]['Body']
        return out, openpyxl.load_workbook(io.BytesIO(body))

    def cells(self, wb):
        return [str(c.value) for ws in wb.worksheets for row in ws.iter_rows() for c in row if c.value is not None]

    def tab_pos(self, wb, name):
        ws = wb[name]
        head = [c.value for c in ws[1]]
        i = head.index('PO #')
        return sorted({str(r[i].value) for r in ws.iter_rows(min_row=2)})


class SheetOld(SheetBase):
    def test_month_sheet_never_lists_an_invoiced_po_as_awaiting(self):
        out, wb = self.build(customer='ZALPHA', month='2026-08', include=['invoiced', 'shipped_pending'])
        self.assertEqual(self.tab_pos(wb, 'Invoiced'), ['66-000001', '77001003'])
        # 77001002 was invoiced Jul 31 and left the book Aug 5: not awaiting.
        self.assertEqual(self.tab_pos(wb, 'Shipped, not yet invoiced'), ['77001008'])
        body = [c for c in self.fake.calls if c[:2] == LOOKUP][0][3]
        self.assertIn('77001002', body['pos'])
        self.assertNotIn('from', body)
        self.assertTrue(any('renumbered' in w for w in out['warnings']))
        self.assertPlain(out)
        for text in self.cells(wb):
            self.assertNotIn(EM, text)
            self.assertNotIn(EN, text)

    def test_padded_po_lands_in_one_tab(self):
        out, wb = self.build(customer='ZALPHA', month='2026-07', include=['invoiced', 'shipped_pending'])
        self.assertEqual(self.tab_pos(wb, 'Invoiced'), ['77001001', '77001002'])
        self.assertNotIn('Shipped, not yet invoiced', wb.sheetnames)

    def test_lookup_failure_leaves_the_tab_out_and_says_so(self):
        self.fake.fail[LOOKUP] = 500
        out, wb = self.build(customer='ZALPHA', month='2026-08')
        self.assertNotIn('Shipped, not yet invoiced', wb.sheetnames)
        self.assertTrue(any('left out so nothing is counted twice' in w for w in out['warnings']))

    def test_archive_only_customer_gets_a_sheet(self):
        out, wb = self.build(customer='Beta Test Group')
        self.assertEqual(out['customer'], 'ZBETA1')
        self.assertIn('On Order', wb.sheetnames)
        self.assertMeta(out)

    def test_summary_tab_states_the_cutoff(self):
        out, wb = self.build(customer='ZALPHA')
        self.assertEqual(wb['Summary']['A3'].value, out['history_label'])
        self.assertEqual(wb['Summary']['A1'].value, 'Alpha Test Stores Sales Sheet')


class SheetNew(SheetBase):
    mode = 'new'

    def test_awaiting_tab_comes_from_the_server(self):
        out, wb = self.build(customer='ZALPHA', include=['invoiced', 'awaiting_invoice', 'on_order'])
        self.assertEqual(self.fake.count(*LOOKUP), 0)
        self.assertEqual(wb.sheetnames, ['Summary', 'Invoiced', 'Shipped, not yet invoiced', 'On Order'])
        ws = wb['Shipped, not yet invoiced']
        self.assertEqual([c.value for c in ws[1]],
                         ['Image', 'Style', 'PO #', 'Status', 'Left the book', 'Est. ship date',
                          'Qty (estimate until invoiced)', 'Est. Price', 'Est. Value'])
        self.assertEqual(self.tab_pos(wb, 'Shipped, not yet invoiced'), ['66-000001', '77001004', '77001010'])
        tab = [t for t in out['tabs'] if t['tab'] == 'Shipped, not yet invoiced'][0]
        self.assertEqual((tab['units'], tab['basis']), (700, 'estimate until invoiced'))
        self.assertEqual(out['history_label'], self.data['history']['label'])
        self.assertIn(self.data['history']['label'], self.cells(wb))
        self.assertPlain(out)

    def pend_lines(self, wb, po):
        ws = wb['Shipped, not yet invoiced']
        head = [c.value for c in ws[1]]
        i_po, i_st, i_q = head.index('PO #'), head.index('Style'), head.index('Qty (estimate until invoiced)')
        return [(r[i_st].value or '', r[i_q].value) for r in ws.iter_rows(min_row=2) if str(r[i_po].value) == po]

    def test_pending_lines_follow_the_peak_rule(self):
        # A server still on the peak estimate sends no estLines: the peak lines add up.
        a4 = self.data['archive'][3]['_new']
        a4.pop('estLines')
        a4.update(estUnits=500, estValue=1500.0, countUnits=500, countValue=1500.0)
        out, wb = self.build(customer='ZALPHA', include=['shipped_pending'])
        self.assertEqual(self.pend_lines(wb, '77001004'), [('ZZTEST001SLS', 500)])

    def test_pending_lines_that_do_not_add_up_become_one_line(self):
        # No style list adds up to the counted 480: one PO line keeps the tab equal to the count.
        self.data['archive'][3]['_new'].pop('estLines')
        out, wb = self.build(customer='ZALPHA', include=['shipped_pending'])
        self.assertEqual(self.pend_lines(wb, '77001004'), [('', 480)])
        tab = [t for t in out['tabs'] if t['tab'] == 'Shipped, not yet invoiced'][0]
        self.assertEqual(tab['units'], 700)

    def test_moved_po_is_not_awaiting(self):
        out, wb = self.build(customer='ZBETA1')
        self.assertNotIn('Shipped, not yet invoiced', wb.sheetnames)
        self.assertIn('On Order', wb.sheetnames)

    def test_cold_view_uses_archive_statuses(self):
        r = self.fake.sales
        def no_pending(params):
            out = r(params)
            out.pop('pending', None)
            return out
        self.fake.sales = no_pending
        out, wb = self.build(customer='ZALPHA', month='2026-09', include=['shipped_pending'])
        self.assertEqual(self.tab_pos(wb, 'Shipped, not yet invoiced'), ['66-000001', '77001004', '77001010'])
        self.assertEqual(self.fake.count(*LOOKUP), 0)


# ── INV-08: tool text ────────────────────────────────────────────────────

class ToolText(unittest.TestCase):
    def setUp(self):
        self.app = load_app()

    def test_schemas(self):
        tools = {t['name']: t for t in self.app._AI_AGENT_TOOLS}
        for name in ('past_orders_lookup', 'sales_history_lookup', 'build_sales_sheet'):
            text = json.dumps(tools[name]).lower()
            self.assertNotIn('cancel', text, name)
            self.assertNotIn('pulled', text, name)
            self.assertNotIn(EM, json.dumps(tools[name], ensure_ascii=False), name)
        past = tools['past_orders_lookup']
        self.assertEqual(past['input_schema']['properties']['status']['enum'],
                         ['all', 'invoiced', 'awaiting_invoice', 'not_counted', 'pending'])
        for name in ('past_orders_lookup', 'sales_history_lookup'):
            d = tools[name]['description']
            self.assertIn('invoices_through', d)
            self.assertIn('history_label', d)
            self.assertIn('ropship', d)
        d = past['description']
        for status in ('invoiced', 'shipped_awaiting_invoice', 'left_pending', 'moved', 'reentered',
                       'left_before_window', 'closed_not_invoiced', 'left_before_window_no_invoice'):
            self.assertIn(status, d)
        self.assertIn('never appears in this archive', d)

    def test_mcp_instructions(self):
        t = self.app._MCP_INSTRUCTIONS
        self.assertNotIn('cancel', t.lower())
        self.assertIn('never in this archive', t)
        self.assertIn('invoice cut-off', t)

    def test_new_code_has_no_long_dashes(self):
        fns = [self.app._ai_tool_past_orders, self.app._ai_tool_sales_history, self.app._ai_tool_build_sales_sheet,
               self.app._past_po_state, self.app._past_po_row, self.app._history_info, self.app._oo_json,
               self.app._fetch_po_history, self.app._po_invoice_lookup, self.app._sales_resolve_customer_ex,
               self.app._sales_style_search, self.app._sheet_pending_rows, self.app._po_key]
        for fn in fns:
            src = inspect.getsource(fn)
            self.assertNotIn(EM, src, fn.__name__)
            self.assertNotIn(EN, src, fn.__name__)


if __name__ == '__main__':
    unittest.main()
