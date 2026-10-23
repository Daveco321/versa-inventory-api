"""Warehouse line sheet tabs print WAREHOUSE availability only (Oct 16 2026).

On a stock:'warehouse' tab without the mailbox's per-warehouse breakdown (the
platform AI chat and the MCP connector) the Total ATS column printed the style's
WHOLE ATS, production on order included: a style with 24,444 units in the
warehouse and 200,028 incoming read 203,969 while the warehouse deck said 18,953.
The number must come from the warehouse photo deck's own computation
(_pres_stock_cards('warehouse', ...), smart routing via _pres_route_sku), so a
sheet and a deck can never disagree. The mailbox's breakdown sheets, overseas
tabs and all-stock tabs stay exactly as they were.

Everything here is synthetic and offline: made-up style numbers (serial 9xx,
fabric ZZ), a made-up feed, ledger and order book, S3 replaced by a recorder.
The shared fixtures come from test_line_sheet_overseas.py.
"""
import io
import os
import sys
import unittest
from datetime import datetime

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
try:
    from tests.test_line_sheet_overseas import (  # noqa: E402
        NOW, ORDERS, S_AS_OS, S_AS_WH, S_NJLOT, S_NONE, S_ROUTED, S_SPLIT, S_VAR, S_WHCOVER, S_WHONLY,
        _Base, _item, _lot, _order)
except ImportError:   # run from inside tests/
    from test_line_sheet_overseas import (  # noqa: E402
        NOW, ORDERS, S_AS_OS, S_AS_WH, S_NJLOT, S_NONE, S_ROUTED, S_SPLIT, S_VAR, S_WHCOVER, S_WHONLY,
        _Base, _item, _lot, _order)

S_WHZERO = 'TJNAZZ951SLS'     # 500 in the warehouse, all of it committed to an October order
S_NJSTOCK = 'TJDKZZ952SLS'    # 400 in NJ, nothing else (admin only)


class _WhBase(_Base):
    EXTRA_FEED = [_item(S_WHZERO, 'NAUTICA', jtw=500, inc=1000, committed=-500),
                  _item(S_NJSTOCK, 'DKNY', nj=400)]
    EXTRA_LEDGER = [_lot(S_WHZERO, 1000, '2026-12-15', prod='ZZ26960')]
    EXTRA_ORDERS = [_order(S_WHZERO, 500, '2026-10-20', '2026-10-30', '9100011')]

    def wh_deck(self, customer_view):
        """What the warehouse photo deck prints per base style for the same inputs."""
        got = self.app._pres_stock_cards('warehouse', customer_view, self.app._pres_filter_spec({}), {})
        out = {}
        for c in got['cards']:
            out[c['base']] = out.get(c['base'], 0) + c['number']
        return out

    def spy_cards(self):
        calls = []
        real = self._saved['_pres_stock_cards']

        def spy(*a, **k):
            calls.append(a)
            return real(*a, **k)
        self.app._pres_stock_cards = spy
        return calls


class WarehouseTabs(_WhBase):

    # a. production on order never reaches a warehouse tab's Total ATS column
    def test_warehouse_tab_prints_only_the_warehouse_share_admin(self):
        out = self.sheet([{'title': 'WH', 'skus': [S_SPLIT], 'stock': 'warehouse'}])
        self.assertNotIn('error', out)
        self.assertEqual(self.built['tabs'][0]['view_mode'], 'ats')
        self.assertEqual(self.printed(), {S_SPLIT: 6732})     # was 7,200 (6,732 + 468 incoming)

    def test_warehouse_tab_prints_only_the_warehouse_share_customer(self):
        self.sheet([{'title': 'WH', 'skus': [S_SPLIT], 'stock': 'warehouse'}], customer_view=True)
        self.assertEqual(self.printed(), {S_SPLIT: 6732})

    def test_sheet_equals_the_warehouse_deck_admin_and_customer(self):
        bases = [S_SPLIT, S_ROUTED, S_WHCOVER, S_AS_OS, S_AS_WH, S_NJLOT, S_NONE, S_WHONLY, S_NJSTOCK]
        for cv in (False, True):
            with self.subTest(customer_view=cv):
                self.sheet([{'title': 'WH', 'skus': bases, 'stock': 'warehouse'}], customer_view=cv)
                deck = self.wh_deck(cv)
                self.assertEqual(self.printed(), {b: deck[b] for b in bases if deck.get(b, 0) > 0})

    # b. smart routing decides which side pays for committed demand
    def test_order_only_production_can_make_is_charged_to_incoming(self):
        # 600 ships in October (warehouse), 1,500 in January (December lot): the
        # warehouse keeps 1,000 - 600. Warehouse-first FIFO would charge all 2,100 to
        # the warehouse first and leave nothing.
        self.sheet([{'title': 'WH', 'skus': [S_ROUTED], 'stock': 'warehouse'}])
        self.assertEqual(self.printed(), {S_ROUTED: 400})
        fifo = max(0, 1000 - min(2100, 1000))
        self.assertGreater(self.printed()[S_ROUTED], fifo)
        route = self.app._pres_route_sku(
            {'jtw': 1000, 'tr': 0, 'dcw': 0, 'qa': 0, 'nj': 0, 'abfi': 0, 'committed': -2100, 'allocated': 0,
             'incoming': 2000},
            [{'units': 2000, 'arr': datetime(2026, 12, 15).date(), 'etd': None, 'fob_flag': False,
              'hidden': False, 'sup': False}],
            [o for o in ORDERS if o['style'] == S_ROUTED], [], [], NOW.date(), sku=S_ROUTED)
        self.assertEqual(route['wh'], 600)

    def test_order_only_the_warehouse_can_make_comes_out_of_the_warehouse(self):
        self.sheet([{'title': 'WH', 'skus': [S_WHCOVER], 'stock': 'warehouse'}])
        self.assertEqual(self.printed(), {S_WHCOVER: 400})     # was 2,400 total ATS

    # c. deduction assignments win over routing
    def test_deduction_assignments_are_respected(self):
        self.sheet([{'title': 'WH', 'skus': [S_AS_OS, S_AS_WH], 'stock': 'warehouse'}])
        self.assertEqual(self.printed(), {S_AS_OS: 1000, S_AS_WH: 400})

    # d. NJ/ABFI stock is admin-only
    def test_customer_view_never_counts_nj_stock(self):
        self.sheet([{'title': 'WH', 'skus': [S_WHONLY], 'stock': 'warehouse'}])
        self.assertEqual(self.printed(), {S_WHONLY: 350})
        self.sheet([{'title': 'WH', 'skus': [S_WHONLY], 'stock': 'warehouse'}], customer_view=True)
        self.assertEqual(self.printed(), {S_WHONLY: 300})

    def test_nj_only_style_is_admin_only(self):
        self.sheet([{'title': 'WH', 'skus': [S_SPLIT, S_NJSTOCK], 'stock': 'warehouse'}])
        self.assertEqual(self.printed(), {S_SPLIT: 6732, S_NJSTOCK: 400})
        # the old customer sheet printed this style with a 0 (400 - 400 NJ)
        out = self.sheet([{'title': 'WH', 'skus': [S_SPLIT, S_NJSTOCK], 'stock': 'warehouse'}],
                         customer_view=True)
        self.assertEqual(self.printed(), {S_SPLIT: 6732})
        self.assertEqual(self.entry(out, 'WH')['styles_without_warehouse_availability'], [S_NJSTOCK])

    # e. nothing left in the warehouse: off the tab, and said so
    def test_style_with_no_warehouse_left_is_named_on_a_curated_tab(self):
        out = self.sheet([{'title': 'WH', 'skus': [S_SPLIT, S_WHZERO, S_VAR], 'stock': 'warehouse'}])
        self.assertEqual(list(self.items()), [S_SPLIT])
        e = self.entry(out, 'WH')
        self.assertEqual(e['styles'], 1)
        self.assertEqual(e['styles_without_warehouse_availability'], [S_WHZERO, S_VAR])

    def test_style_with_no_warehouse_left_is_counted_on_a_filter_tab(self):
        out = self.sheet([{'title': 'NA WH', 'brands': ['NAUTICA'], 'stock': 'warehouse'}])
        self.assertEqual(self.printed(), {S_SPLIT: 6732, S_NONE: 5000, S_WHCOVER: 400, S_ROUTED: 400,
                                          S_WHONLY: 350})
        e = self.entry(out, 'NA WH')
        self.assertEqual(e['styles'], 5)
        self.assertEqual(e['styles_without_warehouse_availability'], 1)   # S_WHZERO

    # row order: a filter tab sorts on the number it prints, a curated tab keeps it
    def test_filter_tab_rows_sort_by_the_warehouse_number_stable(self):
        # The filter orders by gross total ATS: S_NJLOT 3,100 (3,000 of it incoming)
        # first. S_AS_WH and S_NJSTOCK both print 400 and keep the filter's order.
        self.sheet([{'title': 'DK WH', 'brands': ['DKNY'], 'stock': 'warehouse'}])
        self.assertEqual([(s, it['total_ats']) for s, it in self.items().items()],
                         [(S_AS_OS, 1000), (S_AS_WH, 400), (S_NJSTOCK, 400), (S_NJLOT, 100)])

    def test_customer_filter_tab_drops_the_nj_only_style(self):
        out = self.sheet([{'title': 'DK WH', 'brands': ['DKNY'], 'stock': 'warehouse'}], customer_view=True)
        self.assertEqual(list(self.items()), [S_AS_OS, S_AS_WH, S_NJLOT])
        self.assertEqual(self.entry(out, 'DK WH')['styles_without_warehouse_availability'], 1)

    def test_curated_tab_keeps_the_callers_order(self):
        self.sheet([{'title': 'WH', 'skus': [S_WHONLY, S_ROUTED, S_SPLIT], 'stock': 'warehouse'}])
        self.assertEqual(list(self.items()), [S_WHONLY, S_ROUTED, S_SPLIT])

    # min_units: applied again to the routed number on a filter tab
    def test_filter_tab_min_units_applies_to_the_routed_number(self):
        # S_WHCOVER and S_ROUTED pass the filter on 1,000 gross but only 400 is left.
        out = self.sheet([{'title': 'NA WH', 'brands': ['NAUTICA'], 'stock': 'warehouse', 'min_units': 1000}])
        self.assertEqual(self.printed(), {S_SPLIT: 6732, S_NONE: 5000})
        e = self.entry(out, 'NA WH')
        self.assertEqual(e['styles'], 2)
        self.assertEqual(e['styles_below_min_units'], 2)
        self.assertNotIn('styles_without_warehouse_availability', e)

    def test_skus_tab_honors_min_units(self):
        # A minimum is about the number printed, curated tab or not (Oct 23 2026).
        out = self.sheet([{'title': 'WH', 'skus': [S_SPLIT, S_ROUTED], 'stock': 'warehouse', 'min_units': 1000}])
        self.assertEqual(self.printed(), {S_SPLIT: 6732})
        e = self.entry(out, 'WH')
        self.assertEqual(e['styles'], 1)
        self.assertEqual(e['styles_below_min_units'], [S_ROUTED])

    def test_skus_tab_without_min_units_keeps_every_style(self):
        out = self.sheet([{'title': 'WH', 'skus': [S_ROUTED], 'stock': 'warehouse'}])
        self.assertEqual(self.printed(), {S_ROUTED: 400})
        self.assertNotIn('styles_below_min_units', self.entry(out, 'WH'))

    def test_min_units_keeps_a_row_exactly_at_the_minimum(self):
        self.sheet([{'title': 'WH', 'skus': [S_ROUTED], 'stock': 'warehouse', 'min_units': 400}])
        self.assertEqual(self.printed(), {S_ROUTED: 400})

    # emptied tabs
    def test_a_tab_left_empty_is_skipped_with_a_neutral_reason(self):
        out = self.sheet([{'title': 'WH', 'skus': [S_WHZERO], 'stock': 'warehouse'},
                          {'title': 'OS', 'skus': [S_SPLIT], 'stock': 'overseas'}])
        self.assertEqual([t['tab_name'] for t in self.built['tabs']], ['OS'])
        e = self.entry(out, 'WH')
        self.assertTrue(e['skipped'])
        self.assertEqual(e['styles'], 0)
        self.assertIn('warehouse', e['reason'])
        self.assertNotIn('committed', e['reason'])
        self.assertNotIn('window', e['reason'])
        self.assertNotIn('arrival_window', e)

    def test_tab_emptied_by_min_units_is_skipped_with_a_reason(self):
        out = self.sheet([{'title': 'WH', 'search': '932', 'stock': 'warehouse', 'min_units': 500},
                          {'title': 'WH2', 'skus': [S_SPLIT], 'stock': 'warehouse'}])
        self.assertEqual([t['tab_name'] for t in self.built['tabs']], ['WH2'])
        e = self.entry(out, 'WH')
        self.assertTrue(e['skipped'])
        self.assertEqual(e['styles_below_min_units'], 1)
        self.assertIn('500', e['reason'])
        self.assertIn('warehouse', e['reason'])

    def test_every_tab_empty_is_an_error(self):
        out = self.sheet([{'title': 'WH', 'skus': [S_WHZERO, S_VAR], 'stock': 'warehouse'}])
        self.assertIn('error', out)
        self.assertIn('warehouse', out['error'])
        self.assertNotIn('tabs', self.built)
        self.assertNotIn('body', self.built)

    # a routing failure is an error, never a sheet with total ATS
    def test_routing_failure_returns_an_error(self):
        def boom(*a, **k):
            raise RuntimeError('ledger exploded')
        self.app._pres_stock_cards = boom
        out = self.sheet([{'title': 'WH', 'skus': [S_SPLIT], 'stock': 'warehouse'}])
        self.assertIn('error', out)
        self.assertIn('warehouse', out['error'])
        self.assertNotIn('tabs', self.built)
        self.assertNotIn('body', self.built)

    def test_routing_runs_once_per_build(self):
        calls = self.spy_cards()
        out = self.sheet([{'title': 'A', 'skus': [S_SPLIT], 'stock': 'warehouse'},
                          {'title': 'B', 'brands': ['NAUTICA'], 'stock': 'warehouse', 'min_units': 100},
                          {'title': 'C', 'brands': ['DKNY'], 'stock': 'warehouse', 'split_by_brand': True},
                          {'title': 'D', 'skus': [S_SPLIT]}])
        self.assertNotIn('error', out)
        self.assertEqual(len(calls), 1)
        self.assertEqual(calls[0][0], 'warehouse')
        self.assertEqual(calls[0][3], {})   # a warehouse deck has no arrival window

    def test_summary_note_and_order_feed(self):
        tab = {'title': 'WH', 'skus': [S_SPLIT], 'stock': 'warehouse'}
        note = next(e for e in self.sheet([tab])['tabs'] if 'warehouse_ats' in e)
        self.assertIs(note['open_orders_feed_ok'], True)
        self.assertNotIn(chr(0x2014), note['warehouse_ats'])   # no em dash in user-facing text
        self.app._fetch_all_open_orders = lambda: ([], False)
        note = next(e for e in self.sheet([tab])['tabs'] if 'warehouse_ats' in e)
        self.assertIs(note['open_orders_feed_ok'], False)


class OtherTabsStayPut(_WhBase):
    """The mailbox's breakdown warehouse tabs, overseas tabs and all-stock tabs print
    exactly what they printed before (values pinned from 22f2a01's code)."""

    BASES = [S_SPLIT, S_ROUTED, S_NONE, S_WHONLY, S_NJLOT]

    def test_mailbox_breakdown_warehouse_tab_unchanged(self):
        calls = self.spy_cards()
        out = self.sheet([{'title': 'WH', 'skus': self.BASES + [S_WHZERO], 'stock': 'warehouse'}],
                         customer_view=True, breakdown=True)
        self.assertNotIn('error', out)
        self.assertEqual(calls, [])          # the warehouse pass never runs there
        cols = {s: (it['wh_ats_jtw'], it['wh_ats_tr'], it['wh_ats_dcw'], it['wh_ats_qa'], it['wh_ats_total'])
                for s, it in self.items().items()}
        self.assertEqual(cols, {S_SPLIT: (6732, 0, 0, 0, 6732), S_ROUTED: (400, 0, 0, 0, 400),
                                S_NONE: (5000, 0, 0, 0, 5000), S_WHONLY: (300, 0, 0, 0, 300),
                                S_NJLOT: (100, 0, 0, 0, 100)})
        self.assertEqual(self.printed(), {S_SPLIT: 7200, S_ROUTED: 900, S_NONE: 5000, S_WHONLY: 300,
                                          S_NJLOT: 2100})
        e = self.entry(out, 'WH')
        self.assertEqual(e['styles'], 5)
        self.assertNotIn('styles_without_warehouse_availability', e)
        self.assertFalse(any('warehouse_ats' in x for x in out['tabs']))
        bd = next(x for x in out['tabs'] if x.get('warehouse_breakdown'))
        self.assertEqual(bd['styles_dropped_no_warehouse_stock'], 1)   # S_WHZERO, as before

    def test_mailbox_breakdown_workbook_cells_unchanged(self):
        import openpyxl
        self.app.build_multi_brand_excel = self.real_excel
        self.sheet([{'title': 'WH', 'skus': [S_SPLIT, S_ROUTED], 'stock': 'warehouse'}],
                   customer_view=True, breakdown=True)
        ws = openpyxl.load_workbook(io.BytesIO(self.built['body']))['WH']
        head = [c.value for c in ws[1]]
        self.assertNotIn('Total ATS', head)
        sku_col, col = head.index('SKU'), head.index('Warehouse ATS')
        cells = {row[sku_col].value: row[col].value
                 for row in ws.iter_rows(min_row=2) if row[sku_col].value in (S_SPLIT, S_ROUTED)}
        self.assertEqual(cells, {S_SPLIT: 6732, S_ROUTED: 400})

    def test_overseas_tab_unchanged_next_to_a_warehouse_tab(self):
        for cv in (False, True):
            with self.subTest(customer_view=cv):
                alone = self.sheet([{'title': 'OS', 'skus': self.BASES, 'stock': 'overseas'}], customer_view=cv)
                os_alone = self.printed()
                mixed = self.sheet([{'title': 'OS', 'skus': self.BASES, 'stock': 'overseas'},
                                    {'title': 'WH', 'skus': self.BASES, 'stock': 'warehouse'}], customer_view=cv)
                self.assertEqual(self.printed(0), os_alone)
                self.assertEqual(os_alone, {S_SPLIT: 468, S_ROUTED: 500, S_NJLOT: 2000 if cv else 3000})
                self.assertEqual(self.entry(mixed, 'OS'), self.entry(alone, 'OS'))
                note = lambda o: next(x for x in o['tabs'] if 'overseas_ats' in x)
                self.assertEqual(note(mixed), note(alone))

    def test_any_tabs_unchanged(self):
        for cv, pinned in ((False, {S_SPLIT: 7200, S_ROUTED: 900, S_NONE: 5000, S_WHONLY: 350, S_NJLOT: 3100}),
                           (True, {S_SPLIT: 7200, S_ROUTED: 900, S_NONE: 5000, S_WHONLY: 300, S_NJLOT: 2100})):
            with self.subTest(customer_view=cv):
                calls = self.spy_cards()
                out = self.sheet([{'title': 'ALL', 'skus': self.BASES},
                                  {'title': 'WH', 'skus': self.BASES, 'stock': 'warehouse'}], customer_view=cv)
                self.assertEqual(self.built['tabs'][0]['view_mode'], 'all')
                self.assertEqual(self.printed(0), pinned)
                self.assertEqual(self.entry(out, 'ALL'), {'tab': 'ALL', 'styles': 5})
                self.assertEqual([c[0] for c in calls], ['warehouse'])   # never for the all tab
                self.app._pres_stock_cards = self._saved['_pres_stock_cards']

    def test_any_filter_tab_unchanged(self):
        self.sheet([{'title': 'ALL', 'brands': ['DKNY']}])
        self.assertEqual([(s, it['total_ats']) for s, it in self.items().items()],
                         [(S_NJLOT, 3100), (S_AS_OS, 2400), (S_AS_WH, 2400), (S_VAR, 1400), (S_NJSTOCK, 400)])


class WorkbookCells(_WhBase):
    """The Total ATS cell in a real workbook, customer and admin layouts."""

    def _cells(self, customer_view):
        import openpyxl
        self.app.build_multi_brand_excel = self.real_excel
        out = self.sheet([{'title': 'WH', 'skus': [S_SPLIT, S_ROUTED, S_WHONLY], 'stock': 'warehouse'}],
                         customer_view=customer_view)
        self.assertNotIn('error', out)
        ws = openpyxl.load_workbook(io.BytesIO(self.built['body']))['WH']
        head = [c.value for c in ws[1]]
        sku_col = head.index('SKU')
        rows = {}
        for row in ws.iter_rows(min_row=2):
            if row[sku_col].value in (S_SPLIT, S_ROUTED, S_WHONLY):
                rows[row[sku_col].value] = {h: row[i].value for i, h in enumerate(head)}
        return head, rows

    def test_customer_total_ats_cell(self):
        head, rows = self._cells(True)
        self.assertEqual(head[-2:], ['Warehouse', 'Total ATS'])
        self.assertEqual({s: r['Total ATS'] for s, r in rows.items()},
                         {S_SPLIT: 6732, S_ROUTED: 400, S_WHONLY: 300})
        self.assertEqual(rows[S_WHONLY]['Warehouse'], 'JTW')   # NJ never named

    def test_admin_total_ats_cell_and_gross_columns(self):
        head, rows = self._cells(False)
        for h in ('JTW', 'NJ', 'Incoming', 'Total Warehouse', 'Total ATS', 'Delivery'):
            self.assertIn(h, head)
        self.assertEqual({s: r['Total ATS'] for s, r in rows.items()},
                         {S_SPLIT: 6732, S_ROUTED: 400, S_WHONLY: 350})
        # the gross columns stay gross
        self.assertEqual((rows[S_ROUTED]['JTW'], rows[S_ROUTED]['Incoming'], rows[S_ROUTED]['Total Warehouse']),
                         (1000, 2000, 1000))
        self.assertEqual((rows[S_WHONLY]['NJ'], rows[S_WHONLY]['Total Warehouse']), (50, 350))
        self.assertEqual({r['Delivery'] for r in rows.values()}, {'ATS'})


S_BIG = 'TJNAZZ961SLS'        # the reported shape: 24,444 in the warehouse, 200,028 incoming


class ReportedShape(_Base):
    """A warehouse sheet printed 203,969 for 24,444 in the warehouse and 200,028
    incoming; the warehouse deck prints 18,953 (5,491 of the 20,503 committed ship
    before the lot lands, the rest after)."""

    EXTRA_FEED = [_item(S_BIG, 'NAUTICA', jtw=24444, inc=200028, committed=-20503)]
    EXTRA_LEDGER = [_lot(S_BIG, 200028, '2026-12-15', prod='ZZ26970')]
    EXTRA_ORDERS = [_order(S_BIG, 5491, '2026-10-20', '2026-10-30', '9100021'),
                    _order(S_BIG, 15012, '2027-01-10', '2027-01-20', '9100022')]

    def test_reported_shape(self):
        for cv in (False, True):
            with self.subTest(customer_view=cv):
                self.sheet([{'title': 'WH', 'skus': [S_BIG], 'stock': 'warehouse'}], customer_view=cv)
                self.assertEqual(self.printed(), {S_BIG: 18953})       # was 203,969
                self.assertEqual(self.items()[S_BIG]['incoming'], 200028)   # Incoming stays gross
        # FIFO (all 20,503 out of the warehouse first) would say 3,941
        self.assertGreater(18953, 24444 - 20503)


S_SIZEONLY = 'TJNAZZ971SLS'   # warehouse stock only on by-size rows
S_SIZEMIX = 'TJNAZZ972SLS'    # a style row plus one by-size row


class SizeRows(_Base):
    """The deck skips by-size rows (_is_sized_sku), and so did the mailbox breakdown.
    A base whose only warehouse stock sits on size rows drops (and is named); a
    size row next to a style row no longer adds its units."""

    EXTRA_FEED = [_item(S_SIZEONLY + '-M', 'NAUTICA', jtw=120), _item(S_SIZEONLY + '-L', 'NAUTICA', jtw=80),
                  _item(S_SIZEMIX, 'NAUTICA', jtw=1000), _item(S_SIZEMIX + '-M', 'NAUTICA', jtw=146)]

    def test_size_rows_leave_the_warehouse_number(self):
        self.assertTrue(self.app._is_sized_sku(S_SIZEONLY + '-M'))
        out = self.sheet([{'title': 'WH', 'skus': [S_SIZEONLY, S_SIZEMIX], 'stock': 'warehouse'}],
                         customer_view=True)
        self.assertEqual(self.printed(), {S_SIZEMIX: 1000})          # was 1,146
        self.assertEqual(self.entry(out, 'WH')['styles_without_warehouse_availability'], [S_SIZEONLY])
        # the mailbox's breakdown says the same
        self.sheet([{'title': 'WH', 'skus': [S_SIZEONLY, S_SIZEMIX], 'stock': 'warehouse'}],
                   customer_view=True, breakdown=True)
        self.assertEqual({s: it['wh_ats_total'] for s, it in self.items().items()}, {S_SIZEMIX: 1000})


if __name__ == '__main__':
    unittest.main()
