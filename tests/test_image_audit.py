"""Missing Image and Mismatch Report (image_audit.py, Oct 15 2026).

Run from the repo root:  python -B -m unittest discover -s tests -p "test_image_audit.py"
"""
import io
import json
import os
import sys
import types
import unittest

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, ROOT)

import image_audit as ia                     # noqa: E402

from PIL import Image                        # noqa: E402


def jpeg(color=(255, 255, 255), size=(64, 48), stripes=False):
    im = Image.new('RGB', size, color)
    if stripes:
        for x in range(0, size[0], 6):
            for y in range(size[1]):
                im.putpixel((x, y), (0, 0, 0))
    b = io.BytesIO()
    im.save(b, 'JPEG')
    return b.getvalue()


class FakeBody:
    def __init__(self, b):
        self.b = b

    def read(self):
        return self.b


class NoSuchKey(Exception):
    def __init__(self):
        super().__init__('NoSuchKey')
        self.response = {'Error': {'Code': 'NoSuchKey'}}


class FakeS3:
    def __init__(self, objects=None):
        self.objects = dict(objects or {})      # key -> bytes
        self.puts = []

    def list_objects_v2(self, Bucket, Prefix, MaxKeys=1000, ContinuationToken=None):
        keys = sorted(k for k in self.objects if k.startswith(Prefix))
        return {'Contents': [{'Key': k, 'ETag': '"%s"' % ia.hashlib.md5(self.objects[k]).hexdigest(),
                              'Size': len(self.objects[k])} for k in keys], 'IsTruncated': False}

    def get_object(self, Bucket, Key):
        if Key not in self.objects:
            raise NoSuchKey()
        return {'Body': FakeBody(self.objects[Key])}

    def put_object(self, Bucket, Key, Body, ContentType=None):
        self.objects[Key] = Body if isinstance(Body, bytes) else Body.encode()
        self.puts.append(Key)

    def delete_object(self, Bucket, Key):
        self.objects.pop(Key, None)


def listing(names):
    return {n: ('etag-' + n, 10) for n in names}


EMPTY_L = {'ovr': {}, 'brand': {}, 'sync': {}, 'swatch': {}}


class TestCatalogMirror(unittest.TestCase):
    def test_brand_correction(self):
        self.assertEqual(ia.correct_brand({'sku': 'TMNASU201SLS', 'brand': 'NAUTICA'}, {}), 'NAUTICA')
        self.assertEqual(ia.correct_brand({'sku': 'ROKLAW001SLS', 'brand': 'KLP'}, {}), 'KL')
        self.assertEqual(ia.correct_brand({'sku': 'ROXXAW001SLS', 'brand': 'NM'}, {}), 'NICOLE')
        self.assertEqual(ia.correct_brand({'sku': 'LUCK123', 'brand': 'X'}, {}), 'LUCKY')
        # a feed brand the catalog knows wins over the style code (the VEGBWS case)
        self.assertEqual(ia.correct_brand({'sku': 'VEGBWS024RFS-V', 'brand': 'NAUTICA'}, {}), 'NAUTICA')
        self.assertEqual(ia.correct_brand({'sku': 'TMNASU201SLS', 'brand': 'NAUTICA'},
                                          {'TMNASU201SLS': {'brand': 'DKNY'}}), 'DKNY')

    def test_image_code_and_fallback_keys(self):
        self.assertEqual(ia.image_code('TMNASU201SLS', 'NAUTICA'), 'NA_201')
        self.assertEqual(ia.image_code('TMNTSU201SLS', 'NAUTICA'), 'NT_201')
        self.assertEqual(ia.color_fallback_keys('ROUSDPP12SRS', 'USPA'), ['US_P12', 'US_012'])
        self.assertEqual(ia.color_fallback_keys('ROGBPO036SSM', 'BEENE'), ['GB_SW_036', 'GB_036'])

    def test_color_name_order(self):
        cmap = {'NA_201': 'WHT SOLID', 'TMNASU201SLS': 'navy grnd white check', 'NM_002': 'white||navy print'}
        self.assertEqual(ia.color_name({}, cmap, 'TMNASU201SLS', 'NAUTICA'), ('Navy Grnd White Check', False))
        self.assertEqual(ia.color_name({}, cmap, 'RONASU201SLS', 'NAUTICA'), ('White Solid', False))
        self.assertEqual(ia.color_name({'RONASU201SLS': {'color': 'Red'}}, cmap, 'RONASU201SLS-V', 'NAUTICA'),
                         ('Red', True))
        self.assertEqual(ia.color_name({}, cmap, 'BUNMPP002RFS', 'NICOLE'), ('White Navy Print', False))
        self.assertEqual(ia.color_name({}, {}, 'ZZZ', 'NAUTICA'), ('', False))

    def test_group_key(self):
        self.assertEqual(ia.group_key('TMDKPK001SLS 15-15.5 32'), 'TMDKPK001SLS')
        self.assertEqual(ia.group_key('BUJNAW539SLP-V'), 'BUJNAW539SLP-V')
        self.assertEqual(ia.group_key('SSVDR001-SZ10'), 'SSVDR001')
        self.assertEqual(ia.group_key('BEN-22-116'), 'BEN-22-116')

    def test_category_label(self):
        self.assertEqual(ia.category_label('TMNASU201SLS', 'NAUTICA'), 'Dress shirt')
        self.assertEqual(ia.category_label('ROGBPO036SSM', 'BEENE'), 'Sportswear')
        self.assertEqual(ia.category_label('ROUSDPP12SRS', 'USPA'), 'Pants')
        self.assertEqual(ia.category_label('CTH01', 'CHAPS'), 'Accessory')
        self.assertEqual(ia.category_label('BUGBRPB05BRV', 'BEENE'), 'Blazer / Vest')


class TestResolvePhoto(unittest.TestCase):
    def L(self, **kw):
        d = {'ovr': {}, 'brand': {}, 'sync': {}, 'swatch': {}}
        d.update(kw)
        return d

    def test_exact_card_photo_wins(self):
        L = self.L(ovr=listing(['TMNASU201SLS.JPG']), brand=listing(['NAUTICA/NA_201.jpg']))
        ph = ia.resolve_photo({'sku': 'TMNASU201SLS'}, 'NAUTICA', {}, L, set())
        self.assertEqual(ph['key'], 'ALL INVENTORY Photos/STYLE OVERRIDES/TMNASU201SLS.JPG')

    def test_empty_file_is_ignored(self):
        L = self.L(brand={'US POLO/US_224.jpg': ('d41d8cd98f00b204e9800998ecf8427e', 0)},
                   sync=listing(['US_224.jpg']))
        ph = ia.resolve_photo({'sku': 'ROUSPC224SLS'}, 'USPA', {}, L, {'US_224'})
        self.assertEqual(ph['key'], 'ALL INVENTORY Photos/DROPBOX_SYNC/US_224.jpg')

    def test_size_row_uses_base_style_photo(self):
        L = self.L(ovr=listing(['SSVDR001.jpg']), brand=listing(['Von Dutch/VD_001.jpg']))
        ph = ia.resolve_photo({'sku': 'SSVDR001-SZ10'}, 'VD', {}, L, set())
        self.assertEqual(ph['key'], 'ALL INVENTORY Photos/STYLE OVERRIDES/SSVDR001.jpg')

    def test_sync_needs_the_dropbox_index(self):
        L = self.L(sync=listing(['NA_201.jpg']))
        self.assertIsNone(ia.resolve_photo({'sku': 'TMNASU201SLS'}, 'NAUTICA', {}, L, set()))
        self.assertIsNotNone(ia.resolve_photo({'sku': 'TMNASU201SLS'}, 'NAUTICA', {}, L, {'NA_201'}))

    def test_non_shirts_never_borrow_the_brand_photo(self):
        L = self.L(brand=listing(['BEENE/GB_036.jpg']), sync=listing(['GB_036.jpg']))
        self.assertIsNone(ia.resolve_photo({'sku': 'ROGBPO036SSM'}, 'BEENE', {}, L, {'GB_036'}))

    def test_sportswear_server_photo(self):
        ph = ia.resolve_photo({'sku': 'ROGBPO036SSM'}, 'BEENE', {}, self.L(), set(),
                              server_probe=lambda s, b: b'img' if s == 'ROGBPO036SSM' else None)
        self.assertEqual(ph['source'], 'server photo')
        self.assertEqual(ph['etag'], ia.hashlib.md5(b'img').hexdigest())

    def test_accessory_and_swatch(self):
        L = self.L(sync=listing(['CTH03_V.jpg']))
        self.assertEqual(ia.resolve_photo({'sku': 'CTH03-V'}, 'CHAPS', {}, L, set())['key'],
                         'ALL INVENTORY Photos/DROPBOX_SYNC/CTH03_V.jpg')
        L = self.L(swatch=listing(['NA_024.jpg']))
        self.assertEqual(ia.resolve_photo({'sku': 'TMNASU024SLS'}, 'NAUTICA', {}, L, set())['source'], 'swatch')

    def test_accessory_variant_never_borrows_the_base_file(self):
        # resolveImage probes CTH02_V only (review, Oct 15 2026)
        L = self.L(sync=listing(['CTH02.jpg']))
        self.assertIsNone(ia.resolve_photo({'sku': 'CTH02-V'}, 'CHAPS', {}, L, set()))
        L = self.L(sync=listing(['CTH02.jpg']), swatch=listing(['CH_002.jpg']))
        self.assertEqual(ia.resolve_photo({'sku': 'CTH02-V'}, 'CHAPS', {}, L, set())['source'], 'swatch')

    def test_swatch_keys_come_from_the_full_sku(self):
        L = self.L(swatch=listing(['BE_038.jpg']))
        self.assertEqual(ia.resolve_photo({'sku': 'BEN-22-38'}, 'BEN', {}, L, set())['source'], 'swatch')
        L = self.L(swatch=listing(['NA_201.jpg']))
        self.assertIsNone(ia.resolve_photo({'sku': 'AMNASU201SLS-1414.53233'}, 'NAUTICA', {}, L, set()))

    def test_raw_feed_brand_drives_the_folder(self):
        # NM cards look in PHOTOS INVENTORY/NM/ on the desktop, not NICOLE MILLER/
        L = self.L(brand=listing(['NICOLE MILLER/NM_050.jpg']), sync=listing(['NM_050.jpg']))
        ph = ia.resolve_photo({'sku': 'BUNMPP050RFS'}, 'NM', {}, L, {'NM_050'})
        self.assertEqual(ph['key'], 'ALL INVENTORY Photos/DROPBOX_SYNC/NM_050.jpg')

    def test_override_record_hosted_url(self):
        ovr = {'TMBEPU545SLS': {'image': 'https://x.s3.us-east-2.amazonaws.com/ALL+INVENTORY+Photos/STYLE+OVERRIDES/TMBEPU545SLS.jpg'}}
        L = self.L(ovr=listing(['TMBEPU545SLS.jpg']))
        ph = ia.resolve_photo({'sku': 'TMBEPU545SLS'}, 'BEN', ovr, L, set())
        self.assertEqual(ph['source'], 'card photo')
        ph = ia.resolve_photo({'sku': 'TMBEPU545SLS'}, 'BEN', {'TMBEPU545SLS': {'image': 'data:image/jpeg;base64,AAAA'}}, L, set())
        self.assertEqual(ph['source'], 'override record')


class TestNamesAndResults(unittest.TestCase):
    def test_names_conflict(self):
        self.assertFalse(ia.names_conflict('Blue Pinstriepe', 'Blue Pinstripe'))
        self.assertFalse(ia.names_conflict('Grey & Black Check', 'Black Grnd Grey Check'))
        self.assertFalse(ia.names_conflict('Lt Pink Solid / Peacoat Horse', 'Peacoat Horse'))
        self.assertTrue(ia.names_conflict('Black Solid Diamond', 'White Grnd Red Bold Stripe'))
        self.assertTrue(ia.names_conflict('Estate Blue Solid', 'White Grnd Micro Blue Gingham'))
        self.assertTrue(ia.names_conflict('Black Herringbone Dobby', 'Black Jacquard Stripe Dobby'))
        self.assertTrue(ia.names_conflict('Black Solid', 'White Solid'))

    def test_brand_names_cover_every_brand(self):
        self.assertEqual(set(ia.BRAND_FULL), ia.BRAND_MAPPING_KEYS)
        self.assertEqual(ia.BRAND_FULL['BLO'], 'Bloomingdales Private Label')

    def test_seen_text_drops_placeholders(self):
        self.assertEqual(ia.seen_text('white', 'none', 'solid'), 'white (solid)')
        self.assertEqual(ia.seen_text('navy', 'light blue', 'check'), 'navy, light blue (check)')
        self.assertEqual(ia.clean_seen('white, n/a (solid)'), 'white (solid)')

    def test_result_labels(self):
        self.assertEqual(ia.result_label({'verdict': 'mismatch', 'severity': 'high'}), 'Mismatch (high)')
        self.assertEqual(ia.result_label({'verdict': 'unclear', 'kind': 'blank'}), 'Blank photo')
        self.assertEqual(ia.result_label({'verdict': 'unclear', 'kind': 'not_product'}), 'Logo or sketch, not the item')
        self.assertEqual(ia.result_label({'verdict': 'unclear'}), 'Check by eye')
        self.assertEqual(ia.result_label({'verdict': 'match'}), 'OK')

    def test_us_copy(self):
        self.assertEqual(ia.us_copy('Colourblock — colour'), 'Colorblock, color')
        self.assertEqual(ia.us_copy('S – M'), 'S-M')

    def test_fingerprint_groups(self):
        a = ia.fingerprint(jpeg((200, 30, 30), stripes=True))
        b = ia.fingerprint(jpeg((200, 30, 30), stripes=True))
        c = ia.fingerprint(jpeg((20, 30, 200)))
        g = ia.photo_groups({'a': a, 'b': b, 'c': c})
        self.assertEqual(g['a'], g['b'])
        self.assertNotEqual(g['a'], g['c'])


class FakeBlock:
    def __init__(self, text):
        self.type = 'text'
        self.text = text


class FakeResp:
    def __init__(self, payload, stop='end_turn'):
        self.content = [types.SimpleNamespace(type='thinking', thinking=''), FakeBlock(json.dumps(payload))]
        self.stop_reason = stop
        self.stop_details = None


class FakeClient:
    def __init__(self, decide=None, stop='end_turn'):
        self.calls = []
        self.decide = decide or (lambda p: ('match', 'none'))
        self.stop = stop
        self.beta = types.SimpleNamespace(messages=types.SimpleNamespace(create=self._create))

    def with_options(self, **kw):
        return self

    def _create(self, **kw):
        self.calls.append(kw)
        photos = []
        for block in kw['messages'][0]['content']:
            if block['type'] == 'text' and 'id=' in block['text']:
                pid = block['text'].split('id=')[1].split(' ')[0]
                color = block['text'].split('catalog color name: "')[1].split('"')[0]
                v, sev = self.decide(color)
                photos.append({'id': pid, 'seen_main_color': 'white', 'seen_other_colours': '',
                               'seen_other_colors': '', 'seen_pattern': 'solid', 'photo_problem': 'none',
                               'verdict': v, 'severity': sev, 'reason': 'Name says colour — but it is white.'})
        return FakeResp({'photos': photos}, self.stop)


class TestVision(unittest.TestCase):
    def batch(self):
        return [{'id': 'p1', 'color': 'Black Solid', 'hint': 'Dress shirt', 'jpeg': jpeg(), 'etag': 'e1'},
                {'id': 'p2', 'color': 'White Solid', 'hint': 'Dress shirt', 'jpeg': jpeg(), 'etag': 'e2'}]

    def test_request_shape_and_parse(self):
        cl = FakeClient(decide=lambda c: ('mismatch', 'high') if 'Black' in c else ('match', 'none'))
        out = ia.vision_batch(cl, self.batch(), '2026-10-19')
        kw = cl.calls[0]
        self.assertEqual(kw['model'], ia.MODEL)
        self.assertEqual(kw['output_config']['effort'], 'low')
        self.assertEqual(kw['output_config']['format']['type'], 'json_schema')
        self.assertEqual(kw['fallbacks'], 'default')
        self.assertIn('server-side-fallback-2026-07-01', kw['betas'])
        self.assertNotIn('tool_choice', kw)
        self.assertEqual(sum(1 for b in kw['messages'][0]['content'] if b['type'] == 'image'), 2)
        self.assertEqual(out['p1']['verdict'], 'mismatch')
        self.assertEqual(out['p1']['severity'], 'high')
        self.assertEqual(out['p2']['severity'], '')
        self.assertNotIn('—', out['p1']['reason'])
        self.assertIn('color', out['p1']['reason'])

    def test_refusal_returns_nothing(self):
        self.assertEqual(ia.vision_batch(FakeClient(stop='refusal'), self.batch(), 'd'), {})

    def test_swatch_is_marked_and_never_not_product(self):
        cl = FakeClient()
        orig = cl._create

        def create(**kw):
            r = orig(**kw)
            data = json.loads(r.content[1].text)
            for o in data['photos']:
                o['photo_problem'], o['verdict'] = 'not_product', 'unclear'
            r.content[1].text = json.dumps(data)
            return r
        cl.beta.messages.create = create
        b = self.batch()
        b[0]['source'] = 'swatch'
        out = ia.vision_batch(cl, b, 'd')
        self.assertIn('photo type: fabric swatch', cl.calls[0]['messages'][0]['content'][0]['text'])
        self.assertNotIn('fabric swatch', cl.calls[0]['messages'][0]['content'][2]['text'])
        self.assertEqual(out['p1']['kind'], '')
        self.assertEqual(out['p2']['kind'], 'not_product')


class TestWorkbook(unittest.TestCase):
    def test_columns_exactly(self):
        import openpyxl
        x = ia.write_workbook(
            [{'brand': 'Chaps', 'style': 'BUCHOX034SLB', 'color': 'Buttermilk Yellow Solid', 'category': 'Dress shirt'}],
            [{'result': 'Mismatch (high)', 'brand': 'Chaps', 'styles': 'BUCHOX020TBS', 'color': 'Red Singe Solid',
              'seen': 'light blue (solid)', 'why': 'Plainly light blue.', 'link': 'https://e.x/a.jpg', 'thumb': jpeg()}])
        wb = openpyxl.load_workbook(io.BytesIO(x))
        self.assertEqual(wb.sheetnames, ['Missing Images', 'Image vs Color Check'])
        self.assertEqual([c.value for c in wb['Missing Images'][1]], ['Brand', 'Style Number', 'Color Name', 'Category'])
        self.assertEqual([c.value for c in wb['Image vs Color Check'][1]],
                         ['Photo', 'Result', 'Brand', 'Style Numbers', 'Color Name on the Catalog',
                          'What the Photo Shows', 'Why Flagged', 'Photo Link'])
        self.assertEqual([c.value for c in wb['Missing Images'][2]], ['Chaps', 'BUCHOX034SLB', 'Buttermilk Yellow Solid', 'Dress shirt'])
        ws = wb['Image vs Color Check']
        self.assertEqual(ws['B2'].value, 'Mismatch (high)')
        self.assertEqual(ws['H2'].hyperlink.target, 'https://e.x/a.jpg')
        self.assertEqual(len(ws._images), 1)


def make_app(s3, items, ovr=None, cmap=None, dbx=('ZZ_999',), ident=None, client=None, server=None):
    from flask import Flask
    app = Flask('t')
    api = ia.register_image_audit_routes(
        app, get_s3=lambda: s3, s3_bucket='b', region='us-east-2',
        items=lambda: items, overrides=lambda: dict(ovr or {}), dropbox_codes=lambda: set(dbx),
        color_map=lambda: dict(cmap or {}), server_photo=server or (lambda s, b: None),
        anthropic_client=lambda: client or FakeClient(), identity=lambda: ident if ident is not None else {'tier': 'machine'},
        off_loop=None, is_leader=lambda: True, start_scheduler=False)
    return app, api


class TestBuildAndRoutes(unittest.TestCase):
    def setUp(self):
        P = 'ALL INVENTORY Photos/'
        self.s3 = FakeS3({
            P + 'STYLE OVERRIDES/TMCHPT554SLP.jpg': jpeg((240, 240, 240)),
            P + 'PHOTOS INVENTORY/NAUTICA/NA_201.jpg': jpeg((10, 10, 10)),
            P + 'PHOTOS INVENTORY/NAUTICA/NA_202.jpg': jpeg((10, 10, 10)),   # same picture, other name
        })
        self.items = [{'sku': 'TMCHPT554SLP', 'brand': 'CHAPS'}, {'sku': 'TMNASU201SLS', 'brand': 'NAUTICA'},
                      {'sku': 'TMNASU202SLS', 'brand': 'NAUTICA'}, {'sku': 'TMNASU203SLS', 'brand': 'NAUTICA'},
                      {'sku': 'TMNASU203SLS-L', 'brand': 'NAUTICA'},
                      {'sku': 'TMCHPT554SLP-M', 'brand': 'CHAPS'}]
        self.cmap = {'TMCHPT554SLP': 'White Grnd Red & Blue Check', 'NA_201': 'Black Solid', 'NA_202': 'White Solid',
                     'NA_203': 'Navy Solid'}

    def test_build_end_to_end(self):
        import openpyxl
        cl = FakeClient(decide=lambda c: ('mismatch', 'high') if 'Check' in c else ('match', 'none'))
        app, api = make_app(self.s3, self.items, cmap=self.cmap, client=cl)
        xlsx, summary = api['build']('2026-10-19')
        wb = openpyxl.load_workbook(io.BytesIO(xlsx))
        miss = [[c.value for c in r] for r in wb['Missing Images'].iter_rows(min_row=2)]
        self.assertEqual(miss, [['Nautica', 'TMNASU203SLS', 'Navy Solid', 'Dress shirt']])   # sizes fold in
        rows = [[c.value for c in r] for r in wb['Image vs Color Check'].iter_rows(min_row=2)]
        results = sorted(r[1] for r in rows)
        self.assertEqual(results, ['Mismatch (high)', 'Same photo, different colors', 'Same photo, different colors'])
        styles = [r[3] for r in rows if r[1] == 'Mismatch (high)']
        self.assertEqual(styles, ['TMCHPT554SLP'])           # the -M size row folds into its style
        self.assertEqual(summary['missing_colors'], 1)
        self.assertEqual(summary['checked_new'], 3)
        self.assertEqual(summary['unchecked'], 0)
        verd = json.loads(self.s3.objects[ia.VERDICTS_KEY])
        self.assertEqual(len(verd), 3)
        # second run sends nothing new to the model
        cl.calls.clear()
        api['build']('2026-10-26')
        self.assertEqual(cl.calls, [])

    def test_export_route_and_auth(self):
        app, api = make_app(self.s3, self.items, cmap=self.cmap)
        c = app.test_client()
        api['run_build']('2026-10-19', 'test')
        st = json.loads(self.s3.objects[ia.STATUS_KEY])
        self.assertEqual(st['state'], 'ready')
        orig = ia.datetime

        class FixedDT(orig):
            @classmethod
            def now(cls, tz=None):
                return orig(2026, 10, 19, 10, 0, tzinfo=tz)
        ia.datetime = FixedDT
        try:
            r = c.get('/export-image-audit')
            self.assertEqual(r.status_code, 200)
            self.assertEqual(r.data[:2], b'PK')
            self.assertIn('flagged', json.loads(r.headers['X-Image-Audit-Summary']))
            self.assertEqual(c.get('/export-image-audit?format=summary').json['state'], 'ready')
        finally:
            ia.datetime = orig

    def test_routes_refuse_without_machine_or_admin(self):
        app, _ = make_app(self.s3, self.items, cmap=self.cmap, ident={'tier': 'staff', 'is_admin': False})
        c = app.test_client()
        self.assertEqual(c.get('/export-image-audit').status_code, 401)
        self.assertEqual(c.post('/admin/image-audit/build').status_code, 401)
        self.assertEqual(c.post('/admin/image-audit/verdicts', json={}).status_code, 401)
        self.assertEqual(c.get('/admin/image-audit/status').status_code, 401)

    def test_verdict_merge(self):
        app, _ = make_app(self.s3, self.items, cmap=self.cmap)
        r = app.test_client().post('/admin/image-audit/verdicts', json={
            'verdicts': {'abc': {'verdict': 'match'}}, 'photos': {'e1': {'dh': '0' * 64, 'th': '00' * 192}, 'bad': {}}})
        self.assertEqual(r.json, {'verdicts': 1, 'photos': 1, 'rejected': 1})
        r = app.test_client().post('/admin/image-audit/verdicts', json={'verdicts': {'x': {'verdict': 'bogus'}}})
        self.assertEqual(r.json, {'rejected': 1})

    def test_dead_build_does_not_block(self):
        import time as _t
        app, api = make_app(self.s3, self.items, cmap=self.cmap)
        c = app.test_client()
        day = ia.datetime.now(ia.ET).date().isoformat()
        # a live build elsewhere (fresh heartbeat) blocks a second one
        self.s3.objects[ia.STATUS_KEY] = json.dumps({'state': 'building', 'date': day, 'started_at': _t.time() - 3600,
                                                     'heartbeat_at': _t.time() - 30}).encode()
        self.assertEqual(c.post('/admin/image-audit/build').status_code, 409)
        # a dead one (no heartbeat for 11 minutes) does not
        self.s3.objects[ia.STATUS_KEY] = json.dumps({'state': 'building', 'date': day, 'started_at': _t.time() - 900,
                                                     'heartbeat_at': _t.time() - 660}).encode()
        r = c.post('/admin/image-audit/build')
        self.assertEqual(r.status_code, 200)
        for _ in range(200):
            if json.loads(self.s3.objects[ia.STATUS_KEY]).get('state') == 'ready':
                break
            _t.sleep(0.05)
        self.assertEqual(json.loads(self.s3.objects[ia.STATUS_KEY])['state'], 'ready')

    def test_half_loaded_worker_refuses_to_build(self):
        app, api = make_app(self.s3, self.items, cmap={})
        with self.assertRaises(RuntimeError):
            api['build']('2026-10-19')
        app, api = make_app(self.s3, self.items, cmap=self.cmap, dbx=())
        with self.assertRaises(RuntimeError):
            api['build']('2026-10-19')

    def test_unreadable_cache_is_never_rewritten(self):
        app, api = make_app(self.s3, self.items, cmap=self.cmap)
        api['build']('2026-10-19')
        before = dict(self.s3.objects)
        orig = self.s3.get_object

        def flaky(Bucket, Key):
            if Key == ia.VERDICTS_KEY:
                raise TimeoutError('synthetic read timeout')
            return orig(Bucket, Key)
        self.s3.get_object = flaky
        with self.assertRaises(TimeoutError):
            api['build']('2026-10-26')
        self.assertEqual(self.s3.objects[ia.VERDICTS_KEY], before[ia.VERDICTS_KEY])

    def test_bad_bodies_and_status(self):
        app, api = make_app(self.s3, self.items, cmap=self.cmap)
        c = app.test_client()
        self.assertEqual(c.post('/admin/image-audit/verdicts', json=[1, 2]).status_code, 400)
        self.s3.objects[ia.STATUS_KEY] = b'[]'
        self.assertEqual(c.get('/admin/image-audit/status').status_code, 200)

    def test_attempt_cap_stops_on_demand_builds(self):
        app, api = make_app(self.s3, self.items, cmap=self.cmap)
        day = ia.datetime.now(ia.ET).date().isoformat()
        self.s3.objects[ia.STATUS_KEY] = json.dumps({'state': 'building', 'date': day, 'attempts': 4,
                                                     'started_at': 0, 'heartbeat_at': 0}).encode()
        r = app.test_client().get('/export-image-audit')
        self.assertEqual(r.status_code, 503)
        self.assertEqual(r.json['status'], 'error')

    def test_dead_override_file_counts_as_missing(self):
        ovr = {'TMNASU203SLS': {'image': 'https://b.s3.us-east-2.amazonaws.com/ALL+INVENTORY+Photos/STYLE+OVERRIDES/TMNASU203SLS.jpg'}}
        self.assertIsNone(ia.resolve_photo({'sku': 'TMNASU203SLS'}, 'NAUTICA', ovr, dict(EMPTY_L), set()))

    def test_foreign_override_urls_are_never_fetched(self):
        ovr = {'TMNASU203SLS': {'image': 'https://evil.example.com/a.jpg'}}
        app, api = make_app(self.s3, self.items, ovr=ovr, cmap=self.cmap)
        xlsx, summary = api['build']('2026-10-19')
        self.assertEqual(summary['unreadable'], 1)          # shown on the card, but never fetched by us


class TestAppWiring(unittest.TestCase):
    def test_paths_are_on_the_machine_lists(self):
        with open(os.path.join(ROOT, 'app.py'), encoding='utf-8') as f:
            src = f.read()
        block = src.split('_AUTHZ_MACHINE_EXTRA = {', 1)[1].split('}', 1)[0]
        self.assertIn("'/export-image-audit'", block)
        self.assertIn("'/admin/image-audit/status'", block)
        post = src.split("if method == 'POST' and path in ('/admin/claude-uploads'", 1)[1].split('):', 1)[0]
        self.assertIn("'/admin/image-audit/build'", post)
        self.assertIn("'/admin/image-audit/verdicts'", post)
        self.assertIn('register_image_audit_routes(app', src)


if __name__ == '__main__':
    unittest.main()
